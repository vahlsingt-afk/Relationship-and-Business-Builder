#!/usr/bin/env python3
"""
parse_workbook.py — mechanical, sheet-aware xlsx -> JSON extraction for a
Master Account Plan workbook.

RB-2026-08-28: unlike free-text extraction (fabrication risk, guarded
heavily elsewhere this session), this document's sheets have known, labeled,
tabular columns -- parsing a known column header into a known field is
deterministic, the same trust tier as the CSV/JSON auto-decode path already
in server.py's ingest pipeline, not LLM inference over prose. No model
judgment involved anywhere in this module.

Real reference file inspected 2026-08-28:
system/inbox/user_artifacts/Worldpay_Master_Account_Plan_2026-08-14.xlsx --
10 sheets: Executive Summary, Dashboard, Ranked Portfolio, Stack
Intelligence, Evidence Ledger, Conflict Register, RM Portfolio, Scoring
Model, Worldpay Source Snapshot, Change Log.
"""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any, Optional

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


def _normalize_header(h: Any) -> str:
    s = re.sub(r"[^a-z0-9]+", "_", str(h).strip().lower())
    return s.strip("_") or "col"


def _parse_table_sheet(ws) -> list[dict]:
    """Generic: the first row with 2+ non-empty cells is the header row
    (title/subtitle rows above it typically have exactly one cell); every
    subsequent non-empty row becomes one dict keyed by normalized header.
    Unknown/extra columns are kept, never dropped."""
    rows = list(ws.iter_rows(values_only=True))
    headers: list[str] = []
    header_idx: Optional[int] = None
    for i, row in enumerate(rows):
        non_empty = [c for c in row if c is not None and str(c).strip() != ""]
        if len(non_empty) >= 2:
            headers = [_normalize_header(c) if c is not None and str(c).strip() else f"col_{j}"
                       for j, c in enumerate(row)]
            header_idx = i
            break
    if header_idx is None:
        return []
    out: list[dict] = []
    for row in rows[header_idx + 1:]:
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        entry: dict = {}
        for j, val in enumerate(row):
            if val is None:
                continue
            key = headers[j] if j < len(headers) else f"col_{j}"
            entry[key] = val.strip() if isinstance(val, str) else val
        if entry:
            out.append(entry)
    return out


def _split_list_field(val: Any) -> list[str]:
    """Real data uses ';' for some list-shaped columns (Key Vendors) and
    ',' for others (Top Ranked Accounts) -- detect which delimiter the cell
    actually uses rather than assuming one."""
    if not val:
        return []
    text = str(val)
    delimiter = ";" if ";" in text else ","
    return [p.strip() for p in text.split(delimiter) if p.strip()]


_RANKED_PORTFOLIO_DIMENSIONS = (
    "strategic_value", "active_trigger", "access", "lifecycle",
    "deployment_health", "whitespace", "coordination",
)

_RANKED_PORTFOLIO_KEEP = {
    "rank", "priority", "account", "rm", "tier", "locations", "score",
    "trigger_current_state", "lifecycle_posture", "health",
    "opportunity_type", "immediate_next_action", "key_vendors",
    "field_intelligence", "evidence_confidence",
} | set(_RANKED_PORTFOLIO_DIMENSIONS)

_DIMENSION_WEIGHTS = {
    "strategic_value": 0.20,
    "active_trigger": 0.20,
    "access": 0.15,
    "lifecycle": 0.15,
    "deployment_health": 0.10,
    "whitespace": 0.15,
    "coordination": 0.05,
}


def _score_from_dimensions(row: dict) -> Optional[float]:
    """Rebuild a blank formula score from the workbook's published model.

    Some valid xlsx files contain formulas without cached results. openpyxl's
    data_only mode then returns None even though every dimension is present.
    This deterministic fallback preserves the workbook's stated formula; it
    does not infer or change any human-owned dimension rating.
    """
    try:
        values = {key: float(row[key]) for key in _DIMENSION_WEIGHTS}
    except (KeyError, TypeError, ValueError):
        return None
    score = sum((values[key] / 5.0) * weight for key, weight in _DIMENSION_WEIGHTS.items()) * 100
    return round(score, 1)


def _shape_ranked_portfolio(rows: list[dict]) -> list[dict]:
    shaped = []
    for row in rows:
        dimension_scores = {d: row.get(d) for d in _RANKED_PORTFOLIO_DIMENSIONS if d in row}
        extra_fields = {k: v for k, v in row.items() if k not in _RANKED_PORTFOLIO_KEEP}
        shaped.append({
            "rank": row.get("rank"),
            "priority": row.get("priority"),
            "account_name": row.get("account"),
            "rm_name": row.get("rm"),
            "tier": row.get("tier"),
            "locations": row.get("locations"),
            "score": row.get("score") if row.get("score") is not None else _score_from_dimensions(row),
            "dimension_scores": dimension_scores,
            "trigger_current_state": row.get("trigger_current_state"),
            "lifecycle_posture": row.get("lifecycle_posture"),
            "health": row.get("health"),
            "opportunity_type": row.get("opportunity_type"),
            "immediate_next_action": row.get("immediate_next_action"),
            "key_vendors": _split_list_field(row.get("key_vendors")),
            "field_intelligence": row.get("field_intelligence"),
            "evidence_confidence": row.get("evidence_confidence"),
            "linked_blue_sheet_slug": None,        # resolved later by create_plan.py
            "linked_account_research_slug": None,  # resolved later by create_plan.py
            "extra_fields": extra_fields,
        })
    return shaped


_RM_PORTFOLIO_KEEP = {
    "rm", "opportunity_accounts", "p1_accounts", "average_score",
    "top_ranked_accounts", "trigger_summary", "workshop_focus", "required_output",
}


def _shape_rm_portfolios(rows: list[dict]) -> list[dict]:
    shaped = []
    for row in rows:
        extra_fields = {k: v for k, v in row.items() if k not in _RM_PORTFOLIO_KEEP}
        shaped.append({
            "rm_name": row.get("rm"),
            "opportunity_accounts_count": row.get("opportunity_accounts"),
            "p1_accounts_count": row.get("p1_accounts"),
            "average_score": row.get("average_score"),
            "top_ranked_accounts": _split_list_field(row.get("top_ranked_accounts")),
            "trigger_summary": row.get("trigger_summary"),
            "workshop_focus": row.get("workshop_focus"),
            "required_output": row.get("required_output"),
            "extra_fields": extra_fields,
        })
    return shaped


def parse_workbook(content: bytes) -> dict:
    """Returns a dict of extracted sheet data, ready for create_plan.py to
    version and persist. Raises ValueError if openpyxl can't open the file
    or the required sheets are missing (mirrors dataset_classifier's own
    required_sheets check -- this is a second, independent confirmation,
    never trust classification alone before a real write)."""
    if openpyxl is None:
        raise RuntimeError("openpyxl is not installed")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"could not open workbook: {exc}") from exc

    sheet_names = set(wb.sheetnames)
    required = {"Ranked Portfolio", "RM Portfolio"}
    missing = required - sheet_names
    if missing:
        raise ValueError(f"workbook is missing required sheet(s): {sorted(missing)}")

    ranked_portfolio = _shape_ranked_portfolio(_parse_table_sheet(wb["Ranked Portfolio"]))
    rm_portfolios = _shape_rm_portfolios(_parse_table_sheet(wb["RM Portfolio"]))

    dashboard_rows = _parse_table_sheet(wb["Dashboard"]) if "Dashboard" in sheet_names else []
    dashboard = dashboard_rows[0] if dashboard_rows else {}

    stack_intelligence = _parse_table_sheet(wb["Stack Intelligence"]) if "Stack Intelligence" in sheet_names else []
    evidence_ledger = _parse_table_sheet(wb["Evidence Ledger"]) if "Evidence Ledger" in sheet_names else []
    conflict_register = _parse_table_sheet(wb["Conflict Register"]) if "Conflict Register" in sheet_names else []
    scoring_model = _parse_table_sheet(wb["Scoring Model"]) if "Scoring Model" in sheet_names else []
    change_log = _parse_table_sheet(wb["Change Log"]) if "Change Log" in sheet_names else []
    change_log_sheets = [n for n in wb.sheetnames if "change log" in n.lower()]
    if change_log_sheets and not change_log:
        change_log = _parse_table_sheet(wb[change_log_sheets[0]])

    source_snapshot_sheets = [n for n in wb.sheetnames if "source snapshot" in n.lower()]
    source_snapshot = _parse_table_sheet(wb[source_snapshot_sheets[0]]) if source_snapshot_sheets else []

    executive_summary_lines: list[str] = []
    if "Executive Summary" in sheet_names:
        for row in wb["Executive Summary"].iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip() != ""]
            if cells:
                executive_summary_lines.append(" | ".join(cells))

    return {
        "sheet_names": wb.sheetnames,
        "executive_summary_text": "\n".join(executive_summary_lines),
        "dashboard": dashboard,
        "ranked_portfolio": ranked_portfolio,
        "rm_portfolios": rm_portfolios,
        "stack_intelligence": stack_intelligence,
        "evidence_ledger": evidence_ledger,
        "conflict_register": conflict_register,
        "scoring_model": scoring_model,
        "change_log": change_log,
        "source_snapshot": source_snapshot,
    }


def parse_workbook_file(path: Path) -> dict:
    return parse_workbook(Path(path).read_bytes())
