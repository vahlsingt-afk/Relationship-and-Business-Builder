#!/usr/bin/env python3
"""
render.py — targeted in-place xlsx update for a Master Account Plan.

Deliberately NOT a from-scratch multi-tab regeneration (unlike Blue Sheet's
render.py, which is the right approach there because RBB authors Blue
Sheets). Todd's own workbook formatting and tab structure are the real
source of truth here -- this module only edits specific cells/rows after a
safe auto-apply (impact_review.py): appends new Evidence Ledger rows,
refreshes RM Portfolio aggregate columns. Everything else in the workbook
is left untouched.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path
from typing import Optional

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

ENGINE_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(ENGINE_DIR))
import mp_common as common  # noqa: E402

sys.path.insert(0, str(common.ROOT.parent / "system" / "scripts"))
import xlsx_safety  # noqa: E402


def _normalize_header(h) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(h).strip().lower()).strip("_") if h is not None else ""


def _find_header_row(ws) -> tuple[Optional[int], dict]:
    """Returns (1-indexed row number, {normalized_header: 1-indexed col})."""
    for row in ws.iter_rows(min_row=1):
        non_empty = [c for c in row if c.value is not None and str(c.value).strip() != ""]
        if len(non_empty) >= 2:
            headers = {}
            for cell in row:
                if cell.value is not None and str(cell.value).strip():
                    headers[_normalize_header(cell.value)] = cell.column
            return row[0].row, headers
    return None, {}


def append_evidence_rows(vendor_slug: str, new_evidence: list[dict]) -> int:
    """Appends rows to the Evidence Ledger sheet for newly-recorded evidence
    items. `new_evidence` entries match the shape appended to evidence.jsonl
    by impact_review.py. Returns the number of rows appended."""
    if not new_evidence or openpyxl is None:
        return 0
    vendor_dir = common.ROOT / "vendors" / vendor_slug
    current_dir = vendor_dir / "current"
    xlsx_files = list(current_dir.glob("*.xlsx"))
    if not xlsx_files:
        return 0
    path = xlsx_files[0]

    wb = openpyxl.load_workbook(path)
    if "Evidence Ledger" not in wb.sheetnames:
        return 0
    ws = wb["Evidence Ledger"]
    header_row_num, cols = _find_header_row(ws)
    if header_row_num is None:
        return 0

    next_row = ws.max_row + 1
    appended = 0
    for ev in new_evidence:
        row_values = {
            "account": ev.get("account_name"),
            "event_type": ev.get("signal_type"),
            "event_date": ev.get("event_at"),
            "evidence_summary": ev.get("signal_summary"),
            "confidence": "RBB-detected",
            "research_date": common.today(),
        }
        for key, col in cols.items():
            if key in row_values and row_values[key] is not None:
                # RB-SECURITY-2026-09-05: account_name/signal_summary can
                # carry externally-influenced text (a press-release/vendor-
                # relationship signal). Same formula/CSV-injection guard
                # (CWE-1236) already applied to Blue Sheets.
                ws.cell(row=next_row, column=col, value=xlsx_safety.sanitize_cell_value(row_values[key]))
        next_row += 1
        appended += 1

    if appended:
        wb.save(path)
    return appended


def refresh_rm_portfolio_sheet(vendor_slug: str, rm_portfolios: list[dict]) -> int:
    """Updates RM Portfolio's aggregate columns (Opportunity Accounts, P1
    Accounts, Average Score) in place to match the recomputed
    rm_portfolios.json. Returns the number of rows updated."""
    if not rm_portfolios or openpyxl is None:
        return 0
    vendor_dir = common.ROOT / "vendors" / vendor_slug
    current_dir = vendor_dir / "current"
    xlsx_files = list(current_dir.glob("*.xlsx"))
    if not xlsx_files:
        return 0
    path = xlsx_files[0]

    wb = openpyxl.load_workbook(path)
    if "RM Portfolio" not in wb.sheetnames:
        return 0
    ws = wb["RM Portfolio"]
    header_row_num, cols = _find_header_row(ws)
    if header_row_num is None or "rm" not in cols:
        return 0
    rm_col = cols["rm"]

    by_rm = {rm.get("rm_name"): rm for rm in rm_portfolios}
    updated = 0
    for row_num in range(header_row_num + 1, ws.max_row + 1):
        rm_name = ws.cell(row=row_num, column=rm_col).value
        rm = by_rm.get(rm_name)
        if not rm:
            continue
        if "opportunity_accounts" in cols and rm.get("opportunity_accounts_count") is not None:
            ws.cell(row=row_num, column=cols["opportunity_accounts"], value=rm["opportunity_accounts_count"])
        if "p1_accounts" in cols and rm.get("p1_accounts_count") is not None:
            ws.cell(row=row_num, column=cols["p1_accounts"], value=rm["p1_accounts_count"])
        if "average_score" in cols and rm.get("average_score") is not None:
            ws.cell(row=row_num, column=cols["average_score"], value=rm["average_score"])
        updated += 1

    if updated:
        wb.save(path)
    return updated
