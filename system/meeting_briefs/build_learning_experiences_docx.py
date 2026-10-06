from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "2026-10-02-learning-from-your-previous-experiences.md"
OUTPUT = ROOT / "Learning-from-Your-Previous-Experiences-Meeting-Brief.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=140, bottom=110, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def remove_paragraph_border(paragraph_or_style):
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def new_numbering_instance(doc, abstract_num_id=7):
    numbering = doc.part.numbering_part.element
    existing = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    num_id = max(existing, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_num_id))
    num.append(abstract)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.get_or_add_numPr()
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)


def add_inline(paragraph, text, size=10.8, color="222222"):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Aptos Mono"
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos Mono")
            run.font.size = Pt(9.2)
        else:
            run = paragraph.add_run(part)
        run.font.name = run.font.name or "Aptos"
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), run.font.name)
        if run.font.size is None:
            run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)


def style_document(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = RGBColor(34, 34, 34)
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(7)
    remove_paragraph_border(title)

    for name, size, before, after in (
        ("Heading 1", 16.5, 15, 6),
        ("Heading 2", 13.2, 11, 4),
        ("Heading 3", 11.2, 8, 3),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        st = styles[name]
        st.font.name = "Aptos"
        st.font.size = Pt(10.6)
        if name == "List Number":
            st.paragraph_format.left_indent = Inches(0.4)
            st.paragraph_format.first_line_indent = Inches(-0.24)
        else:
            st.paragraph_format.left_indent = Inches(0.25)
            st.paragraph_format.first_line_indent = Inches(-0.18)
        st.paragraph_format.space_after = Pt(3.5)


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.85)
    table.columns[1].width = Inches(4.65)
    table.rows[0].cells[0].width = Inches(1.85)
    table.rows[0].cells[1].width = Inches(4.65)
    table.rows[0].cells[0].text = rows[0][0]
    table.rows[0].cells[1].text = rows[0][1]
    set_repeat_table_header(table.rows[0])
    for cell in table.rows[0].cells:
        set_cell_shading(cell, "1F4E78")
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.name = "Aptos"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for idx, row_data in enumerate(rows[1:], start=1):
        cells = table.add_row().cells
        for j, value in enumerate(row_data):
            cells[j].width = Inches(1.85 if j == 0 else 4.65)
            cells[j].text = value
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cells[j])
            set_cell_margins(cells[j])
            if idx % 2 == 0:
                set_cell_shading(cells[j], "F3F6F9")
            for paragraph in cells[j].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(25, 25, 25)
                    if j == 0:
                        run.font.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("Learning from Your Previous Experiences")
    remove_paragraph_border(title)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(11)
    r = subtitle.add_run("Meeting Brief")
    r.font.name = "Aptos Display"
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(70, 70, 70)

    metadata = [
        ("Meeting", "Friday, October 2, 2026, 1:30-2:30 PM CT"),
        ("Organizer", "Michael Estabrooks"),
        ("Participants", "Todd Vahlsing, Michael Estabrooks, Tracy Gallimore; Michael Schwartz and Ryan Hildebrand invited"),
        ("Purpose", "Translate lessons from PAR and consulting into practical improvements for GP/Genius sales flow, documentation, commercial scoping, pricing, and downstream handoff."),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2.5)
        rr = p.add_run(label + ": ")
        rr.bold = True
        rr.font.name = "Aptos"
        rr.font.size = Pt(10.4)
        vv = p.add_run(value)
        vv.font.name = "Aptos"
        vv.font.size = Pt(10.4)

    doc.add_paragraph()

    i = 7
    active_numbering_id = None
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            active_numbering_id = None
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(line[3:], style="Heading 1")
            keep_with_next(p)
            if line[3:] == "Evidence and cautions":
                p.paragraph_format.page_break_before = True
        elif line.startswith("### "):
            p = doc.add_paragraph(line[4:], style="Heading 2")
            keep_with_next(p)
        elif line.startswith("| "):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            parsed = []
            for t in table_lines:
                cells = [c.strip() for c in t.strip().strip("|").split("|")]
                if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    continue
                parsed.append(cells)
            add_table(doc, parsed)
            continue
        elif re.match(r"^\d+\. ", line):
            if active_numbering_id is None:
                active_numbering_id = new_numbering_instance(doc)
            p = doc.add_paragraph(style="List Number")
            apply_numbering(p, active_numbering_id)
            add_inline(p, re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        else:
            p = doc.add_paragraph()
            add_inline(p, line)
        i += 1

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("Learning from Your Previous Experiences  |  Meeting Brief")
    fr.font.name = "Aptos"
    fr.font.size = Pt(8.5)
    fr.font.color.rgb = RGBColor(100, 100, 100)

    core = doc.core_properties
    core.title = "Learning from Your Previous Experiences Meeting Brief"
    core.subject = "Meeting preparation notes for October 2, 2026"
    core.author = "Todd Vahlsing"
    core.keywords = "meeting brief, sales process, account governance, CPQ, franchisee sales"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
