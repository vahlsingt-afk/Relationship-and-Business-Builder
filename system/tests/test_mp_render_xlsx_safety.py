"""
test_mp_render_xlsx_safety.py

RB-SECURITY-2026-09-05: mp_render.py's append_evidence_rows() writes
account_name/signal_summary into a real Master Account Plan workbook --
both can carry externally-influenced text (a vendor-relationship press
release, a signal summary). Same formula/CSV-injection guard (CWE-1236)
already applied to Blue Sheets.
"""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT_DIR = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT_DIR / "master_account_plans" / "_engine"))
sys.path.insert(0, str(ROOT_DIR / "system" / "scripts"))

import openpyxl  # noqa: E402

import mp_render as render  # noqa: E402


class TestAppendEvidenceRowsSanitizes(unittest.TestCase):
    def _workbook_with_evidence_ledger(self) -> tuple[Path, Path]:
        wb = openpyxl.Workbook()
        wb.active.title = "Evidence Ledger"
        ws = wb.active
        ws.append(("account", "event_type", "event_date", "evidence_summary", "confidence", "research_date"))
        root = Path(tempfile.mkdtemp())
        tmp_dir = root / "vendors" / "test-vendor" / "current"
        tmp_dir.mkdir(parents=True)
        path = tmp_dir / "Test_Vendor_Master_Account_Plan.xlsx"
        wb.save(path)
        return path, root

    def test_signal_summary_formula_trigger_is_sanitized(self):
        path, root = self._workbook_with_evidence_ledger()
        import mp_common as common
        orig_root = common.ROOT
        common.ROOT = root
        try:
            n = render.append_evidence_rows("test-vendor", [{
                "account_name": "=1+1",
                "signal_type": "vendor_relationship_formed",
                "event_at": "2026-09-05",
                "signal_summary": "-2+2 formula-shaped summary text",
            }])
            self.assertEqual(n, 1)
            wb = openpyxl.load_workbook(path)
            ws = wb["Evidence Ledger"]
            values = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
            self.assertIn("'=1+1", values)
            self.assertIn("'-2+2 formula-shaped summary text", values)
            self.assertNotIn("=1+1", values)
        finally:
            common.ROOT = orig_root


if __name__ == "__main__":
    unittest.main()
