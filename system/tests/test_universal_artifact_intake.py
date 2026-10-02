from __future__ import annotations

import base64
import io
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

from system.api import server


def test_generic_user_artifact_is_persisted_mutated_and_receipted(
    tmp_path: Path, monkeypatch
):
    inbox = tmp_path / "system" / "inbox"
    cache = tmp_path / "system" / ".cache"
    inbox.mkdir(parents=True)
    cache.mkdir(parents=True)

    # RB-SECURITY-2026-09-03 made server._auth() fail *closed* (503) when
    # API_KEY is unset, so the old `monkeypatch.setattr(server, "API_KEY",
    # None)` trick to disable auth for these direct-function-call tests now
    # 503s instead of passing through. Bypass _auth() itself instead --
    # same isolation pattern test_vendor_list.py already uses for
    # rbb_chat._auth_flexible.
    monkeypatch.setattr(server, "_auth", lambda *a, **k: None)
    monkeypatch.setattr(server.core, "PROJECT_DIR", tmp_path)
    monkeypatch.setattr(server.core, "INBOX_DIR", inbox)
    monkeypatch.setattr(server.core, "CACHE_DIR", cache)
    monkeypatch.setattr(server, "_load_artifact_registry", lambda: {})
    # RB-2026-08-31: uploadAndIngestFile's "intelligence" pipeline now also
    # persists the full extraction (uploaded_document_store) and registers
    # it (intelligence_index) -- both modules compute their storage paths
    # once at import time from the real core.SYSTEM_DIR, so they need the
    # same explicit isolation as core.PROJECT_DIR/INBOX_DIR/CACHE_DIR above,
    # or this test's tmp-path receipt would try to relative_to() against a
    # real production path.
    monkeypatch.setattr(server.uploaded_document_store, "STORE_DIR", tmp_path / "system" / "uploaded_documents")
    monkeypatch.setattr(server.intelligence_index, "INDEX_PATH", tmp_path / "system" / "intelligence_index.json")
    monkeypatch.setattr(server.intelligence_index, "UPDATE_LOG_PATH", tmp_path / "system" / "intelligence_index_updates.jsonl")
    monkeypatch.setattr(
        server.intelligence_triage,
        "triage_input",
        lambda *args, **kwargs: {
            "triage_id": "TRG-TEST",
            "type_count": 2,
            "noise_only": False,
            "identified_types": [{"intelligence_type": "macro_signal"}],
        },
    )
    monkeypatch.setattr(
        server.intelligence_mutation_engine,
        "run",
        lambda *args, **kwargs: {
            "trust_stats": {
                "total_mutations": 3,
                "confidence": "high",
            },
            "apply_result": {"applied": 2},
        },
    )
    monkeypatch.setattr(
        server.subprocess,
        "run",
        lambda *args, **kwargs: SimpleNamespace(
            returncode=0, stdout='{"results":[]}', stderr=""
        ),
    )
    monkeypatch.setattr(server.daily_brief, "build_report", lambda *args: {"ok": True})
    monkeypatch.setattr(server.publish, "publish_brief", lambda **kwargs: {"ok": True})

    receipt = server.post_ingest_upload(
        server.FileUploadIn(
            filename="../operator-note.md",
            extracted_text="Toast is expanding with a new enterprise relationship signal.",
            source_type="note",
        ),
        x_api_key=None,
    )

    assert receipt["processing_status"] == "success"
    assert receipt["pipeline_run"] == "intelligence"
    assert receipt["delta"]["mutations_generated"] == 3
    assert receipt["delta"]["mutations_applied"] == 2
    assert receipt["freshness_recorded"] is True
    assert receipt["brief_rebuilt"] is True
    assert receipt["file_written"] == "system/inbox/user_artifacts/operator-note.md"
    assert (inbox / "user_artifacts" / "operator-note.md").exists()

    ledger = json.loads((inbox / "user_artifacts.json").read_text(encoding="utf-8"))
    assert ledger["items"][-1]["ingestion_id"] == receipt["ingestion_id"]
    history = (cache / "ingest_upload_history.jsonl").read_text(encoding="utf-8")
    assert receipt["ingestion_id"] in history


def _nsn_workbook_base64() -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "StoreTech"
    ws.append(["NSN", "Field Office", "Market", "Entity Name"])
    for sheet in ["Markets", "FO OTM-STIM", "RFM", "COOP2", "Entity"]:
        wb.create_sheet(sheet)
    buf = io.BytesIO()
    wb.save(buf)
    return base64.b64encode(buf.getvalue()).decode()


def test_recognized_xlsx_workbook_routes_to_micro_graph_pipeline(
    tmp_path: Path, monkeypatch
):
    """RB-DEFECT-065: an NSN-fingerprinted xlsx upload is recognized by
    dataset_classifier and routed to the micro-graph builder automatically,
    instead of falling through to the generic intelligence/text pipeline
    (which can't process xlsx bytes at all).
    """
    inbox = tmp_path / "system" / "inbox"
    cache = tmp_path / "system" / ".cache"
    inbox.mkdir(parents=True)
    cache.mkdir(parents=True)

    monkeypatch.setattr(server, "_auth", lambda *a, **k: None)
    monkeypatch.setattr(server.core, "PROJECT_DIR", tmp_path)
    monkeypatch.setattr(server.core, "INBOX_DIR", inbox)
    monkeypatch.setattr(server.core, "CACHE_DIR", cache)
    monkeypatch.setattr(
        server.subprocess,
        "run",
        lambda *args, **kwargs: SimpleNamespace(
            returncode=0,
            stdout=json.dumps({
                "ok": True,
                "written": True,
                "graph_id": "micro_ecosystem:mcdonalds_us_ops",
                "counts": {"nodes": 42, "edges": 90},
            }),
            stderr="",
        ),
    )
    monkeypatch.setattr(server.daily_brief, "build_report", lambda *args: {"ok": True})
    monkeypatch.setattr(server.publish, "publish_brief", lambda **kwargs: {"ok": True})

    receipt = server.post_ingest_upload(
        server.FileUploadIn(
            filename="NSN Lookup 2026-07 JUL.xlsx",
            content_base64=_nsn_workbook_base64(),
        ),
        x_api_key=None,
    )

    assert receipt["pipeline_run"] == "micro_graph"
    assert receipt["processing_status"] == "success"
    assert receipt["ingest_result"]["dataset_type"] == "mcdonalds_nsn_lookup_workbook"
    assert receipt["delta"]["records_added"] == 42
    assert receipt["delta"]["records_changed"] == 90
    assert receipt["file_written"] == "system/inbox/micro_graph_sources/NSN Lookup 2026-07 JUL.xlsx"
    assert (inbox / "micro_graph_sources" / "NSN Lookup 2026-07 JUL.xlsx").exists()


def test_unrecognized_xlsx_falls_back_to_intelligence_review(tmp_path, monkeypatch):
    """A workbook that doesn't match any known signature (missing sheet
    fingerprint) should NOT be auto-routed to a mutation-writing pipeline —
    it stays on the conservative review-first path instead.
    """
    inbox = tmp_path / "system" / "inbox"
    cache = tmp_path / "system" / ".cache"
    inbox.mkdir(parents=True)
    cache.mkdir(parents=True)

    wb = Workbook()
    wb.active.title = "Sheet1"
    buf = io.BytesIO()
    wb.save(buf)
    content_b64 = base64.b64encode(buf.getvalue()).decode()

    monkeypatch.setattr(server, "_auth", lambda *a, **k: None)
    monkeypatch.setattr(server.core, "PROJECT_DIR", tmp_path)
    monkeypatch.setattr(server.core, "INBOX_DIR", inbox)
    monkeypatch.setattr(server.core, "CACHE_DIR", cache)
    monkeypatch.setattr(server, "_load_artifact_registry", lambda: {})
    monkeypatch.setattr(
        server.intelligence_triage,
        "triage_input",
        lambda *args, **kwargs: {
            "triage_id": "TRG-TEST", "type_count": 0, "noise_only": True, "identified_types": [],
        },
    )
    monkeypatch.setattr(
        server.intelligence_mutation_engine,
        "run",
        lambda *args, **kwargs: {"trust_stats": {"total_mutations": 0}, "apply_result": {"applied": 0}},
    )
    monkeypatch.setattr(server.daily_brief, "build_report", lambda *args: {"ok": True})
    monkeypatch.setattr(server.publish, "publish_brief", lambda **kwargs: {"ok": True})

    receipt = server.post_ingest_upload(
        server.FileUploadIn(filename="random_other.xlsx", content_base64=content_b64),
        x_api_key=None,
    )

    assert receipt["pipeline_run"] == "intelligence"
    assert receipt["file_written"] == "system/inbox/user_artifacts/random_other.xlsx"


def test_truncated_linkedin_zip_content_raises_413_with_watch_folder(monkeypatch):
    """RB 2026-08-26: confirmed live in the actual GPT — the platform can also
    deliver content that decodes successfully but is truncated to a tiny
    fragment (12 bytes seen live, for a real ~1.9MB export). This is
    technically valid non-empty content, so it must be caught by a minimum
    plausible-size check rather than silently 'succeeding' with a useless
    0-processed/unknown_classification result."""
    monkeypatch.setattr(server, "_auth", lambda *a, **k: None)
    tiny_payload_b64 = base64.b64encode(b"PK\x03\x04tiny").decode()

    with pytest.raises(HTTPException) as exc_info:
        server.post_ingest_upload(
            server.FileUploadIn(
                filename="Complete_LinkedInDataExport_08-19-2026.zip",
                content_base64=tiny_payload_b64,
            ),
            x_api_key=None,
        )

    assert exc_info.value.status_code == 413
    assert "system/inbox/linkedin_exports/" in exc_info.value.detail


def test_small_real_file_below_zip_threshold_still_ingests_normally(monkeypatch, tmp_path):
    """Regression guard: the minimum-size check is scoped to known
    large-export filename patterns (.zip/linkedin-ish, .txt|.zip/whatsapp-ish)
    only. A genuinely small, legitimate file of a different type (e.g. a vcf
    contact card, as verified live today) must still ingest normally and not
    get swept into the 413 path."""
    inbox = tmp_path / "system" / "inbox"
    cache = tmp_path / "system" / ".cache"
    inbox.mkdir(parents=True)
    cache.mkdir(parents=True)

    monkeypatch.setattr(server, "_auth", lambda *a, **k: None)
    monkeypatch.setattr(server.core, "PROJECT_DIR", tmp_path)
    monkeypatch.setattr(server.core, "INBOX_DIR", inbox)
    monkeypatch.setattr(server.core, "CACHE_DIR", cache)
    monkeypatch.setattr(server, "_load_artifact_registry", lambda: {})

    import sys

    scripts_dir = Path(server.__file__).resolve().parent.parent / "scripts"
    sys.path.insert(0, str(scripts_dir))
    import contacts_ingest

    monkeypatch.setattr(
        contacts_ingest,
        "ingest_file",
        lambda *args, **kwargs: {"ok": True, "scanned": 1, "mutations": []},
    )
    monkeypatch.setattr(server.daily_brief, "build_report", lambda *args: {"ok": True})
    monkeypatch.setattr(server.publish, "publish_brief", lambda **kwargs: {"ok": True})

    tiny_vcf = b"BEGIN:VCARD\nVERSION:3.0\nFN:Test\nEND:VCARD\n"
    receipt = server.post_ingest_upload(
        server.FileUploadIn(
            filename="tiny_contact.vcf",
            content_base64=base64.b64encode(tiny_vcf).decode(),
        ),
        x_api_key=None,
    )

    assert receipt["pipeline_run"] == "contacts"
    assert receipt["processing_status"] == "success"


def test_corrupted_linkedin_zip_base64_raises_413_with_watch_folder(monkeypatch):
    """RB 2026-08-26: confirmed live in the actual GPT — for a large export,
    the platform doesn't always drop the payload to empty; it can also arrive
    non-empty but corrupted/truncated, failing b64decode. Same underlying
    transfer failure as the empty-payload case, must get the same 413 +
    watch-folder guidance rather than falling through to the generic
    'invalid base64' 400."""
    monkeypatch.setattr(server, "_auth", lambda *a, **k: None)

    with pytest.raises(HTTPException) as exc_info:
        server.post_ingest_upload(
            server.FileUploadIn(
                filename="Complete_LinkedInDataExport_08-19-2026.zip",
                content_base64="not-valid-base64-!!!",
            ),
            x_api_key=None,
        )

    assert exc_info.value.status_code == 413
    assert "system/inbox/linkedin_exports/" in exc_info.value.detail


def test_empty_linkedin_zip_payload_raises_413_with_watch_folder(monkeypatch):
    """RB 2026-08-26: a Custom GPT Action call for a large LinkedIn export ZIP
    can arrive with no content at all (the platform drops the base64 payload
    before it reaches this endpoint). This must fail loudly with a specific,
    actionable message the GPT can relay verbatim — not the generic 400,
    and never a fabricated success."""
    monkeypatch.setattr(server, "_auth", lambda *a, **k: None)

    with pytest.raises(HTTPException) as exc_info:
        server.post_ingest_upload(
            server.FileUploadIn(filename="Complete_LinkedInDataExport_08-19-2026.zip"),
            x_api_key=None,
        )

    assert exc_info.value.status_code == 413
    assert "system/inbox/linkedin_exports/" in exc_info.value.detail
    assert "do not retry" in exc_info.value.detail.lower() or "retry" in exc_info.value.detail.lower()


def test_empty_whatsapp_export_payload_raises_413_with_watch_folder(monkeypatch):
    monkeypatch.setattr(server, "_auth", lambda *a, **k: None)

    with pytest.raises(HTTPException) as exc_info:
        server.post_ingest_upload(
            server.FileUploadIn(filename="WhatsApp Chat with Sarah McAngus.txt"),
            x_api_key=None,
        )

    assert exc_info.value.status_code == 413
    assert "system/inbox/whatsapp_exports/" in exc_info.value.detail


def test_empty_payload_for_unrecognized_filename_still_raises_generic_400(monkeypatch):
    """Regression guard: the new 413 path is scoped to known large-export
    filename patterns only. Any other empty-payload call must keep raising
    the original generic 400 — do not silently widen the 413 fallback to
    every failed upload."""
    monkeypatch.setattr(server, "_auth", lambda *a, **k: None)

    with pytest.raises(HTTPException) as exc_info:
        server.post_ingest_upload(
            server.FileUploadIn(filename="random_notes.txt"),
            x_api_key=None,
        )

    assert exc_info.value.status_code == 400
