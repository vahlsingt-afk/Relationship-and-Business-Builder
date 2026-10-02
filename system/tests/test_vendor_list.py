"""
test_vendor_list.py — RB-2026-09-01.

Coverage for the canonical vendor list Todd asked to download: `listVendors`
(system/api/server.py), ecosystem_intelligence.vendor_export_rows(), and the
xlsx export served via rbb_chat.py's /vendor-list/download route (CSV
initially, switched to xlsx per Todd's same-day follow-up).
"""
from __future__ import annotations

import importlib.util
import json
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = ROOT / "system" / "scripts"

spec = importlib.util.spec_from_file_location("ecosystem_intelligence", SCRIPTS_DIR / "ecosystem_intelligence.py")
ei = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(ei)


def _graph_with_vendors() -> dict:
    return {
        "version": 1, "contract": "rb_ecosystem_intelligence_v1", "last_updated": "2026-09-01",
        "entities": [
            {"id": "brand-blaze-pizza", "name": "Blaze Pizza", "entity_type": "brand",
             "aliases": [], "attributes": {}, "sources": [], "confidence": {}, "domains": ["restaurants"]},
            {"id": "vendor-toast", "name": "Toast", "entity_type": "vendor", "aliases": ["Toast POS"],
             "attributes": {"primary_category": "pos"}, "status": "active",
             "sources": [], "confidence": {}, "domains": ["restaurants"],
             "created_at": "2026-08-01T00:00:00", "updated_at": "2026-08-01T00:00:00"},
            {"id": "vendor-untracked", "name": "UntrackedVendor", "entity_type": "vendor", "aliases": [],
             "attributes": {"primary_category": "inventory"}, "status": "active",
             "sources": [], "confidence": {}, "domains": ["restaurants"],
             "created_at": "2026-08-01T00:00:00", "updated_at": "2026-08-01T00:00:00"},
        ],
        "relationships": [{
            "id": "rel-1", "from_entity_id": "brand-blaze-pizza", "to_entity_id": "vendor-toast",
            "relationship_type": "uses_vendor_for_category", "category": "pos", "status": "active",
        }],
        "signals": [], "sources": [], "assessments": [], "user_relevance": [], "strategic_recommendations": [],
    }


class TestVendorExportRows(unittest.TestCase):
    def setUp(self):
        self._tmpdir = tempfile.mkdtemp()
        self._graph_path = Path(self._tmpdir) / "ecosystem_intelligence.json"
        self._orig_path = ei.core.ECOSYSTEM_INTELLIGENCE_PATH
        ei.core.ECOSYSTEM_INTELLIGENCE_PATH = self._graph_path
        self._graph_path.write_text(json.dumps(_graph_with_vendors()), encoding="utf-8")

        # Isolate the competitor_intelligence cross-reference too --
        # vendor_export_rows() imports competitor_intelligence_common
        # (a real, already-loaded module -- Python caches it in
        # sys.modules, so re-importing inside the function returns the
        # SAME object) and reads its module-level ROOT; without
        # redirecting that too, a fake vendor id that happens to collide
        # with a REAL tracked competitor's real vendor_entity_id (e.g.
        # "vendor-toast") would get flagged is_tracked_competitor from the
        # real production registry, not this test's isolated graph.
        import sys as _sys
        _sys.path.insert(0, str(SCRIPTS_DIR))
        import competitor_intelligence_common as cic
        self._cic = cic
        self._orig_cic_root = cic.ROOT
        comp_root = Path(self._tmpdir) / "competitor_intelligence"
        (comp_root / "_portfolio").mkdir(parents=True)
        (comp_root / "_portfolio" / "competitor_registry.json").write_text(
            json.dumps({"registry": []}), encoding="utf-8",
        )
        cic.ROOT = comp_root

    def tearDown(self):
        ei.core.ECOSYSTEM_INTELLIGENCE_PATH = self._orig_path
        self._cic.ROOT = self._orig_cic_root
        shutil.rmtree(self._tmpdir, ignore_errors=True)

    def test_returns_only_vendor_entities_sorted_by_name(self):
        rows = ei.vendor_export_rows()
        names = [r["name"] for r in rows]
        self.assertEqual(names, sorted(names, key=str.lower))
        self.assertNotIn("Blaze Pizza", names)  # a brand, not a vendor
        self.assertIn("Toast", names)
        self.assertIn("UntrackedVendor", names)

    def test_brand_relationship_count_reflects_real_relationships(self):
        rows = ei.vendor_export_rows()
        toast = next(r for r in rows if r["name"] == "Toast")
        untracked = next(r for r in rows if r["name"] == "UntrackedVendor")
        self.assertEqual(toast["brand_relationship_count"], 1)
        self.assertEqual(untracked["brand_relationship_count"], 0)

    def test_aliases_joined_as_string(self):
        rows = ei.vendor_export_rows()
        toast = next(r for r in rows if r["name"] == "Toast")
        self.assertEqual(toast["aliases"], "Toast POS")

    def test_untracked_vendor_is_not_flagged_as_competitor(self):
        rows = ei.vendor_export_rows()
        untracked = next(r for r in rows if r["name"] == "UntrackedVendor")
        self.assertFalse(untracked["is_tracked_competitor"])

    def test_real_competitor_link_flags_is_tracked_competitor(self):
        # Register "toast" as a tracked competitor pointing at vendor-toast,
        # via this test's own isolated competitor_intelligence root.
        comp_dir = self._cic.ROOT / "competitors" / "toast"
        comp_dir.mkdir(parents=True)
        (comp_dir / "competitor.json").write_text(json.dumps({
            "competitor_id": "comp-toast", "competitor_slug": "toast", "display_name": "Toast",
            "vendor_entity_id": "vendor-toast",
        }), encoding="utf-8")
        (self._cic.ROOT / "_portfolio" / "competitor_registry.json").write_text(json.dumps({
            "registry": [{"competitor_slug": "toast", "display_name": "Toast"}],
        }), encoding="utf-8")

        rows = ei.vendor_export_rows()
        toast = next(r for r in rows if r["name"] == "Toast")
        untracked = next(r for r in rows if r["name"] == "UntrackedVendor")
        self.assertTrue(toast["is_tracked_competitor"])
        self.assertFalse(untracked["is_tracked_competitor"])


class TestListVendorsEndpoint(unittest.TestCase):
    def test_endpoint_wraps_export_rows(self):
        import sys
        sys.path.insert(0, str(ROOT / "system" / "api"))
        sys.modules["ecosystem_intelligence"] = ei
        import server

        tmpdir = tempfile.mkdtemp()
        try:
            graph_path = Path(tmpdir) / "ecosystem_intelligence.json"
            graph_path.write_text(json.dumps(_graph_with_vendors()), encoding="utf-8")
            orig_path = server.ecosystem_intelligence.core.ECOSYSTEM_INTELLIGENCE_PATH
            server.ecosystem_intelligence.core.ECOSYSTEM_INTELLIGENCE_PATH = graph_path
            try:
                with patch.object(server, "_auth", lambda *a, **k: None):
                    result = server.get_vendors_list(x_api_key=None)
                self.assertEqual(result["vendor_count"], 2)
                self.assertEqual(result["contract"], "rb_vendor_list_v1")
            finally:
                server.ecosystem_intelligence.core.ECOSYSTEM_INTELLIGENCE_PATH = orig_path
        finally:
            shutil.rmtree(tmpdir, ignore_errors=True)


class TestVendorListXlsxDownload(unittest.TestCase):
    """RB-2026-09-01: Todd's same-day follow-up -- 'this should be in excel
    format' -- switched the /vendor-list/download route from CSV to a real
    xlsx workbook. Exercises the actual route function, with
    _execute_operation mocked so this doesn't need a live server."""

    def setUp(self):
        import sys
        sys.path.insert(0, str(ROOT / "system" / "api"))
        import rbb_chat
        self.rbb_chat = rbb_chat
        self._orig_execute = rbb_chat._execute_operation
        self._orig_auth = rbb_chat._auth_flexible
        rbb_chat._auth_flexible = lambda *a, **k: None

        def _fake_execute(operation_id, arguments):
            self.assertEqual(operation_id, "listVendors")
            return 200, {"vendors": [{
                "vendor_id": "vendor-toast", "name": "Toast", "aliases": "Toast POS",
                "primary_category": "pos", "status": "active", "brand_relationship_count": 1,
                "is_tracked_competitor": True, "created_at": "2026-08-01T00:00:00", "updated_at": "2026-08-01T00:00:00",
            }]}
        rbb_chat._execute_operation = _fake_execute

    def tearDown(self):
        self.rbb_chat._execute_operation = self._orig_execute
        self.rbb_chat._auth_flexible = self._orig_auth

    def test_returns_real_xlsx_workbook(self):
        import io
        from openpyxl import load_workbook

        resp = self.rbb_chat.get_vendor_list_download(passcode=None, x_chat_passcode=None)
        self.assertEqual(
            resp.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertTrue(resp.headers["content-disposition"].endswith('.xlsx"'))

        wb = load_workbook(io.BytesIO(resp.body))
        ws = wb.active
        self.assertEqual(ws.max_row, 2)  # header + 1 vendor row
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(header, [
            "vendor_id", "name", "aliases", "primary_category", "status",
            "brand_relationship_count", "is_tracked_competitor", "created_at", "updated_at",
        ])
        row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
        self.assertEqual(row[1], "Toast")
        self.assertEqual(row[6], True)  # is_tracked_competitor stays a real bool, not a CSV string

    def test_sanitizes_formula_injection_in_name_and_aliases(self):
        """RB-SECURITY-2026-09-05: name/aliases can carry externally-
        influenced text (a captured company name) into a real user-
        downloaded export. Same formula/CSV-injection guard (CWE-1236)
        already applied to Blue Sheets."""
        import io
        from openpyxl import load_workbook

        def _evil_execute(operation_id, arguments):
            return 200, {"vendors": [{
                "vendor_id": "vendor-evil", "name": "=1+1", "aliases": "-2+2",
                "primary_category": "pos", "status": "active", "brand_relationship_count": 0,
                "is_tracked_competitor": False, "created_at": "2026-08-01T00:00:00", "updated_at": "2026-08-01T00:00:00",
            }]}
        self.rbb_chat._execute_operation = _evil_execute

        resp = self.rbb_chat.get_vendor_list_download(passcode=None, x_chat_passcode=None)
        wb = load_workbook(io.BytesIO(resp.body))
        ws = wb.active
        row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
        self.assertEqual(row[1], "'=1+1")
        self.assertEqual(row[2], "'-2+2")

    def test_propagates_upstream_error(self):
        self.rbb_chat._execute_operation = lambda op, args: (500, {"error": "boom"})
        with self.assertRaises(Exception):
            self.rbb_chat.get_vendor_list_download(passcode=None, x_chat_passcode=None)


if __name__ == "__main__":
    unittest.main()
