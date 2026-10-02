"""
test_team_sales_tools.py

Regression coverage for system/scripts/team_sales_tools.py (Blue Sheet
Center + Green Sheet Center on the Team Portal). Same discipline as
test_team_portal.py's tech-stack whitelist tests: a synthetic fixture
account carrying distinctive, unmistakable "canary" strings in every
judgment/tactical field, asserting those canaries never reach the
team-safe output even though they are genuinely present on the underlying
account.json/brand_profile.json.
"""
from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "system" / "scripts"))
sys.path.insert(0, str(ROOT / "blue_sheets" / "_engine"))

import team_sales_tools as tst  # noqa: E402
import common as bs_common  # noqa: E402
import customers_prospects_common as cpc  # noqa: E402
import green_sheet  # noqa: E402
import artifact_vault_common as avc  # noqa: E402

_CANARY = "TODD-PRIVATE-CANARY"
_UNVERIFIED_NAME = "Fake Contact"


def _fixture_account() -> dict:
    return {
        "account_id": "acct-fixture-brand",
        "account_slug": "fixture-brand",
        "display_name": "Fixture Brand",
        "aliases": [],
        "owners": [{"name": "Todd Vahlsing"}],
        "updated_at": "2026-09-01T00:00:00Z",
        "template_version": "1.0",
        "engagement_tier": "active",
        "portfolio_status": "active_pursuit",
        "opportunities": [
            {
                "opportunity_id": "opp-fixture-1",
                "name": "Fixture Opportunity",
                "customer_stated_objective": {"value": "Public: reduce checkout friction"},
                "single_sales_objective": {
                    "value": f"{_CANARY}: undercut Zeta on renewal timing", "status": "confirmed",
                },
                "commercial_hypothesis": {
                    "value": f"{_CANARY}: price 8% below list to win", "status": "confirmed",
                },
            }
        ],
        "qualification": {
            "criteria": [
                {"answer": "Account potential", "current_read": f"{_CANARY}: weak, don't overinvest",
                 "next_step": f"{_CANARY}: quietly deprioritize -- loop in {_UNVERIFIED_NAME} first"},
            ],
        },
        "strategic_position": {
            "euphoria_panic": {
                "current_state": {"value": f"{_CANARY}: panic, losing ground", "status": "confirmed"},
                "reason": {"value": f"{_CANARY}: competitor about to win"},
                "timing": {"value": f"{_CANARY}: next 30 days"},
            },
            "competition": {"value": f"{_CANARY}: Zeta is entrenched, exploit their support gaps"},
            "position": {
                "place_in_funnel": f"{_CANARY}: early",
                "customer_priority": f"{_CANARY}: low",
                "critical_test": f"{_CANARY}: get a pilot signed",
                "position_vs_competition": {"value": f"{_CANARY}: behind"},
                "immediate_move": f"{_CANARY}: pressure-test with {_UNVERIFIED_NAME}",
            },
            "strengths": [
                {"value": "Public: strong uptime record", "possible_action": "Public: cite uptime SLA",
                 "best_action_plan": f"{_CANARY}: use uptime to peel off IT champion",
                 "owner": "Todd Vahlsing", "target": "2026-10-01"},
            ],
            "red_flags": [
                {"value": "Public: pricing perception risk", "possible_action": "Public: clarify pricing",
                 "best_action_plan": f"{_CANARY}: avoid pricing talk until champion is locked",
                 "owner": "Todd Vahlsing", "target": "2026-10-15"},
            ],
        },
        "buying_influences": [
            {
                "name": "Jane Person", "title": "VP Technology", "location": "HQ",
                "role_etuc": {"value": f"{_CANARY}: Economic Buyer", "evidence_ids": ["ev-fixture-real-0001"]},
                "influence": "Public: senior IT role",
                "mode": {"value": f"{_CANARY}: skeptical, needs private handling"},
                "personal_win": {"value": f"{_CANARY}: wants a promotion, exploit this", "status": "confirmed"},
                "business_result": "Public: cost reduction",
                "competitive_preference": {"value": f"{_CANARY}: leans toward Zeta", "status": "confirmed"},
                "rating": {"value": f"{_CANARY}: not yet a real champion"},
                "current_read": f"{_CANARY}: privately hostile to us",
                "access": f"{_CANARY}: only reachable through her assistant",
                "next_step": f"{_CANARY}: use Bob to get a side meeting",
                "owner": "Todd Vahlsing",
            },
            {
                # Real 2026-09-29 McDonald's finding, reproduced: a person
                # whose only evidence is Todd's own unverified LinkedIn
                # contact index, not account-specific research -- must be
                # dropped from the team-safe export entirely, and their
                # name must not leak through any other free-text field
                # either (qualification/strategic_position above, actions/
                # latest_review below).
                "name": _UNVERIFIED_NAME, "title": "Unrelated Co, Some Title", "location": None,
                "role_etuc": {"value": "Coach (unverified)", "evidence_ids": ["ev-fixture-baseline-0001"]},
                "influence": "Unknown", "mode": {"value": "Unknown"},
                "personal_win": {"value": "Unknown", "status": "unknown"},
                "business_result": "Unknown",
                "competitive_preference": {"value": "Unknown", "status": "unknown"},
                "rating": {"value": "Unknown"}, "current_read": "Unknown", "access": "Unknown",
                "next_step": "Unknown", "owner": "Todd Vahlsing",
            },
        ],
        "technology_stack": [
            {"layer": "POS", "vendor": "Zeta POS", "current_state": "Public: deployed 2024",
             "confidence": "high", "status": "Confirmed"},
        ],
        "commercial_models": [
            {"item": "License fee", "value": "1200", "unit": "per unit/mo", "confidence": "medium", "status": "Proposed"},
        ],
        "latest_review": {
            "blue_sheet_owner": "Todd Vahlsing", "last_reviewed": "2026-09-01",
            "core_contributors": "Todd Vahlsing", "review_status": "Active",
            "current_critical_test": "Public: land a pilot",
            "next_formal_review": f"After first pressure-test with {_UNVERIFIED_NAME}",
        },
        "bottom_line": f"{_CANARY}: this account only matters as leverage against Zeta elsewhere",
    }


def _fixture_evidence() -> list:
    return [
        {"evidence_id": "ev-fixture-real-0001", "description": "Jane Person / Fixture Brand JPR call intelligence",
         "excerpt": "Real call transcript with Jane Person.", "captured_at": "2026-09-01"},
        {"evidence_id": "ev-fixture-baseline-0001", "description": "LinkedIn baseline relationship map",
         "excerpt": "Section 10 -- baseline_index.json, 3,082 records.", "captured_at": "2026-08-01"},
    ]


def _fixture_brand_profile() -> dict:
    item = {
        "category": "growth_goal", "value": "Public: expanding to 50 new units",
        "strategic_implication": f"{_CANARY}: they'll need budget flexibility we can exploit",
        "confidence": "high", "as_of": "2026-09-01",
    }
    return {
        "identity_ownership_footprint": [item],
        "leadership": [],
        "technology_payment_landscape": {"highlights": []},
        "brand_digital_cx_strategy": [],
        "account_economics_scale": [],
        "relationship_history": [],
    }


class TestBlueSheetCenterFirewall(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmpdir.cleanup)
        tmp = Path(self.tmpdir.name)

        self.accounts_root = tmp / "customers_prospects"
        acct_dir = self.accounts_root / "accounts" / "fixture-brand"
        acct_dir.mkdir(parents=True)
        (acct_dir / "account.json").write_text(json.dumps(_fixture_account()), encoding="utf-8")
        (acct_dir / "brand_profile.json").write_text(json.dumps(_fixture_brand_profile()), encoding="utf-8")
        (acct_dir / "actions.json").write_text(json.dumps({
            "account_id": "acct-fixture-brand", "parent_rbb_loop_id": None,
            "actions": [
                {"type": "follow_up", "issue": "Public issue", "description": "Public: send pricing",
                 "owner": "Todd Vahlsing", "target": "2026-10-01", "status": "open", "blocker": ""},
                {"type": "follow_up", "issue": f"Pressure-test with {_UNVERIFIED_NAME}",
                 "description": f"Internal enablement with {_UNVERIFIED_NAME}",
                 "owner": "Todd Vahlsing", "target": "2026-10-01", "status": "open", "blocker": ""},
            ],
        }), encoding="utf-8")
        (acct_dir / "contradictions.json").write_text("{}", encoding="utf-8")
        (acct_dir / "source_index.json").write_text("{}", encoding="utf-8")
        (acct_dir / "evidence.jsonl").write_text(
            "\n".join(json.dumps(e) for e in _fixture_evidence()) + "\n", encoding="utf-8",
        )

        self._bs_root_patch = patch.object(bs_common, "CUSTOMERS_PROSPECTS_ROOT", self.accounts_root)
        self._bs_root_patch.start()
        self.addCleanup(self._bs_root_patch.stop)

    def test_coverage_report_flags_unmade_judgment_calls_as_needing_input(self):
        """The fixture never designates a Coach anywhere, so that check
        must fail -- distinct from the export-time masking test below:
        coverage reports whether Todd has MADE a judgment call at all
        (status != unknown), the export masks the VALUE of that call
        either way. A field can legitimately be both "populated" here
        (Todd has decided) and masked in the workbook (a teammate still
        can't see what he decided)."""
        cov = tst.get_blue_sheet_coverage("fixture-brand")
        self.assertIn("Coach", cov["needs_team_input"])
        self.assertIn("Brand and opportunity identity", cov["populated"])

    def test_coverage_report_flags_missing_judgment_calls_when_status_is_unknown(self):
        account = _fixture_account()
        account["buying_influences"][0]["personal_win"] = {"value": "Unknown", "status": "unknown"}
        account["buying_influences"][0]["competitive_preference"] = {"value": "Unknown", "status": "unknown"}
        account["strategic_position"]["euphoria_panic"]["current_state"] = {"value": "Unknown", "status": "unknown"}
        (self.accounts_root / "accounts" / "fixture-brand" / "account.json").write_text(
            json.dumps(account), encoding="utf-8",
        )
        cov = tst.get_blue_sheet_coverage("fixture-brand")
        for label in ("Personal wins", "Competitive preference", "Euphoria/Panic"):
            self.assertIn(label, cov["needs_team_input"])

    def test_unknown_slug_raises_not_found(self):
        with self.assertRaises(FileNotFoundError):
            tst.get_blue_sheet_coverage("does-not-exist")

    def test_rendered_workbook_never_contains_canary(self):
        import openpyxl
        buf, display_name = tst.render_team_blue_sheet_workbook("fixture-brand")
        self.assertEqual(display_name, "Fixture Brand")

        wb = openpyxl.load_workbook(buf)
        blob = "\n".join(
            str(cell.value) for ws in wb.worksheets for row in ws.iter_rows() for cell in row
            if cell.value is not None
        )
        self.assertNotIn(_CANARY, blob, "private judgment/tactical content leaked into the team-safe workbook")
        # Public/factual content must still be present -- this is a
        # whitelist, not a wipe.
        self.assertIn("Public: strong uptime record", blob)
        self.assertIn("Zeta POS", blob)
        self.assertIn("Jane Person", blob)
        self.assertIn("Team input required", blob)

    def test_rendered_workbook_never_attributes_rows_to_todd_by_name(self):
        """Real 2026-09-29 production-deployment feedback: a downloaded
        team-safe workbook must not carry Todd's name as the owner of
        every buying influence / strength / red flag / action, nor as
        Account Owner / Blue Sheet Owner / Core Contributors -- the
        fixture sets all of these to "Todd Vahlsing" specifically so a
        regression here fails loudly rather than by coincidence."""
        import openpyxl
        buf, _ = tst.render_team_blue_sheet_workbook("fixture-brand")
        wb = openpyxl.load_workbook(buf)
        blob = "\n".join(
            str(cell.value) for ws in wb.worksheets for row in ws.iter_rows() for cell in row
            if cell.value is not None
        )
        self.assertNotIn("Todd Vahlsing", blob)

    def test_render_does_not_touch_canonical_blue_sheet_files(self):
        canonical_dir = Path(__file__).resolve().parents[2] / "blue_sheets" / "accounts"
        before = set(canonical_dir.rglob("*.xlsx")) if canonical_dir.exists() else set()
        tst.render_team_blue_sheet_workbook("fixture-brand")
        after = set(canonical_dir.rglob("*.xlsx")) if canonical_dir.exists() else set()
        self.assertEqual(before, after)

    def test_unverified_contact_excluded_from_buying_influences_table(self):
        """Real 2026-09-29 McDonald's finding, reproduced: a buying
        influence whose only evidence traces to Todd's own unverified
        LinkedIn baseline contact index (not account-specific research)
        must not appear as a stakeholder in the team-safe workbook at
        all -- masking its fields to "Team input required" is not
        enough, since the wrong PERSON, not just their judgment fields,
        is the leak."""
        import openpyxl
        buf, _ = tst.render_team_blue_sheet_workbook("fixture-brand")
        wb = openpyxl.load_workbook(buf)
        blob = "\n".join(
            str(cell.value) for ws in wb.worksheets for row in ws.iter_rows() for cell in row
            if cell.value is not None
        )
        self.assertNotIn(_UNVERIFIED_NAME, blob)
        self.assertIn("Jane Person", blob)  # the verified contact must still appear

    def test_unverified_contact_name_scrubbed_from_free_text_elsewhere(self):
        """The unverified name also appears verbatim in qualification.
        criteria[].next_step, strategic_position.position.immediate_move,
        actions[].issue/description, and latest_review.next_formal_review
        -- none of those are the buying_influences table, so excluding
        the row alone doesn't touch them. The blanket name-scrub pass
        must catch all of them."""
        import openpyxl
        buf, _ = tst.render_team_blue_sheet_workbook("fixture-brand")
        wb = openpyxl.load_workbook(buf)
        blob = "\n".join(
            str(cell.value) for ws in wb.worksheets for row in ws.iter_rows() for cell in row
            if cell.value is not None
        )
        self.assertNotIn(_UNVERIFIED_NAME, blob)
        self.assertIn(tst._SCRUB_PLACEHOLDER, blob)

    def test_coverage_report_excludes_unverified_contact_from_stakeholder_count_and_checks(self):
        """The fixture's unverified contact claims role_etuc "Coach
        (unverified)" -- if the coverage report counted it, "Coach" would
        wrongly disappear from needs_team_input, and the stakeholder
        count would include a person who isn't real."""
        cov = tst.get_blue_sheet_coverage("fixture-brand")
        self.assertIn("Coach", cov["needs_team_input"])
        self.assertIn("1 known stakeholders", cov["populated"])


class TestGreenSheetCenter(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmpdir.cleanup)
        tmp = Path(self.tmpdir.name)

        acct_dir = tmp / "customers_prospects" / "accounts" / "fixture-brand"
        acct_dir.mkdir(parents=True)
        (acct_dir / "account.json").write_text(json.dumps(_fixture_account()), encoding="utf-8")
        evidence_with_participants = _fixture_evidence() + [
            {"evidence_id": "ev-fixture-real-0002", "description": "Real call with Jane Person",
             "participants": ["Jane Person"],
             "excerpt": f"{_CANARY}: Jane privately admitted budget is frozen until Q2.",
             "captured_at": "2026-09-05"},
        ]
        (acct_dir / "evidence.jsonl").write_text(
            "\n".join(json.dumps(e) for e in evidence_with_participants) + "\n", encoding="utf-8",
        )
        (acct_dir / "discovery_questions.json").write_text("{}", encoding="utf-8")

        self._cpc_root_patch = patch.object(cpc, "ROOT", tmp / "customers_prospects")
        self._vault_patch = patch.object(avc, "VAULT_ROOT", tmp / "artifact_vault")
        self._cpc_root_patch.start()
        self._vault_patch.start()
        self.addCleanup(self._cpc_root_patch.stop)
        self.addCleanup(self._vault_patch.stop)

    def test_missing_call_purpose_is_rejected(self):
        with self.assertRaises(tst.ValidationError):
            tst.preview_team_green_sheet("fixture-brand", call_purpose="  ", attendees=["Jane Person"])

    def test_missing_attendees_is_rejected(self):
        with self.assertRaises(tst.ValidationError):
            tst.preview_team_green_sheet("fixture-brand", call_purpose="Discovery call", attendees=[])

    def test_unknown_slug_raises_not_found(self):
        with self.assertRaises(tst.NotFoundError):
            tst.preview_team_green_sheet("does-not-exist", call_purpose="Discovery call", attendees=["Jane"])

    def test_preview_stamps_footer(self):
        result = tst.preview_team_green_sheet(
            "fixture-brand", call_purpose="Discovery call", attendees=["Jane Person"],
            requested_by="zak-bertram",
        )
        self.assertIn("Generated by: RBB Team Portal (requested by zak-bertram)", result["markdown"])
        self.assertIn("Requires human validation", result["markdown"])

    def test_preview_never_touches_the_canonical_artifact_vault(self):
        """render_green_sheet() is called directly -- never generate_
        green_sheet(), which would persist a new version into the same
        artifact_vault store Todd's own real Green Sheets live in."""
        vault_dir = Path(avc.VAULT_ROOT) / "green_sheets" / "fixture-brand"
        self.assertFalse(vault_dir.exists())
        tst.preview_team_green_sheet(
            "fixture-brand", call_purpose="Discovery call", attendees=["Jane Person"],
        )
        self.assertFalse(vault_dir.exists(), "Team Portal preview must not write a version to the artifact vault")

    def test_preview_never_contains_matched_attendees_judgment_fields(self):
        """The On This Call table matches Jane Person (an attendee) and
        would otherwise render her real role_etuc/current_read/next_step
        and the account's single_sales_objective straight from the
        fixture -- the exact leak found live 2026-09-29."""
        result = tst.preview_team_green_sheet(
            "fixture-brand", call_purpose="Discovery call", attendees=["Jane Person"],
        )
        self.assertNotIn(_CANARY, result["markdown"])
        self.assertIn("On This Call", result["markdown"])
        self.assertIn("Jane Person", result["markdown"])
        self.assertIn("Team input required", result["markdown"])

    def test_preview_never_contains_unverified_attendee_as_a_real_contact(self):
        """The unverified contact is a real attendee name a caller could
        legitimately type -- the Green Sheet must show "no on-file
        record" for them, not fabricate a row from unverified data."""
        result = tst.preview_team_green_sheet(
            "fixture-brand", call_purpose="Discovery call", attendees=["Jane Person", _UNVERIFIED_NAME],
        )
        self.assertIn(f"No on-file buying-influence record for: {_UNVERIFIED_NAME}", result["markdown"])

    def test_preview_never_renders_raw_evidence_excerpts(self):
        """Real 2026-09-29 finding: evidence.jsonl excerpts are raw,
        unvetted free text -- render_relevant_evidence_section renders
        them verbatim for any attendee whose name matches a record's
        participants[]. Since that text can't be reliably field-masked,
        the team-safe dossier drops the evidence list entirely."""
        result = tst.preview_team_green_sheet(
            "fixture-brand", call_purpose="Discovery call", attendees=["Jane Person"],
        )
        self.assertNotIn(_CANARY, result["markdown"])
        self.assertNotIn("budget is frozen", result["markdown"])
        self.assertIn("No prior evidence on file", result["markdown"])


if __name__ == "__main__":
    unittest.main()
