"""
test_capture_ingest_document_extraction.py — RB-2026-09-20.

Extends the image-OCR capture work to PDF/Word/Excel: a document dropped
into a watched folder (e.g. settings.json's icloud_screenshots source,
whose patterns now include .pdf/.docx/.xlsx/.xlsm) is mechanically
extracted at batch-sweep time instead of relayed live through rbb-chat's
/upload endpoint. Covers _extract_text()'s document dispatch, each
extractor's real-file happy path and missing-package failure mode, and the
capture_type default fix (document_extraction -> "pasted_content", never
_detect_capture_type()'s meeting-oriented fallback) alongside image_ocr.
"""
from __future__ import annotations

import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "system" / "scripts"))

import capture_ingest as ci  # noqa: E402


def _real_xlsx_bytes() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Value"])
    ws.append(["Widget", 42])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _real_docx_bytes() -> bytes:
    import docx
    doc = docx.Document()
    doc.add_paragraph("Hello from a real test document.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestExtractTextDocumentDispatch(unittest.TestCase):
    def test_document_suffixes_route_to_document_extraction(self):
        for suffix in (".pdf", ".docx", ".xlsx", ".xlsm"):
            with tempfile.TemporaryDirectory() as d:
                path = Path(d) / f"doc{suffix}"
                path.write_bytes(b"irrelevant -- extraction itself is mocked")
                with patch.object(ci, "_extract_document_text", MagicMock(return_value="extracted text")) as mock_extract:
                    text, method = ci._extract_text(path, "pre_transcribed")

            mock_extract.assert_called_once_with(path)
            self.assertEqual(text, "extracted text")
            self.assertEqual(method, "document_extraction")


class TestExtractTextFromXlsxReal(unittest.TestCase):
    def test_real_workbook_flattened_to_text(self):
        content = _real_xlsx_bytes()
        text = ci._extract_text_from_xlsx(content)
        self.assertIn("Sheet1", text)
        self.assertIn("Widget", text)
        self.assertIn("42", text)

    def test_missing_openpyxl_raises_runtime_error(self):
        with patch.dict(sys.modules, {"openpyxl": None}):
            with self.assertRaises(RuntimeError):
                ci._extract_text_from_xlsx(b"irrelevant")


class TestExtractTextFromDocxReal(unittest.TestCase):
    def test_real_document_flattened_to_text(self):
        content = _real_docx_bytes()
        text = ci._extract_text_from_docx(content)
        self.assertIn("Hello from a real test document.", text)

    def test_missing_python_docx_raises_runtime_error(self):
        with patch.dict(sys.modules, {"docx": None}):
            with self.assertRaises(RuntimeError):
                ci._extract_text_from_docx(b"irrelevant")


class TestExtractTextFromPdf(unittest.TestCase):
    """No PDF-writing library is available in this environment, so pypdf's
    own PdfReader is mocked -- this tests our wrapper's call/join/failure
    handling, not pypdf's real parsing correctness."""

    def test_pages_joined_in_order(self):
        page1 = MagicMock()
        page1.extract_text.return_value = "Page one text"
        page2 = MagicMock()
        page2.extract_text.return_value = "Page two text"
        mock_reader = MagicMock()
        mock_reader.pages = [page1, page2]

        with patch("pypdf.PdfReader", MagicMock(return_value=mock_reader)):
            text = ci._extract_text_from_pdf(b"irrelevant")

        self.assertEqual(text, "Page one text\nPage two text")

    def test_scanned_pdf_with_no_text_layer_returns_empty_not_raise(self):
        page = MagicMock()
        page.extract_text.return_value = None
        mock_reader = MagicMock()
        mock_reader.pages = [page]

        with patch("pypdf.PdfReader", MagicMock(return_value=mock_reader)):
            text = ci._extract_text_from_pdf(b"irrelevant")

        self.assertEqual(text, "")

    def test_missing_pypdf_raises_runtime_error(self):
        with patch.dict(sys.modules, {"pypdf": None}):
            with self.assertRaises(RuntimeError):
                ci._extract_text_from_pdf(b"irrelevant")


class TestSweepDocumentCaptureTypeDefault(unittest.TestCase):
    def setUp(self):
        self._tmpdir = tempfile.TemporaryDirectory()
        self.folder = Path(self._tmpdir.name).resolve()
        self.pending_dir = self.folder / "captures" / "pending"

    def tearDown(self):
        self._tmpdir.cleanup()

    def test_pdf_capture_defaults_to_pasted_content_not_meeting(self):
        (self.folder / "report.pdf").write_bytes(b"x" * 20)
        source = {
            "id": "test_documents",
            "label": "Test Documents",
            "folder": str(self.folder),
            "patterns": ["**/*.pdf"],
            "transcription": "pre_transcribed",
        }

        with patch.object(ci, "_load_sources", MagicMock(return_value=[source])), \
             patch.object(ci, "_extract_document_text", MagicMock(return_value="document body text")), \
             patch.object(ci, "SYSTEM_DIR", self.folder / "system_fake"), \
             patch.object(ci, "PENDING_DIR", self.pending_dir), \
             patch.object(ci, "CAPTURES_DIR", self.folder / "captures"), \
             patch.object(ci, "REGISTRY_PATH", self.folder / "captures" / ".registry.json"), \
             patch.object(ci.audit_log, "append_event", MagicMock()):
            result = ci.sweep()

        self.assertEqual(result["files_queued"], 1)
        pending_files = list(self.pending_dir.glob("*.json"))
        self.assertEqual(len(pending_files), 1)
        payload = json.loads(pending_files[0].read_text())
        self.assertEqual(payload["capture_type"], "pasted_content")
        self.assertEqual(payload["transcription_method"], "document_extraction")
        self.assertEqual(payload["transcript"], "document body text")


if __name__ == "__main__":
    unittest.main()
