"""
test_uploaded_document_retrieval.py — RB-2026-08-31.

Todd's defect report (Pollo Campero RFP response document, due 2026-09-04):
uploadAndIngestFile extracted a DOCX/PDF's full text mechanically, used it
in-memory for triage/mutations/reference-linking, then discarded it -- only
a 400-char excerpt ever survived. There was no way to retrieve "what does
this document actually say" after that one call, so the GPT could not do a
section-by-section review of an uploaded response document against known
account requirements (TDR/SOW). This closes that gap: uploadAndIngestFile
now persists the full extraction (uploaded_document_store) and returns a
document_id; getUploadedDocument retrieves it, in full or by section/page.
"""
from __future__ import annotations

import base64
import io
import json
from pathlib import Path
from types import SimpleNamespace

import docx
import pytest

from system.api import server


def _docx_base64(sections: list[tuple[str, str]]) -> str:
    doc = docx.Document()
    for heading, body in sections:
        doc.add_heading(heading, level=1)
        doc.add_paragraph(body)
    buf = io.BytesIO()
    doc.save(buf)
    return base64.b64encode(buf.getvalue()).decode()


@pytest.fixture()
def isolated_upload(tmp_path: Path, monkeypatch):
    inbox = tmp_path / "system" / "inbox"
    cache = tmp_path / "system" / ".cache"
    doc_store = tmp_path / "system" / "uploaded_documents"
    index_path = tmp_path / "system" / "intelligence_index.json"
    update_log_path = tmp_path / "system" / "intelligence_index_updates.jsonl"
    inbox.mkdir(parents=True)
    cache.mkdir(parents=True)

    # RB-SECURITY-2026-09-03 made server._auth() fail *closed* on an unset
    # API_KEY -- bypass _auth() itself rather than setting API_KEY=None,
    # same pattern test_vendor_list.py uses for rbb_chat._auth_flexible.
    monkeypatch.setattr(server, "_auth", lambda *a, **k: None)
    monkeypatch.setattr(server.core, "PROJECT_DIR", tmp_path)
    monkeypatch.setattr(server.core, "INBOX_DIR", inbox)
    monkeypatch.setattr(server.core, "CACHE_DIR", cache)
    monkeypatch.setattr(server, "_load_artifact_registry", lambda: {})
    monkeypatch.setattr(server.uploaded_document_store, "STORE_DIR", doc_store)
    monkeypatch.setattr(server.intelligence_index, "INDEX_PATH", index_path)
    monkeypatch.setattr(server.intelligence_index, "UPDATE_LOG_PATH", update_log_path)
    monkeypatch.setattr(
        server.intelligence_triage,
        "triage_input",
        lambda *args, **kwargs: {
            "triage_id": "TRG-TEST",
            "type_count": 1,
            "noise_only": False,
            "identified_types": [{"intelligence_type": "macro_signal"}],
        },
    )
    monkeypatch.setattr(
        server.intelligence_mutation_engine,
        "run",
        lambda *args, **kwargs: {
            "trust_stats": {"total_mutations": 0, "confidence": "medium"},
            "apply_result": {"applied": 0},
        },
    )
    return SimpleNamespace(inbox=inbox, cache=cache, doc_store=doc_store, index_path=index_path)


def test_docx_upload_returns_retrievable_document_id(isolated_upload):
    content_b64 = _docx_base64([
        ("Pricing", "Fixed per-transaction fee, no monthly minimum."),
        ("Implementation Timeline", "Go-live within 90 days of signature."),
    ])
    receipt = server.post_ingest_upload(
        server.FileUploadIn(
            filename="RFP Response - Fictional Account.docx",
            content_base64=content_b64,
            dry_run=True,
        ),
        x_api_key=None,
    )

    assert receipt["pipeline_run"] == "intelligence"
    document_id = receipt["ingest_result"]["document_id"]
    assert document_id == receipt["ingestion_id"]
    assert receipt["ingest_result"]["extraction_status"] == "ok"
    assert "Pricing" in receipt["ingest_result"]["extracted_text_preview"]

    stored = server.uploaded_document_store.load_document(document_id)
    assert stored is not None
    assert "Fixed per-transaction fee" in stored["extracted_text"]
    assert [s["heading"] for s in stored["sections"]] == ["Pricing", "Implementation Timeline"]

    full = server.get_uploaded_document(document_id=document_id, section_index=None, x_api_key=None)
    assert full["filename"] == "RFP Response - Fictional Account.docx"
    assert "Go-live within 90 days" in full["extracted_text"]
    assert full["truncated"] is False
    assert full["sections"] == [
        {"index": 0, "heading": "Pricing"},
        {"index": 1, "heading": "Implementation Timeline"},
    ]

    section = server.get_uploaded_document(document_id=document_id, section_index=1, x_api_key=None)
    assert section["section"]["heading"] == "Implementation Timeline"
    assert "Go-live within 90 days" in section["section"]["text"]


def test_pdf_upload_extracts_pages(isolated_upload):
    from system.tests.test_docx_pdf_ingest_extraction import _make_minimal_pdf

    content = _make_minimal_pdf(["Section A text", "Section B text"])
    receipt = server.post_ingest_upload(
        server.FileUploadIn(
            filename="RFP Response.pdf",
            content_base64=base64.b64encode(content).decode(),
            dry_run=True,
        ),
        x_api_key=None,
    )

    document_id = receipt["ingest_result"]["document_id"]
    full = server.get_uploaded_document(document_id=document_id, section_index=None, x_api_key=None)
    assert full["content_type"] == "application/pdf"
    assert [s["heading"] for s in full["sections"]] == ["Page 1", "Page 2"]
    assert "Section A text" in full["extracted_text"]
    assert "Section B text" in full["extracted_text"]


def test_scanned_pdf_reports_specific_extraction_failure_reason(isolated_upload):
    """RB-2026-08-31: acceptance criteria -- a PDF with no extractable text
    layer must fail with a specific, actionable reason, not the generic
    'extracted_text is required' message alone, and the GPT must not pretend
    to review text that doesn't exist."""
    from system.tests.test_docx_pdf_ingest_extraction import _make_minimal_pdf

    blank_page_pdf = _make_minimal_pdf([""])
    receipt = server.post_ingest_upload(
        server.FileUploadIn(
            filename="Scanned RFP.pdf",
            content_base64=base64.b64encode(blank_page_pdf).decode(),
            dry_run=True,
        ),
        x_api_key=None,
    )
    assert receipt["ingest_result"]["ok"] is False
    assert receipt["ingest_result"]["extraction_status"] == "failed"
    assert "scanned" in receipt["ingest_result"]["error"].lower()


def test_retrieval_404_for_unknown_document_id(isolated_upload):
    from fastapi import HTTPException

    with pytest.raises(HTTPException) as exc_info:
        server.get_uploaded_document(document_id="not-a-real-id", section_index=None, x_api_key=None)
    assert exc_info.value.status_code == 404


def test_uploaded_document_registered_in_intelligence_index(isolated_upload):
    content_b64 = _docx_base64([("Overview", "Fictional Account proposal overview.")])
    receipt = server.post_ingest_upload(
        server.FileUploadIn(
            filename="RFP Response - Fictional Account.docx",
            content_base64=content_b64,
            dry_run=True,
        ),
        x_api_key=None,
    )
    document_id = receipt["ingest_result"]["document_id"]

    matches = server.intelligence_index.find("RFP Response - Fictional Account")
    assert any(m["path"].endswith(f"{document_id}.json") for m in matches)
    entry = next(m for m in matches if m["path"].endswith(f"{document_id}.json"))
    assert entry["resource_type"] == "uploaded_source_document"


def test_large_document_text_truncated_but_sections_intact(isolated_upload):
    big_paragraph = "Filler content. " * 5000  # well past the 40k inline cap
    content_b64 = _docx_base64([("Big Section", big_paragraph)])
    receipt = server.post_ingest_upload(
        server.FileUploadIn(
            filename="Large RFP Response.docx",
            content_base64=content_b64,
            dry_run=True,
        ),
        x_api_key=None,
    )
    document_id = receipt["ingest_result"]["document_id"]

    full = server.get_uploaded_document(document_id=document_id, section_index=None, x_api_key=None)
    assert full["truncated"] is True
    assert full["note"] is not None

    section = server.get_uploaded_document(document_id=document_id, section_index=0, x_api_key=None)
    assert len(section["section"]["text"]) == len(big_paragraph.strip())
