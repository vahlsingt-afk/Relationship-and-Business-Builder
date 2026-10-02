"""
test_docx_pdf_ingest_extraction.py — RB-2026-08-28.

Same gap as test_xlsx_ingest_extraction.py's incident, found while auditing
that fix: .docx and .pdf never appeared anywhere in uploadAndIngestFile's
suffix-dispatch chain (server.py post_ingest_upload) either, so any such
upload fell through to the generic "intelligence" pipeline, which -- before
this fix -- could not read either format's bytes and always failed with
"extracted_text is required," identical to the .xlsx incident. Fix:
_extract_text_from_docx() (python-docx, paragraphs + table cells) and
_extract_text_from_pdf() (pypdf, per-page text) extract mechanically, never
LLM-guessed, before falling back to the error.
"""
from __future__ import annotations

import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "system" / "scripts"))
sys.path.insert(0, str(ROOT / "system" / "api"))

try:
    import docx
    _HAS_DOCX = True
except ImportError:
    _HAS_DOCX = False

try:
    import pypdf
    _HAS_PYPDF = True
except ImportError:
    _HAS_PYPDF = False

import server  # noqa: E402


@unittest.skipUnless(_HAS_DOCX, "python-docx not installed")
class TestExtractTextFromDocx(unittest.TestCase):
    def _make_docx_bytes(self, paragraphs: list[str], table_rows: list[tuple] | None = None) -> bytes:
        import io
        doc = docx.Document()
        for p in paragraphs:
            doc.add_paragraph(p)
        if table_rows:
            table = doc.add_table(rows=0, cols=len(table_rows[0]))
            for row in table_rows:
                cells = table.add_row().cells
                for i, val in enumerate(row):
                    cells[i].text = val
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def test_extracts_paragraph_text(self):
        content = self._make_docx_bytes(["Account Plan: Worldpay", "Renewal date: 2026-09-01"])
        text = server._extract_text_from_docx(content)
        self.assertIn("Account Plan: Worldpay", text)
        self.assertIn("Renewal date: 2026-09-01", text)

    def test_extracts_table_cells(self):
        content = self._make_docx_bytes(
            ["Stakeholders"],
            table_rows=[("Name", "Role"), ("Jane Doe", "CFO")],
        )
        text = server._extract_text_from_docx(content)
        self.assertIn("Name | Role", text)
        self.assertIn("Jane Doe | CFO", text)

    def test_skips_empty_paragraphs(self):
        content = self._make_docx_bytes(["First", "", "   ", "Second"])
        text = server._extract_text_from_docx(content)
        lines = text.splitlines()
        self.assertEqual(lines, ["First", "Second"])

    def test_invalid_bytes_return_empty_string_not_raise(self):
        self.assertEqual(server._extract_text_from_docx(b"not a real docx file"), "")

    def test_empty_document_returns_empty_string(self):
        content = self._make_docx_bytes([])
        self.assertEqual(server._extract_text_from_docx(content), "")

    def test_real_incident_file_if_present(self):
        real_file = ROOT / "system" / "inbox" / "user_artifacts"
        matches = list(real_file.glob("*.docx")) if real_file.exists() else []
        if not matches:
            self.skipTest("no real .docx incident file present in this environment")
        text = server._extract_text_from_docx(matches[0].read_bytes())
        self.assertGreater(len(text), 0, "real .docx artifact should extract some text")


@unittest.skipUnless(_HAS_DOCX, "python-docx not installed")
class TestExtractSectionsFromDocx(unittest.TestCase):
    """RB-2026-08-31: Todd's defect report -- an uploaded RFP response
    document's full text was extracted but discarded, with no section
    structure retrievable for a section-by-section review. Real Word
    heading styles ("Heading 1".."Heading 9", "Title") are a mechanical,
    non-guessed signal python-docx exposes directly."""

    def _make_docx_with_headings(self, blocks: list[tuple[str | None, str]],
                                  table_after: int | None = None,
                                  table_rows: list[tuple] | None = None) -> bytes:
        import io
        doc = docx.Document()
        for i, (heading, body) in enumerate(blocks):
            if heading is not None:
                doc.add_heading(heading, level=1)
            doc.add_paragraph(body)
            if table_after == i and table_rows:
                table = doc.add_table(rows=0, cols=len(table_rows[0]))
                for row in table_rows:
                    cells = table.add_row().cells
                    for j, val in enumerate(row):
                        cells[j].text = val
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def test_splits_on_heading_styles(self):
        content = self._make_docx_with_headings([
            ("Pricing", "Fixed per-transaction fee."),
            ("Timeline", "90 days to go-live."),
        ])
        sections = server._extract_sections_from_docx(content)
        self.assertEqual([s["heading"] for s in sections], ["Pricing", "Timeline"])
        self.assertIn("Fixed per-transaction fee.", sections[0]["text"])
        self.assertIn("90 days to go-live.", sections[1]["text"])

    def test_body_before_first_heading_has_none_heading(self):
        content = self._make_docx_with_headings([
            (None, "Cover page text with no heading yet."),
            ("Section One", "Real content."),
        ])
        sections = server._extract_sections_from_docx(content)
        self.assertEqual(sections[0]["heading"], None)
        self.assertIn("Cover page text", sections[0]["text"])
        self.assertEqual(sections[1]["heading"], "Section One")

    def test_table_attaches_to_preceding_section(self):
        content = self._make_docx_with_headings(
            [("Requirements", "See the table below.")],
            table_after=0,
            table_rows=[("Requirement", "Response"), ("Uptime SLA", "99.9%")],
        )
        sections = server._extract_sections_from_docx(content)
        self.assertEqual(len(sections), 1)
        self.assertIn("Uptime SLA | 99.9%", sections[0]["text"])

    def test_invalid_bytes_return_empty_list_not_raise(self):
        self.assertEqual(server._extract_sections_from_docx(b"not a real docx file"), [])

    def test_no_headings_returns_single_untitled_section(self):
        content = self._make_docx_with_headings([(None, "Just plain body text, no headings at all.")])
        sections = server._extract_sections_from_docx(content)
        self.assertEqual(len(sections), 1)
        self.assertEqual(sections[0]["heading"], None)


def _make_minimal_pdf(page_texts: list[str]) -> bytes:
    """Hand-build a minimal single-content-stream-per-page PDF (standard
    Helvetica font, no embedded font file needed) so the test has no
    dependency on a PDF-authoring library like reportlab -- only pypdf
    (already vendored for reading) is needed, and only to read this back."""
    import io as _io

    objects: list[bytes] = []
    n_pages = len(page_texts)
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{3 + i} 0 R" for i in range(n_pages))
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    font_obj_num = 3 + n_pages
    content_obj_start = font_obj_num + 1
    for i in range(n_pages):
        content_obj_num = content_obj_start + i
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_obj_num} 0 R >> >> "
            f"/Contents {content_obj_num} 0 R >>".encode()
        )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for text in page_texts:
        stream = f"BT /F1 24 Tf 72 700 Td ({text}) Tj ET".encode("latin-1")
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")

    out = _io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode())
        out.write(obj)
        out.write(b"\nendobj\n")
    xref_start = out.tell()
    total = len(objects) + 1
    out.write(f"xref\n0 {total}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets[1:]:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {total} /Root 1 0 R >>\n".encode())
    out.write(b"startxref\n")
    out.write(f"{xref_start}\n".encode())
    out.write(b"%%EOF")
    return out.getvalue()


@unittest.skipUnless(_HAS_PYPDF, "pypdf not installed")
class TestExtractTextFromPdf(unittest.TestCase):
    def _make_pdf_bytes(self, page_texts: list[str]) -> bytes:
        return _make_minimal_pdf(page_texts)

    def test_extracts_page_text(self):
        content = self._make_pdf_bytes(["Account Plan: Worldpay"])
        text = server._extract_text_from_pdf(content)
        self.assertIn("Worldpay", text)

    def test_multiple_pages_concatenated(self):
        content = self._make_pdf_bytes(["Page one content", "Page two content"])
        text = server._extract_text_from_pdf(content)
        self.assertIn("Page one content", text)
        self.assertIn("Page two content", text)

    def test_invalid_bytes_return_empty_string_not_raise(self):
        self.assertEqual(server._extract_text_from_pdf(b"not a real pdf file"), "")

    def test_empty_pdf_bytes_return_empty_string(self):
        self.assertEqual(server._extract_text_from_pdf(b""), "")

    def test_real_incident_file_if_present(self):
        real_dir = ROOT / "system" / "inbox" / "user_artifacts"
        matches = list(real_dir.glob("*.pdf")) if real_dir.exists() else []
        if not matches:
            self.skipTest("no real .pdf incident file present in this environment")
        text = server._extract_text_from_pdf(matches[0].read_bytes())
        self.assertGreater(len(text), 0, "real .pdf artifact should extract some text")


@unittest.skipUnless(_HAS_PYPDF, "pypdf not installed")
class TestExtractPagesFromPdf(unittest.TestCase):
    """RB-2026-08-31: per-page text is the one structural signal a PDF
    reliably carries without OCR/layout inference -- used by
    getUploadedDocument to let a large uploaded document be retrieved by
    page instead of only as one flat string."""

    def test_one_section_per_page(self):
        content = _make_minimal_pdf(["Page one content", "Page two content"])
        pages = server._extract_pages_from_pdf(content)
        self.assertEqual([p["heading"] for p in pages], ["Page 1", "Page 2"])
        self.assertIn("Page one content", pages[0]["text"])
        self.assertIn("Page two content", pages[1]["text"])

    def test_invalid_bytes_return_empty_list_not_raise(self):
        self.assertEqual(server._extract_pages_from_pdf(b"not a real pdf file"), [])

    def test_empty_bytes_return_empty_list(self):
        self.assertEqual(server._extract_pages_from_pdf(b""), [])


@unittest.skipUnless(_HAS_PYPDF, "pypdf not installed")
class TestDiagnosePdfExtractionFailure(unittest.TestCase):
    """RB-2026-08-31: Todd's defect report acceptance criteria explicitly
    asked for a specific reason (scanned/no-OCR, encrypted, parsing error)
    instead of one generic 'extracted_text is required' message regardless
    of cause -- so the GPT never pretends to review text that doesn't
    exist, and can tell Todd exactly why."""

    def test_unparseable_bytes_report_parsing_error(self):
        reason = server._diagnose_pdf_extraction_failure(b"not a real pdf file")
        self.assertIn("parsing error", reason)

    def test_blank_page_reports_scanned_or_no_text_layer(self):
        content = _make_minimal_pdf([""])
        self.assertEqual(server._extract_text_from_pdf(content), "")
        reason = server._diagnose_pdf_extraction_failure(content)
        self.assertIn("scanned", reason.lower())

    def test_encrypted_pdf_reports_encrypted(self):
        import io
        import pypdf

        reader = pypdf.PdfReader(io.BytesIO(_make_minimal_pdf(["secret content"])))
        writer = pypdf.PdfWriter()
        writer.append(reader)
        writer.encrypt(user_password="hunter2")
        buf = io.BytesIO()
        writer.write(buf)
        encrypted_bytes = buf.getvalue()

        self.assertEqual(server._extract_text_from_pdf(encrypted_bytes), "")
        reason = server._diagnose_pdf_extraction_failure(encrypted_bytes)
        self.assertIn("encrypted", reason.lower())


if __name__ == "__main__":
    unittest.main()
