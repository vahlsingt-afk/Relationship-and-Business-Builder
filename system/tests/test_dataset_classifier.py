#!/usr/bin/env python3
"""
Tests for dataset_classifier.py — generalized structured-dataset recognition
(RB-DEFECT-064 Phase 1).

Test IDs and coverage:

RB-CLASSIFIER-001: HubSpot export (filename + distinctive headers) classifies with high confidence
RB-CLASSIFIER-002: LinkedIn Connections.csv classifies as linkedin_connections_export
RB-CLASSIFIER-003: bare name-only CSV does not clear the auto-ingest confidence bar
RB-CLASSIFIER-004: non-tabular file classifies as unknown_structured_dataset, confidence 0.0
RB-CLASSIFIER-005: should_auto_ingest respects a configured confidence_threshold override
RB-CLASSIFIER-006: a signature with no ingest_script never triggers auto-ingest regardless of confidence
RB-CLASSIFIER-007: McDonald's NSN Lookup workbook classifies via sheet-name fingerprint (RB-DEFECT-065)
RB-CLASSIFIER-008: an xlsx workbook missing the NSN sheet fingerprint does not misclassify as it
RB-CLASSIFIER-009: classify_bytes matches classify(path) for the same content + filename
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook

_SCRIPTS = Path(__file__).resolve().parent.parent / "scripts"
sys.path.insert(0, str(_SCRIPTS))

import dataset_classifier as dc  # noqa: E402


def _write_csv(tmp_path: Path, name: str, headers: list[str]) -> Path:
    path = tmp_path / name
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerow(["Jane" for _ in headers])
    return path


def _write_nsn_workbook(tmp_path: Path, name: str = "NSN Lookup 2026-07 JUL.xlsx") -> Path:
    path = tmp_path / name
    wb = Workbook()
    ws = wb.active
    ws.title = "StoreTech"
    ws.append(["NSN", "Field Office", "Market", "Entity Name"])
    for sheet in ["Markets", "FO OTM-STIM", "RFM", "COOP2", "Entity"]:
        wb.create_sheet(sheet)
    wb.save(path)
    return path


def test_hubspot_export_high_confidence(tmp_path):
    """RB-CLASSIFIER-001."""
    path = _write_csv(
        tmp_path,
        "hubspot-crm-exports-all-contacts-2026-07-07.csv",
        ["First Name", "Last Name", "Email", "Company Name", "Job Title",
         "Contact owner", "Lead Status", "Marketing contact status"],
    )
    result = dc.classify(path)
    assert result.dataset_type == "hubspot_crm_export"
    assert result.confidence >= dc.DEFAULT_CONFIDENCE_THRESHOLD
    assert result.ingest_script == "system/scripts/hubspot_ingest.py"


def test_linkedin_export_classified(tmp_path):
    """RB-CLASSIFIER-002."""
    path = _write_csv(
        tmp_path, "Connections.csv",
        ["First Name", "Last Name", "URL", "Email Address", "Connected On"],
    )
    result = dc.classify(path)
    assert result.dataset_type == "linkedin_connections_export"
    assert result.already_handled_by == "P-002_linkedin_ingest.md"


def test_bare_name_csv_does_not_clear_bar(tmp_path):
    """RB-CLASSIFIER-003."""
    path = _write_csv(tmp_path, "some_export.csv", ["First Name", "Last Name"])
    result = dc.classify(path)
    assert not dc.should_auto_ingest(result, settings={"structured_ingest": {}})


def test_non_tabular_file_is_unknown(tmp_path):
    """RB-CLASSIFIER-004."""
    path = tmp_path / "notes.txt"
    path.write_text("just some prose, not a dataset")
    result = dc.classify(path)
    assert result.dataset_type == dc.UNKNOWN_TYPE
    assert result.confidence == 0.0


def test_should_auto_ingest_respects_configured_threshold():
    """RB-CLASSIFIER-005."""
    result = dc.Classification(
        dataset_type="hubspot_crm_export", confidence=0.5, purpose="historical_relationship_database",
        recommended_action="run_hubspot_ingest", ingest_script="system/scripts/hubspot_ingest.py",
    )
    # Below a strict configured threshold...
    assert not dc.should_auto_ingest(result, settings={"structured_ingest": {"confidence_threshold": 0.85}})
    # ...but a deliberately lowered threshold clears it, proving the gate
    # reads live config rather than a hardcoded constant.
    assert dc.should_auto_ingest(result, settings={"structured_ingest": {"confidence_threshold": 0.4}})


def test_no_ingest_script_never_auto_ingests(tmp_path):
    """RB-CLASSIFIER-006."""
    path = _write_csv(
        tmp_path, "salesforce_export.csv",
        ["Account ID", "Account Name", "Lead Source", "Contact Owner"],
    )
    result = dc.classify(path)
    assert result.dataset_type == "salesforce_crm_export"
    assert not dc.should_auto_ingest(result, settings={"structured_ingest": {"confidence_threshold": 0.0}})


def test_mcdonalds_nsn_workbook_classified_by_sheet_fingerprint(tmp_path):
    """RB-CLASSIFIER-007."""
    path = _write_nsn_workbook(tmp_path)
    result = dc.classify(path)
    assert result.dataset_type == "mcdonalds_nsn_lookup_workbook"
    assert result.ingest_script == "system/scripts/micro_graph_mcdonalds.py"
    assert dc.should_auto_ingest(result)


def test_xlsx_missing_nsn_fingerprint_is_unknown(tmp_path):
    """RB-CLASSIFIER-008."""
    path = tmp_path / "random_other.xlsx"
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.save(path)
    result = dc.classify(path)
    assert result.dataset_type != "mcdonalds_nsn_lookup_workbook"
    assert not dc.should_auto_ingest(result)


def test_classify_bytes_matches_classify_path(tmp_path):
    """RB-CLASSIFIER-009."""
    path = _write_nsn_workbook(tmp_path)
    from_path = dc.classify(path)
    from_bytes = dc.classify_bytes(path.read_bytes(), path.name)
    assert from_bytes.to_dict() == from_path.to_dict()
