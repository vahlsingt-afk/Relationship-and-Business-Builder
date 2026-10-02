"""
test_xlsx_ingest_extraction.py — RB-2026-08-28.

Real bug: Todd uploaded Worldpay_Master_Account_Plan_2026-08-14.xlsx through
rbb-chat and asked the CoS to assess/update it. uploadAndIngestFile's
"intelligence" pipeline (server.py) only auto-decoded a whitelist of
plain-text suffixes (.txt/.md/.csv/.json/...) -- for an unstructured/
narrative .xlsx that dataset_classifier doesn't recognize as a known
structured type (RB-DEFECT-065's own comment already flagged this: "the
generic intelligence/text pipeline, which can't even read xlsx bytes"), it
always failed with "extracted_text is required," no matter what the file
actually contained. Fix: _extract_text_from_xlsx() flattens the workbook to
plain text (openpyxl, mechanical row-by-row extraction, never LLM-guessed)
before falling back to the error.
"""
from __future__ import annotations

import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "system" / "scripts"))
sys.path.insert(0, str(ROOT / "system" / "api"))

try:
    import openpyxl
    _HAS_OPENPYXL = True
except ImportError:
    _HAS_OPENPYXL = False

import server  # noqa: E402


@unittest.skipUnless(_HAS_OPENPYXL, "openpyxl not installed")
class TestExtractTextFromXlsx(unittest.TestCase):
    def _make_workbook_bytes(self, rows: list[tuple], sheet_title: str = "Sheet1") -> bytes:
        import io
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_title
        for row in rows:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_extracts_real_row_content(self):
        content = self._make_workbook_bytes([
            ("Rank", "Account", "Score"),
            (1, "DEL TACO LLC", 83),
            (2, "FIVE GUYS", 82),
        ], sheet_title="Executive Summary")
        text = server._extract_text_from_xlsx(content)
        self.assertIn("## Sheet: Executive Summary", text)
        self.assertIn("Rank | Account | Score", text)
        self.assertIn("DEL TACO LLC", text)
        self.assertIn("FIVE GUYS", text)

    def test_skips_fully_empty_rows(self):
        content = self._make_workbook_bytes([
            ("A", "B"),
            (None, None),
            ("C", "D"),
        ])
        text = server._extract_text_from_xlsx(content)
        lines = [l for l in text.splitlines() if l and not l.startswith("##")]
        self.assertEqual(lines, ["A | B", "C | D"])

    def test_multiple_sheets_each_get_a_header(self):
        import io
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Summary"
        ws1.append(("x",))
        ws2 = wb.create_sheet("Detail")
        ws2.append(("y",))
        buf = io.BytesIO()
        wb.save(buf)
        text = server._extract_text_from_xlsx(buf.getvalue())
        self.assertIn("## Sheet: Summary", text)
        self.assertIn("## Sheet: Detail", text)

    def test_invalid_bytes_return_empty_string_not_raise(self):
        self.assertEqual(server._extract_text_from_xlsx(b"not a real xlsx file"), "")

    def test_empty_workbook_returns_empty_string(self):
        content = self._make_workbook_bytes([])
        self.assertEqual(server._extract_text_from_xlsx(content), "")

    def test_real_worldpay_file_if_present(self):
        """The actual file from the live incident, if still in the inbox --
        confirms the fix against the real artifact, not just a synthetic one."""
        real_file = ROOT / "system" / "inbox" / "user_artifacts" / "Worldpay_Master_Account_Plan_2026-08-14.xlsx"
        if not real_file.exists():
            self.skipTest("real incident file not present in this environment")
        text = server._extract_text_from_xlsx(real_file.read_bytes())
        self.assertGreater(len(text), 1000, "real account-plan workbook should extract substantial text")
        self.assertIn("Worldpay", text)


if __name__ == "__main__":
    unittest.main()
