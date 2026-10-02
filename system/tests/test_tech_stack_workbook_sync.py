#!/usr/bin/env python3
"""test_tech_stack_workbook_sync.py — RB-2026-08-29.

tech_stack_workbook_sync.py bridges RB's ecosystem_intelligence.json graph
and Todd's shared "Restaurant Tech Coverage" workbook (backfill: RB -> the
workbook's blank cells; ingest: the workbook's "Evidence Ledger" sheet ->
the graph, reusing ecosystem_intelligence.py's proven conflict-resolution
engine).

Building and live-verifying this against the real workbook found three
real, confirmed bugs on the first --confirm run, all covered here:
  1. ecosystem_intelligence.py's _xlsx_sheet_name_map() stripped the
     required "xl/" prefix from a sheet's zip-member path, breaking every
     named-sheet load (migrate-workbook, import-phase2-evidence, and this
     module) against a workbook whose rels didn't happen to already carry
     that prefix.
  2. Evidence Ledger rows have no "Units"/location-count column at all --
     feeding that gap straight into reconcile_workbook_row produced a
     source_assertion with live_locations=None that _assertion_key's
     content-hash dedup correctly read as a genuine correction, appending a
     redundant, strictly-worse duplicate to 53 of 54 touched relationships.
  3. A row that can't infer a specific vendor_role correctly falls back to
     "unknown" -- but the shared merge treats that as a real, overwriting
     value rather than "no new signal," downgrading 4 already-specific
     vendor_roles to "unknown".

Both #2 and #3 were also entangled with a fourth, separate pre-existing
issue: "Checkers" and "Checkers & Rally's" are two undeduped brand
entities in the live graph, so a plain exact-name lookup silently resolves
to the wrong (empty) duplicate instead of falling through to disambiguation.
"""
from __future__ import annotations

import io
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

_SCRIPTS = Path(__file__).resolve().parent.parent / "scripts"
sys.path.insert(0, str(_SCRIPTS))

import openpyxl  # noqa: E402

import ecosystem_intelligence as eco  # noqa: E402
import tech_stack_workbook_sync as sync  # noqa: E402


def _minimal_graph() -> dict:
    return {
        "version": 1, "contract": "rb_ecosystem_intelligence_v1", "last_updated": "2026-08-29",
        "entities": [], "relationships": [], "sources": [],
    }


class TestXlsxSheetNameMapPrefixDefect(unittest.TestCase):
    """RB-DEFECT-2026-08-29: _xlsx_sheet_name_map must always return a
    real, zf.read()-able archive path, regardless of whether the
    underlying rels Target happened to already carry the "xl/" prefix."""

    def _multi_sheet_workbook(self) -> Path:
        wb = openpyxl.Workbook()
        wb.active.title = "First"
        wb.active.append(("a", "b"))
        for name in ("Second", "Third", "Fourth", "Fifth"):
            ws = wb.create_sheet(name)
            ws.append((f"{name} col1", f"{name} col2"))
        tmp = Path(tempfile.mkdtemp()) / "multi.xlsx"
        wb.save(tmp)
        return tmp

    def test_named_sheet_loads_regardless_of_position(self):
        path = self._multi_sheet_workbook()
        name_map = eco._xlsx_sheet_name_map(path)
        self.assertEqual(len(name_map), 5)
        for target in name_map.values():
            self.assertTrue(target.startswith("xl/"), f"non-archive-relative path: {target}")

    def test_load_xlsx_sheet_by_name_reads_real_content(self):
        path = self._multi_sheet_workbook()
        rows = eco.load_rows(path, sheet_name="Fourth", required_header_keys=("fourth_col1",))
        self.assertEqual(rows, [])  # header-only sheet, no data rows -- but must not raise KeyError


class _IsolatedGraphMixin:
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        tmp_dir = Path(self._tmp.name)
        self._graph_path = tmp_dir / "ecosystem_intelligence.json"
        self._snap_dir = tmp_dir / "_snapshots"
        self._orig_graph_path = eco.core.ECOSYSTEM_INTELLIGENCE_PATH
        self._orig_snap_dir = eco.core.SNAPSHOTS_DIR
        eco.core.ECOSYSTEM_INTELLIGENCE_PATH = self._graph_path
        eco.core.SNAPSHOTS_DIR = self._snap_dir
        self._patch_validator = patch("subprocess.run")
        mock_run = self._patch_validator.start()
        mock_run.return_value.returncode = 0

    def tearDown(self):
        self._patch_validator.stop()
        eco.core.ECOSYSTEM_INTELLIGENCE_PATH = self._orig_graph_path
        eco.core.SNAPSHOTS_DIR = self._orig_snap_dir
        self._tmp.cleanup()

    def _write_graph(self, graph: dict) -> None:
        import json
        self._graph_path.write_text(json.dumps(graph), encoding="utf-8")


class TestBackfill(_IsolatedGraphMixin, unittest.TestCase):
    def _workbook_with_canonical_tech_stack(self, existing_pos: str | None = None) -> Path:
        wb = openpyxl.Workbook()
        wb.active.title = "Canonical Tech Stack"
        ws = wb.active
        ws.append(("Canonical Restaurant Technology Stack",))
        ws.append(("desc",))
        ws.append(())
        headers = ["Company"] + [None] * 4 + ["POS", "POS Confidence"] + [None] * 45 + [
            "Last Verified", "Evidence Sources", "Stack Coverage", "Research State", "Notes",
        ]
        ws.append(headers)
        row = ["Test Brand"] + [None] * 4 + [existing_pos, 1.0 if existing_pos else None] + [None] * 45 + [
            None, None, None, "Research gap", None,
        ]
        ws.append(row)
        tmp = Path(tempfile.mkdtemp()) / "wb.xlsx"
        wb.save(tmp)
        return tmp

    def test_blank_pos_cell_filled_from_graph(self):
        graph = _minimal_graph()
        graph["entities"] = [
            {"id": "brand-test-brand", "name": "Test Brand", "entity_type": "brand", "aliases": []},
        ]
        graph["relationships"] = [{
            "from_entity_id": "brand-test-brand", "to_entity_id": "vendor-realpos",
            "status": "active", "category": "pos", "product": "RealPOS",
            "confidence": {"level": "high"},
        }]
        self._write_graph(graph)
        wb_path = self._workbook_with_canonical_tech_stack()
        out_path = wb_path.with_name("out.xlsx")
        result = sync.backfill(wb_path, out_path)
        self.assertEqual(result["cells_filled"], 1)
        out_wb = openpyxl.load_workbook(out_path)
        ws = out_wb["Canonical Tech Stack"]
        self.assertEqual(ws.cell(5, 6).value, "RealPOS")

    def test_existing_cell_never_overwritten(self):
        graph = _minimal_graph()
        graph["entities"] = [
            {"id": "brand-test-brand", "name": "Test Brand", "entity_type": "brand", "aliases": []},
        ]
        graph["relationships"] = [{
            "from_entity_id": "brand-test-brand", "to_entity_id": "vendor-realpos",
            "status": "active", "category": "pos", "product": "RealPOS",
            "confidence": {"level": "high"},
        }]
        self._write_graph(graph)
        wb_path = self._workbook_with_canonical_tech_stack(existing_pos="AlreadyFilledPOS")
        out_path = wb_path.with_name("out.xlsx")
        result = sync.backfill(wb_path, out_path)
        self.assertEqual(result["cells_filled"], 0)
        out_wb = openpyxl.load_workbook(out_path)
        ws = out_wb["Canonical Tech Stack"]
        self.assertEqual(ws.cell(5, 6).value, "AlreadyFilledPOS")


class TestUpdateOnlyRoundTrip(unittest.TestCase):
    def _master_workbook(self) -> Path:
        wb = openpyxl.Workbook()
        wb.active.title = "Brand Technology Updates"
        updates = wb.active
        updates.append(("Brand Technology Update Queue",))
        updates.append(("desc",))
        updates.append(("workflow",))
        updates.append(())
        updates.append([
            "Update ID", "Brand", "Technology Field", "Current Value", "Current Confidence",
            "Proposed Value", "Proposed Confidence", "Evidence Source / URL", "Source Date",
            "Rationale / Notes", "Requested By", "Review Status", "Applied At", "Applied By",
        ])
        updates.append([
            '=IF(AND($B6<>"",$C6<>""),"BTU-"&TEXT(ROW()-5,"0000"),"")',
            "Old Brand", "POS", "", "", "Old POS", 0.5, "https://old.example", "2026-01-01",
            "old note", "tester", "Draft", "", "",
        ])
        canon = wb.create_sheet("Canonical Tech Stack")
        canon.append(("Canonical Restaurant Technology Stack",))
        canon.append(("desc",))
        canon.append(())
        canon.append([
            "Company", "Units", "Franchisee-Owned Units", "Company-Owned Units", "Ownership Confidence",
            "POS", "POS Confidence", "Back Office", "Back Office Confidence", "Payments", "Payments Confidence",
            "POS Hardware", "POS Hardware Confidence", "Drive-Thru Timers", "Timer Confidence",
            "Digital Menu Boards", "DMB Confidence", "Loyalty", "Loyalty Confidence",
            "Online Ordering", "Ordering Confidence", "KDS", "KDS Confidence",
            "Last Verified", "Evidence Sources", "Stack Coverage", "Research State", "Notes",
        ])
        canon.append(["Alpha Brand", 10, 0, 10, 1.0, "Existing POS", 0.9])
        canon.append(["Beta Brand", 20, 0, 20, 1.0])
        fmap = wb.create_sheet("Field Map")
        fmap.append(("Technology Field", "Confidence Field"))
        fmap.append(("POS", "POS Confidence"))
        fmap.append(("Loyalty", "Loyalty Confidence"))
        fmap.append(("Online Ordering", "Ordering Confidence"))
        tmp = Path(tempfile.mkdtemp()) / "master.xlsx"
        wb.save(tmp)
        return tmp

    def test_export_updates_creates_lightweight_editable_workbook(self):
        master = self._master_workbook()
        out = master.with_name("updates_only.xlsx")
        result = sync.export_updates(master, out)
        self.assertEqual(result["sheet"], "Brand Technology Updates")

        wb = openpyxl.load_workbook(out)
        self.assertIn("Brand Technology Updates", wb.sheetnames)
        self.assertIn("_Lists", wb.sheetnames)
        self.assertEqual(wb["_Lists"].sheet_state, "hidden")
        ws = wb["Brand Technology Updates"]
        self.assertTrue(ws.protection.sheet)
        self.assertFalse(ws["B6"].protection.locked)
        self.assertFalse(ws["H6"].protection.locked)
        self.assertFalse(ws["K6"].protection.locked)
        self.assertTrue(ws["D6"].protection.locked)
        self.assertFalse(str(ws["D6"].value or "").startswith("="))
        self.assertFalse(str(ws["E6"].value or "").startswith("="))
        self.assertGreaterEqual(len(ws.data_validations.dataValidation), 3)

    def test_import_updates_only_copies_user_fields_and_preserves_canonical(self):
        master = self._master_workbook()
        export_path = master.with_name("updates_only.xlsx")
        sync.export_updates(master, export_path)

        wb = openpyxl.load_workbook(export_path)
        ws = wb["Brand Technology Updates"]
        ws["B6"] = "Alpha Brand"
        ws["C6"] = "Loyalty"
        ws["F6"] = "Official Alpha rewards app"
        ws["G6"] = 0.65
        ws["H6"] = "https://apps.example/alpha"
        ws["I6"] = "2026-08-31"
        ws["J6"] = "App-store clue only."
        ws["K6"] = "Field user"
        ws["L6"] = "Ready for Review"
        wb.save(export_path)

        out = master.with_name("master_imported.xlsx")
        result = sync.import_updates(master, export_path, out)
        self.assertEqual(result["rows_imported"], 1)

        imported = openpyxl.load_workbook(out, data_only=False)
        updates = imported["Brand Technology Updates"]
        self.assertEqual(updates["B6"].value, "Alpha Brand")
        self.assertEqual(updates["C6"].value, "Loyalty")
        self.assertEqual(updates["F6"].value, "Official Alpha rewards app")
        self.assertEqual(updates["H6"].value, "https://apps.example/alpha")
        self.assertTrue(str(updates["D6"].value).startswith("=IFERROR(INDEX"))
        self.assertFalse(updates["H6"].protection.locked)

        canon = imported["Canonical Tech Stack"]
        self.assertEqual(canon["F5"].value, "Existing POS")

    def test_import_updates_sanitizes_contributor_text_without_corrupting_formulas(self):
        """RB-SECURITY-2026-09-05: a contributor's free-text field (F6) is
        real, externally-influenceable input flowing back into the
        canonical master -- same formula/CSV-injection guard (CWE-1236)
        already applied to Blue Sheets. The regenerated D/E formula
        columns (this file's own intentional formula cells) must survive
        completely untouched -- the fix must not blanket-sanitize
        everything in the row, only the genuine user-entry columns."""
        master = self._master_workbook()
        export_path = master.with_name("updates_only.xlsx")
        sync.export_updates(master, export_path)

        wb = openpyxl.load_workbook(export_path)
        ws = wb["Brand Technology Updates"]
        ws["B6"] = "Alpha Brand"
        ws["C6"] = "Loyalty"
        ws["F6"] = "=1+1"  # a formula-trigger-prefixed contributor free-text value
        ws["G6"] = 0.65
        ws["L6"] = "Ready for Review"
        wb.save(export_path)

        out = master.with_name("master_imported_evil.xlsx")
        sync.import_updates(master, export_path, out)

        imported = openpyxl.load_workbook(out, data_only=False)
        updates = imported["Brand Technology Updates"]
        self.assertEqual(updates["F6"].value, "'=1+1", "must be written as literal text, not a live formula")
        # The regenerated formula columns must be completely unaffected.
        self.assertTrue(str(updates["D6"].value).startswith("=IFERROR(INDEX"))
        self.assertTrue(str(updates["E6"].value).startswith("=IFERROR(INDEX"))
        self.assertTrue(str(updates["A6"].value).startswith("=IF(AND("))


class TestIngestUnitsInheritance(_IsolatedGraphMixin, unittest.TestCase):
    """RB-DEFECT-2026-08-29 bugs #2-4."""

    def _evidence_ledger_workbook(self, *, brand="Checkers", vendor="Presto", tech_category="Drive-Thru Voice AI") -> Path:
        wb = openpyxl.Workbook()
        wb.active.title = "Evidence Ledger"
        ws = wb.active
        ws.append(("Technology Evidence Ledger",))
        ws.append(("desc",))
        ws.append(())
        ws.append([
            "Brand", "Parent / Platform", "Tech Category", "AI Application", "Vendor",
            "Product / Module", "Relationship Scope", "Lifecycle / Current-State Flag",
            "Evidence Type", "Verification Status", "Evidence Date", "Source Strength",
            "Scope Clarity", "Recency", "Corroboration", "Confidence Score",
            "Confidence Label", "Graph Eligible", "Source URL", "Notes", "Stack Key",
        ])
        ws.append([
            # Deliberately neutral relationship-scope text -- no keyword
            # _infer_workbook_vendor_role() recognizes (no "franchisee",
            # "resell", "approved hardware", etc.) -- this row must fall
            # through to the generic "unknown" default so the test actually
            # exercises the "don't let unknown overwrite a known role" fix,
            # not a real (if different) role inference.
            brand, "", tech_category, "", vendor, f"{vendor} Voice",
            "Named customer reference", "Historical / superseded",
            "investor_filing_deck", "verified", "2024", 90, 45, 100, 85,
            0.7, "medium", "Yes", "https://example.com/filing", "Some evidence note.",
            f"{brand} | {tech_category} | {vendor} | {vendor} Voice",
        ])
        tmp = Path(tempfile.mkdtemp()) / "evidence.xlsx"
        wb.save(tmp)
        return tmp

    def _graph_with_dedupe_and_existing_relationship(self) -> dict:
        graph = _minimal_graph()
        graph["entities"] = [
            # An empty duplicate entity, same shape as the real
            # "Checkers"/"Checkers & Rally's" split -- a plain exact-name
            # index would wrongly resolve to this one.
            {"id": "brand-checkers", "name": "Checkers", "entity_type": "brand", "aliases": []},
            {"id": "brand-checkers-rally-s", "name": "Checkers & Rally's", "entity_type": "brand",
             "aliases": ["Checkers & Rallys"]},
        ]
        graph["relationships"] = [{
            "id": "rel-brand-checkers-rally-s-drive-thru-ai-hardware-reseller-service-provider-vendor-presto",
            "from_entity_id": "brand-checkers-rally-s", "to_entity_id": "vendor-presto",
            "status": "historical", "category": "drive_thru_ai", "product": "Presto Voice",
            "vendor_role": "hardware_reseller_service_provider",
            "confidence": {"level": "medium", "score": 1.0},
            "sources": ["src-existing"],
            "source_assertions": [{
                "source_id": "src-existing", "url": "https://example.com/filing",
                "live_locations": 484.0, "posture": "historical",
            }],
        }]
        return graph

    def test_missing_units_inherited_not_left_null(self):
        graph = self._graph_with_dedupe_and_existing_relationship()
        self._write_graph(graph)
        wb_path = self._evidence_ledger_workbook()
        result = sync.ingest(wb_path, confirm=True)
        self.assertEqual(result["rows_processed"], 1)

        import json
        updated_graph = json.loads(self._graph_path.read_text())
        rel = next(r for r in updated_graph["relationships"] if r["to_entity_id"] == "vendor-presto")
        # No spurious duplicate source_assertion with live_locations=None.
        null_loc_entries = [a for a in rel["source_assertions"] if a.get("live_locations") is None]
        self.assertEqual(null_loc_entries, [], "units should have been inherited, not left null")

    def test_vendor_role_not_downgraded_to_unknown(self):
        graph = self._graph_with_dedupe_and_existing_relationship()
        self._write_graph(graph)
        wb_path = self._evidence_ledger_workbook()
        sync.ingest(wb_path, confirm=True)

        import json
        updated_graph = json.loads(self._graph_path.read_text())
        rel = next(r for r in updated_graph["relationships"] if r["to_entity_id"] == "vendor-presto")
        self.assertEqual(rel["vendor_role"], "hardware_reseller_service_provider")

    def test_resolves_short_brand_name_to_canonical_deduped_entity(self):
        """"Checkers" (Evidence Ledger's raw text) must resolve to the same
        entity the real relationship lives under ("Checkers & Rally's"),
        not the separate, empty "Checkers" duplicate."""
        graph = self._graph_with_dedupe_and_existing_relationship()
        self._write_graph(graph)
        wb_path = self._evidence_ledger_workbook()
        sync.ingest(wb_path, confirm=True)

        import json
        updated_graph = json.loads(self._graph_path.read_text())
        empty_dup_rels = [
            r for r in updated_graph["relationships"] if r["from_entity_id"] == "brand-checkers"
        ]
        self.assertEqual(empty_dup_rels, [], "must not create/attach a relationship to the wrong duplicate entity")

    def test_dry_run_does_not_write(self):
        graph = self._graph_with_dedupe_and_existing_relationship()
        self._write_graph(graph)
        wb_path = self._evidence_ledger_workbook()
        before = self._graph_path.read_text()
        sync.ingest(wb_path, confirm=False)
        after = self._graph_path.read_text()
        self.assertEqual(before, after)


if __name__ == "__main__":
    unittest.main()
