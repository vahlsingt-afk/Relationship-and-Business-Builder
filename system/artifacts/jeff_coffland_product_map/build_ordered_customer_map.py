import importlib.util
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

spec = importlib.util.spec_from_file_location("base", "system/artifacts/jeff_coffland_product_map/build_product_map.py")
b = importlib.util.module_from_spec(spec)
spec.loader.exec_module(b)

OUT = "system/artifacts/jeff_coffland_product_map/Genius_Product_Map_Jeff_Coffland.docx"
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.62); sec.bottom_margin = Inches(0.58)
sec.left_margin = Inches(0.72); sec.right_margin = Inches(0.72)
normal = doc.styles["Normal"]
normal.font.name = "Arial"; normal.font.size = Pt(9.2)
normal.font.color.rgb = RGBColor.from_string(b.INK)
normal.paragraph_format.space_after = Pt(4); normal.paragraph_format.line_spacing = 1.05
for level, size in ((1, 22), (2, 15), (3, 11)):
    st = doc.styles[f"Heading {level}"]
    st.font.name = "Arial"; st.font.size = Pt(size); st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(b.NAVY if level < 3 else b.BLUE)
    st.paragraph_format.space_before = Pt(7 if level > 1 else 0)
    st.paragraph_format.space_after = Pt(5)
hp = sec.header.paragraphs[0]
b.add_text(hp, "globalpayments", 13, b.NAVY, True)
b.add_text(hp, "  /  GENIUS", 9, b.BLUE, True)
b.add_page_number(sec.footer.paragraphs[0])

# Cover
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(42); p.paragraph_format.space_after = Pt(4)
b.add_text(p, "PRODUCT PORTFOLIO MAP", 9, b.CYAN, True)
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(10)
b.add_text(p, "Genius restaurant\ntechnology", 31, b.NAVY, True)
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(22)
b.add_text(p, "McDonald’s capability overview", 17, b.BLUE, True)
b.callout(doc, "One connected restaurant technology portfolio", "Genius brings enterprise commerce, kitchen operations, digital experience, drive-thru, self-service, purpose-built hardware, payments and back-office capabilities together within a flexible platform.")
t = doc.add_table(rows=3, cols=2); t.autofit = False
t.columns[0].width = Inches(1.35); t.columns[1].width = Inches(4.8)
for row, values in enumerate((("Prepared for", "Jeff Coffland"), ("As of", "August 12, 2026"), ("Company", "Global Payments"))):
    for col, value in enumerate(values):
        c=t.cell(row,col); b.margins(c,70,80,70,80); b.shade(c,b.LIGHT if row%2==0 else b.WHITE)
        b.add_text(c.paragraphs[0],value,8.6,b.NAVY if col==0 else b.INK,col==0)
b.heading(doc,"Portfolio at a glance",2)
for x in ("Genius POS, Kitchen, Digital Menu Board, Drive Thru, Kiosk, Hardware, Payments and Back Office.",
          "Each entry summarizes product heritage, capabilities, relevance to McDonald’s and selected experience.",
          "CosMc’s is presented first where it provides the most directly relevant experience."):
    b.bullet(doc,x)

def page(title, number):
    doc.add_page_break(); b.heading(doc,title,1,f"Product map {number} of 5")

# 1. POS / 2. Kitchen
page("Commerce and kitchen operations",1)
b.product(doc,"Genius POS","", "Xenial enterprise POS and Xenial Ordering, now part of the Genius Cloud portfolio.",
    "Cloud enterprise POS with centralized management, real-time synchronization, device and operating-system flexibility, rapid terminal failover and offline operation for up to 30 days.",
    "A flexible restaurant-commerce and integration platform that can connect front-of-house channels with enterprise services.",
    "CosMc’s selected Genius Cloud POS as part of its integrated restaurant technology platform. Additional enterprise experience includes A&W Canada across 1,032 restaurants and 7 Brew across more than 530 locations.", "")
b.product(doc,"Genius Kitchen","", "Xenial Kitchen Management (XKM), Kitchen Management and Order Management System (OMS), now within the Genius portfolio.",
    "Kitchen displays route items to configured production stations, support drive-thru and other destination-specific views, provide order-ready status and support configurable touchscreen or bump-bar workflows.",
    "Connects order channels with production workflows to support throughput, order accuracy and operational visibility.",
    "Genius Kitchen is supported by active product documentation and a continuing Xenial Kitchen release history.", "")

# 3. Digital Menu Board / 4. Drive Thru
page("Digital menu boards and drive-thru",2)
b.product(doc,"Genius Digital Menu Board","", "SICOM DMB, first developed in 2009, was replaced by the next-generation Genius digital menu platform following Global Payments’ acquisition of SICOM.",
    "Centralized content management, scheduled dayparts, local and corporate control, indoor/outdoor menus and order-confirmation boards. System-on-chip displays can retain scheduled content during an internet outage.",
    "Offers a flexible digital-menu architecture for targeted indoor, drive-thru and order-confirmation requirements.",
    "CosMc’s selected Genius indoor and drive-thru digital menu boards as part of its integrated restaurant technology platform. Yoshinoya deployed digital menu boards at 106 locations.", "")
b.product_visual(doc,f"{b.PRODUCT_ASSET_DIR}/s21_shape2.png","Representative Genius system-on-chip digital menu board configuration",3.8)
b.product(doc,"Genius Drive Thru","", "SICOM Timer → Drive-Thru Director → Vision, the Genius next-generation drive-thru timer platform.",
    "Lane events from detectors, cameras and order-confirmation units; real-time vehicle visualization; single/dual-lane workflows; reporting and APIs; zone configuration for queue, pickup and pull-forward activity.",
    "Supports speed-of-service visibility and a more complete operational view of the drive-thru journey.",
    "CosMc’s selected Genius drive-thru timer technology as part of its integrated platform. At Yoshinoya, 34 drive-thru locations paired menu boards with a timer and reported a 35% improvement in service time.", "")

# 5. Kiosk
page("Self-service ordering",3)
b.product(doc,"Genius Kiosk","", "Xenial Kiosk hardware and software and Xenial Ordering Kiosk, now within the Genius portfolio.",
    "Indoor/outdoor self-service ordering with multiple mounting options, favorites and repeat orders, multilingual and accessibility features, dynamic splash screens and pickup workflows.",
    "Provides modular self-service and integration options for restaurant environments, subject to McDonald’s store and accessibility standards.",
    "CosMc’s selected Genius indoor and outdoor kiosks as part of its integrated platform. Genius brings more than 20 years of kiosk development and 15,000+ self-service installations worldwide.", "")
b.product_visual(doc,f"{b.PRODUCT_ASSET_DIR}/s22_shape4.png","Freestanding Genius self-service kiosks in a restaurant environment",5.3)

# 6. Hardware
page("Genius hardware",4)
b.product(doc,"Genius Hardware","", "Xenial and SICOM POS hardware, portable tablets and rugged terminals, unified under the Genius hardware identity.",
    "Global Payments designs and owns the core Genius hardware specification, with Flytech as manufacturing partner. The portfolio supports modular countertop, handheld, kiosk and tethered configurations.",
    "Provides a consistent device experience across counter, line-busting and self-service workflows, subject to McDonald’s hardware, certification and service requirements.",
    "CosMc’s selected Genius hardware across its cloud POS and indoor/outdoor self-service environment.", "")
t=doc.add_table(rows=1,cols=3); t.autofit=False
visuals=(("Countertop",f"{b.HW_DIR}/s14_shape1_Google Shape;2008;p267.png","Modular merchant and customer-facing displays."),
         ("Handheld",f"{b.HW_DIR}/s17_shape3_Google Shape;2049;p270.png","Portable POS with integrated payment."),
         ("Counter and kiosk",f"{b.HW_DIR}/s18_shape17_Google Shape;2091;p271.png","Multiple configurations within one hardware family."))
for i,(label,path,cap) in enumerate(visuals):
    c=t.cell(0,i); b.margins(c,120,105,120,105); b.shade(c,b.WHITE)
    p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; b.add_text(p,label,10.5,b.NAVY,True)
    p=c.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(path,width=Inches(1.75))
    p=c.add_paragraph(); b.add_text(p,cap,8,b.INK)

# 7. Payments / 8. Back Office
page("Payments and back office",5)
b.product(doc,"Genius Payments","", "Global Payments Integrated and Portico capabilities, with Merchantware/Cayan heritage and Global Payments + Worldpay processing reach.",
    "Integrated acceptance across counter, drive-thru, kiosk and digital commerce; EMV/contactless and wallets; card-present and online acceptance; P2PE/tokenization; terminal management and store-and-forward.",
    "Supports cross-channel orchestration, processor flexibility, centralized device control, offline resilience and token continuity.",
    "CosMc’s used Global Payments for payment processing; FreedomPay supplied the gateway for its unattended-payment requirement.", "")
b.product(doc,"Genius Back Office","", "The Portal RTI, RTI Financial Accounting, Back Office and enterprise reporting heritage.",
    "Back-office and enterprise reporting capabilities spanning restaurant operations, inventory, labor, financial information and enterprise-level visibility.",
    "Connects restaurant activity with operational and enterprise reporting to support more consistent decision-making across the system.",
    "Global Payments continues to support customers across The Portal RTI, RTI Financial Accounting and Back Office product families.", "")

doc.core_properties.title="Genius Restaurant Technology Product Map"
doc.core_properties.subject="McDonald’s capability overview for Jeff Coffland"
doc.core_properties.author="Todd Vahlsing"
doc.save(OUT)
print(OUT)
