from pptx import Presentation


deck = Presentation("/Users/toddvahlsing/Downloads/Global Payments Integrated.pptx")
print("slides", len(deck.slides))
for index, slide in enumerate(deck.slides, 1):
    text = []
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text.strip():
            text.append(shape.text.strip().replace("\n", " | "))
    print(f"{index:03d}: " + " || ".join(text))
