from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = "system/artifacts/jeff_coffland_product_map/Genius_Product_Map_Jeff_Coffland.docx"
HW_DIR = "system/artifacts/jeff_coffland_product_map/genius_n_review/hardware_assets"
PRODUCT_ASSET_DIR = "system/artifacts/jeff_coffland_product_map/genius_n_review/product_assets"

NAVY = "14233B"
BLUE = "2764FF"
CYAN = "00A6D6"
PALE = "EAF0FF"
LIGHT = "F4F6F8"
MID = "64748B"
GREEN = "17745A"
AMBER = "9A6700"
RED = "A12626"
WHITE = "FFFFFF"
INK = "1D2939"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc.get_or_add_tcPr()
    tc_mar = tc.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = tc_mar.find(qn(f"w:{tag}"))
        if el is None:
            el = OxmlElement(f"w:{tag}")
            tc_mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("CONFIDENTIAL  •  ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_text(p, text, size=9.2, color=INK, bold=False, italic=False):
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    r.bold = bold
    r.italic = italic
    return r


def heading(doc, text, level=1, kicker=None):
    if kicker:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        add_text(p, kicker.upper(), 8, CYAN, True)
        keep_with_next(p)
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    keep_with_next(p)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2.5)
    add_text(p, text)
    return p


def callout(doc, title, body, fill=PALE, accent=BLUE):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.columns[0].width = Inches(0.10)
    t.columns[1].width = Inches(6.40)
    shade(t.cell(0, 0), accent)
    shade(t.cell(0, 1), fill)
    margins(t.cell(0, 0), 0, 0, 0, 0)
    margins(t.cell(0, 1), 120, 150, 120, 150)
    p = t.cell(0, 1).paragraphs[0]
    add_text(p, title + "\n", 9.2, NAVY, True)
    add_text(p, body, 9.2, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def product(doc, name, current, lineage, validated, relevance, proof, confirm, status="PUBLICLY GROUNDED"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    keep_with_next(p)
    add_text(p, name, 13, NAVY, True)

    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.columns[0].width = Inches(1.28)
    t.columns[1].width = Inches(5.22)
    items = [
        ("Heritage", lineage),
        ("Capabilities", validated),
        ("McDonald’s relevance", relevance),
        ("Selected experience", proof),
    ]
    for i, (label, value) in enumerate(items):
        cells = t.add_row().cells
        for c in cells:
            margins(c, 70, 95, 70, 95)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        shade(cells[0], LIGHT if i % 2 == 0 else "E9EDF2")
        shade(cells[1], WHITE if i % 2 == 0 else "FAFBFC")
        p0 = cells[0].paragraphs[0]
        add_text(p0, label, 8, NAVY, True)
        p1 = cells[1].paragraphs[0]
        add_text(p1, value, 8.35, INK)


def product_visual(doc, path, caption, width=4.8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(path, width=Inches(width))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    add_text(p, caption, 8.2, MID, False, True)


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.62)
sec.bottom_margin = Inches(0.58)
sec.left_margin = Inches(0.72)
sec.right_margin = Inches(0.72)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(9.2)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.05
for level, size in ((1, 22), (2, 15), (3, 11)):
    st = styles[f"Heading {level}"]
    st.font.name = "Arial"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(NAVY if level < 3 else BLUE)
    st.paragraph_format.space_before = Pt(7 if level > 1 else 0)
    st.paragraph_format.space_after = Pt(5)

header = sec.header
header.is_linked_to_previous = False
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
add_text(hp, "globalpayments", 13, NAVY, True)
add_text(hp, "  /  GENIUS", 9, BLUE, True)
footer = sec.footer
add_page_number(footer.paragraphs[0])

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(42)
p.paragraph_format.space_after = Pt(4)
add_text(p, "PRODUCT PORTFOLIO MAP", 9, CYAN, True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
add_text(p, "Genius restaurant\ntechnology", 31, NAVY, True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(22)
add_text(p, "McDonald’s routing reference", 17, BLUE, True)

callout(doc, "One connected restaurant technology portfolio", "Genius brings enterprise commerce, self-service, kitchen, digital experience, drive-thru, payments and purpose-built hardware together within a flexible platform.")

t = doc.add_table(rows=4, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.LEFT
t.autofit = False
t.columns[0].width = Inches(1.35)
t.columns[1].width = Inches(4.8)
for row, values in enumerate([
    ("Prepared for", "Jeff Coffland — McDonald’s capability overview"),
    ("As of", "August 12, 2026"),
    ("Portfolio", "Genius restaurant technology"),
    ("Company", "Global Payments"),
]):
    for col, value in enumerate(values):
        c = t.cell(row, col)
        margins(c, 70, 80, 70, 80)
        shade(c, LIGHT if row % 2 == 0 else WHITE)
        add_text(c.paragraphs[0], value, 8.6, NAVY if col == 0 else INK, col == 0)

doc.add_paragraph().paragraph_format.space_after = Pt(8)
heading(doc, "How to read this map", 2)
bullet(doc, "Each entry identifies the current product, its former Xenial/SICOM/RTI heritage, supported capabilities and relevance to McDonald’s.")
bullet(doc, "Customer examples show demonstrated fit; they do not imply McDonald’s approval or a production entitlement.")

doc.add_page_break()

# Core commerce
heading(doc, "Core commerce and self-service", 1, "Product map 1 of 4")
product(
    doc,
    "Genius Enterprise POS",
    "Genius Enterprise POS",
    "Xenial enterprise POS and Xenial Ordering, now part of the Genius Cloud portfolio.",
    "Cloud enterprise POS with centralized management, real-time synchronization, device and operating-system flexibility, rapid terminal failover and offline operation for up to 30 days.",
    "Relevant as an integration and restaurant-commerce platform reference. Do not position it as a proposed NEWPOS replacement without an explicit McDonald’s request.",
    "CosMc’s selected Genius Cloud POS as part of its integrated restaurant technology platform. Additional enterprise experience includes A&W Canada across 1,032 restaurants and 7 Brew across more than 530 locations.",
    "Current sellable edition/SKU; exact Xenial lineage; offline/failover claims; supported markets and service model; permission to use 7 Brew metrics; what POS version CosMc’s used.",
)
product(
    doc,
    "Genius Kiosk",
    "Genius Kiosk / enterprise self-service ordering",
    "Former public label: Xenial Kiosk hardware and software / Xenial Ordering Kiosk.",
    "Self-service ordering with indoor/outdoor and multiple mounting options, favorites and repeat orders, multilingual and accessibility features, dynamic splash screens and pickup workflows.",
    "Relevant where McDonald’s wants modular self-service or integration options, subject to global store standards and accessibility requirements.",
    "More than 20 years of kiosk development and 15,000+ self-service solutions installed worldwide.",
    "Current external name/configurations; validation and external-use permission for the 20+ year and 15,000+ claims; vendors; accessibility/security; payments; CosMc’s use; service coverage.",
)

doc.add_page_break()

# Operations
heading(doc, "Kitchen, back office and payments", 1, "Product map 2 of 4")
product(
    doc,
    "Genius Kitchen",
    "Genius Kitchen / Kitchen Display System (KDS)",
    "Xenial Kitchen Management (XKM), Kitchen Management and Order Management System (OMS), now within the Genius portfolio.",
    "Kitchen displays route items to configured production stations, support destination-specific views such as drive-thru, and can provide order-ready status. Public documentation describes touchscreen and bump-bar interaction plus configurable ticket views.",
    "Relevant to end-to-end order flow, throughput and integration conversations—not as an assumed replacement for McDonald’s kitchen production systems.",
    "Current Xenial documentation shows active Genius Kitchen product documentation and Xenial Kitchen release notes.",
    "Which KDS/OMS components are current, their external names, integration interfaces, hardware requirements, enterprise deployments and any CosMc’s footprint.",
)
product(
    doc,
    "Genius Back Office",
    "Genius Back Office and enterprise reporting",
    "Historical labels visible in Xenial support materials: The Portal RTI, RTI Financial Accounting, Back Office and enterprise reporting.",
    "Back-office and enterprise reporting capabilities spanning restaurant operations, inventory, labor, financial information and enterprise-level visibility.",
    "Connects restaurant activity with operational and enterprise reporting for more consistent decision-making across the system.",
    "Xenial’s current support directory still routes customers for The Portal RTI, RTI Financial Accounting and Back Office.",
    "Current product names and module boundaries; RTI lineage; supported functions and integrations; customer references; roadmap/status; owner and delivery coverage.",
    status="CONFIRM BEFORE EXTERNAL USE",
)
product(
    doc,
    "Integrated payments and gateway",
    "Global Payments Integrated (GPI) and Portico Gateway; processor connections include Global Payments and Worldpay/Vantiv",
    "Merchantware and Cayan heritage, now represented through Global Payments Integrated and Portico capabilities.",
    "One semi-integrated/API layer for counter, drive-thru, kiosk and digital commerce: EMV/contactless and wallets, card-present and online acceptance, P2PE/tokenization, hosted payment components, remote terminal management and device-based store-and-forward. Listed processor connections include Global Payments, Worldpay/Vantiv, Fiserv, Elavon and Chase.",
    "The McDonald’s story is cross-channel orchestration, processor flexibility, centralized device control, offline resilience and token continuity—not the full feature catalog.",
    "CosMc’s used Global Payments for payment processing; FreedomPay supplied the gateway for its unattended-payment requirement.",
    "Arjun/Payments: confirm names and lineage; GP versus Worldpay roles; processor/device certifications by market; unattended status; current Verifone/iMin availability; and what may be shared with Jeff.",
    status="CONFIRM BEFORE EXTERNAL USE",
)
# Experience
p = doc.add_paragraph()
p.paragraph_format.page_break_before = True
p.paragraph_format.space_after = Pt(0)
heading(doc, "Digital experience and drive-thru", 1, "Product map 3 of 4")
product(
    doc,
    "Genius Digital Displays",
    "Working current label: Genius digital menu technology / digital menu screens",
    "SICOM DMB was first developed in 2009. After Global Payments acquired SICOM in 2018, the next-generation Genius DMB was developed as a complete replacement—not an extension of the former SICOM software.",
    "Centralized content management, scheduled dayparts, local and corporate control, indoor/outdoor menus and order-confirmation boards. LG system-on-chip displays can retain scheduled content during an internet outage.",
    "A credible alternative architecture to evaluate for targeted needs. Do not characterize McDonald’s as replacing an incumbent or running a formal second-source process.",
    "CosMc’s selected Genius indoor and drive-thru digital menu boards as part of its integrated restaurant technology platform. At Yoshinoya, 106 locations deployed digital menu boards, including 34 drive-thrus paired with a timer.",
    "Current name; CMS/monitoring model; LG system-on-chip and offline-content claims; supported playerless versus controller configurations; partners; service coverage; CosMc’s footprint; reference permission.",
)
product(
    doc,
    "Genius Drive Thru",
    "Genius Drive Thru; current documentation includes Drive Thru Director and Vision",
    "SICOM Timer (1989) → rewrite (2001) → new monitor/UI solution (2014) → Drive-Thru Director with continued Genius enhancements. Vision, developed around 2022, is the 100% Genius next-generation timer.",
    "Lane events from detectors, cameras and intelligent order-confirmation units; real-time vehicle visualization; single/dual-lane workflows; gamification; reporting/API. The pitch deck shows Vision zone configuration for stack size, pickup parking and pull-forward tracking across the customer journey.",
    "Supports speed-of-service visibility and a more complete view of the drive-thru journey, including queue, pull-forward and pickup activity.",
    "CosMc’s selected Genius drive-thru timer technology as part of its integrated platform. At Yoshinoya, 34 drive-thru locations paired digital menu boards with a timer and reported a 35% improvement in service time.",
    "Current sellable components versus roadmap; supported sensor/camera/OCU architecture; reporting/API; AI claims; CosMc’s requirements and deployed scope; customer references; install/field-service ownership.",
)

# Hardware is intentionally last so the portfolio narrative leads and the device
# family closes with a tangible visual reference.
doc.add_page_break()
heading(doc, "Genius hardware", 1, "Product map 4 of 4 — hardware reference")
product(
    doc,
    "A modular in-store hardware family",
    "Genius countertop, handheld, kiosk, tethered and GP One hardware families",
    "Heritage: Xenial and SICOM POS hardware, portable tablets and rugged terminals. Configuration and naming may vary by Genius edition and geography.",
    "Global Payments designs and owns the core Genius hardware specification, with Flytech as manufacturing partner. The portfolio supports modular countertop, handheld, kiosk and tethered configurations. Payment terminals and the iMin-based Genius All In One are separate device families.",
    "Provides a consistent Genius device experience across counter, line-busting and self-service workflows, subject to McDonald’s hardware standards, certifications and support requirements.",
    "CosMc’s selected Genius hardware across its cloud POS and indoor/outdoor self-service environment. The broader portfolio provides one hardware identity across restaurant workflows.",
    "",
    status="CONFIRM BEFORE EXTERNAL USE",
)

t = doc.add_table(rows=1, cols=4)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
visuals = [
    ("Modular countertop", f"{HW_DIR}/s14_shape1_Google Shape;2008;p267.png", "Merchant and customer-facing displays with countertop, low-profile and wall configurations."),
    ("Enterprise handheld", f"{HW_DIR}/s17_shape3_Google Shape;2049;p270.png", "Portable enterprise POS for line-busting and mobile service, shown with integrated payment."),
    ("Portrait kiosk", f"{HW_DIR}/s18_shape19_Google Shape;2093;p271.png", "Portrait self-service ordering and payment configuration."),
    ("Kiosk portfolio", f"{HW_DIR}/s18_shape17_Google Shape;2091;p271.png", "Multiple kiosk and tethered configurations within the Genius hardware family."),
]
for i, (label, path, caption) in enumerate(visuals):
    c = t.cell(0, i)
    margins(c, 120, 105, 120, 105)
    shade(c, WHITE)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, label, 10.5, NAVY, True)
    p = c.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(1.38))
    p = c.add_paragraph()
    add_text(p, caption, 8.0, INK)

doc.core_properties.title = "Genius Restaurant Technology Product Map"
doc.core_properties.subject = "McDonald’s routing reference for Jeff Coffland"
doc.core_properties.author = "Todd Vahlsing"
doc.core_properties.keywords = "Genius, Global Payments, Worldpay, Xenial, RTI, McDonald's, CosMc's"
doc.save(OUT)
print(OUT)
