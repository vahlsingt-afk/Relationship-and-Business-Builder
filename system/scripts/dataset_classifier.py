#!/usr/bin/env python3
"""
dataset_classifier.py — generalized structured-dataset recognition (RB-DEFECT-064 Phase 1).

Runs on any uploaded/dropped structured file (CSV/XLSX) before conversational
handling, so RB doesn't need to be told what a file is or what to do with it.
Returns a classification (dataset type, purpose, confidence, recommended
action) which callers compare against `structured_ingest.confidence_threshold`
in settings.json to decide whether to auto-execute the matching ingest
pipeline or surface the file for operator confirmation.

Signature-based, not per-source hardcoding scattered across the codebase:
each known dataset type is one entry in DATASET_SIGNATURES describing the
filename pattern and header/sheet vocabulary that identifies it. Adding a new
recognized source is adding a signature, not writing a new detection path.

Some dataset types (LinkedIn export, Apple Contacts) already have dedicated
ingest pipelines (linkedin_ingest.py, contacts_ingest.py) predating this
classifier; they're registered here so `classify()` gives one unified answer
across every structured upload, but this module does not re-implement their
detection logic.

XLSX support (RB-DEFECT-065 Phase 1): multi-sheet workbooks are classified by
sheet-name fingerprint (e.g. the McDonald's NSN Lookup workbook's "StoreTech"/
"Markets"/"RFM" tabs) in addition to the per-row header vocabulary CSVs use,
since a single workbook's sheets rarely share one header row.

CLI:
    python3 system/scripts/dataset_classifier.py --file PATH
    python3 system/scripts/dataset_classifier.py --smoke
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import sys
from dataclasses import dataclass, field
from pathlib import Path

SCRIPTS_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPTS_DIR))
import rb_core as core  # noqa: E402

DEFAULT_CONFIDENCE_THRESHOLD = 0.85
UNKNOWN_TYPE = "unknown_structured_dataset"


@dataclass
class DatasetSignature:
    dataset_type: str
    purpose: str
    recommended_action: str
    ingest_script: str
    filename_hints: tuple[str, ...] = ()
    required_headers: tuple[str, ...] = ()          # every one must be present
    distinctive_headers: tuple[str, ...] = ()        # scored, not required
    required_sheets: tuple[str, ...] = ()            # xlsx only; every one must be present
    distinctive_sheets: tuple[str, ...] = ()         # xlsx only; scored, not required
    already_handled_by: str | None = None            # existing dedicated pipeline, if any


@dataclass
class Classification:
    dataset_type: str
    confidence: float
    purpose: str
    recommended_action: str
    ingest_script: str
    matched_signals: list[str] = field(default_factory=list)
    already_handled_by: str | None = None

    def to_dict(self) -> dict:
        return {
            "dataset_type": self.dataset_type,
            "confidence": round(self.confidence, 4),
            "purpose": self.purpose,
            "recommended_action": self.recommended_action,
            "ingest_script": self.ingest_script,
            "matched_signals": self.matched_signals,
            "already_handled_by": self.already_handled_by,
        }


# ---------------------------------------------------------------------------
# Signature registry
# ---------------------------------------------------------------------------

DATASET_SIGNATURES: tuple[DatasetSignature, ...] = (
    DatasetSignature(
        dataset_type="linkedin_connections_export",
        purpose="historical_relationship_database",
        recommended_action="run_linkedin_ingest",
        ingest_script="system/scripts/linkedin_ingest.py",
        filename_hints=("linkedindataexport", "connections"),
        required_headers=("First Name", "Last Name"),
        distinctive_headers=("URL", "Connected On"),
        already_handled_by="P-002_linkedin_ingest.md",
    ),
    DatasetSignature(
        dataset_type="apple_contacts_export",
        purpose="baseline_identity_enrichment",
        recommended_action="run_contacts_ingest",
        ingest_script="system/scripts/contacts_ingest.py",
        filename_hints=("contacts",),
        required_headers=(),
        distinctive_headers=("First Name", "Last Name", "Phone", "Company"),
        already_handled_by="RB-DEFECT-025",
    ),
    DatasetSignature(
        dataset_type="hubspot_crm_export",
        purpose="historical_relationship_database",
        recommended_action="run_hubspot_ingest",
        ingest_script="system/scripts/hubspot_ingest.py",
        filename_hints=("hubspot",),
        required_headers=("First Name", "Last Name"),
        distinctive_headers=("Contact owner", "Lead Status", "Marketing contact status"),
    ),
    DatasetSignature(
        dataset_type="salesforce_crm_export",
        purpose="historical_relationship_database",
        recommended_action="review_before_ingest",  # no dedicated pipeline yet (RB-DEFECT-064 Phase 2+)
        ingest_script="",
        filename_hints=("salesforce", "sfdc"),
        required_headers=(),
        distinctive_headers=("Account ID", "Account Name", "Lead Source", "Contact Owner"),
    ),
    DatasetSignature(
        dataset_type="conference_attendee_list",
        purpose="event_relationship_capture",
        recommended_action="run_campaign_reconcile",
        ingest_script="system/scripts/campaign_engine.py",
        filename_hints=("attendee", "registration", "roster"),
        required_headers=(),
        distinctive_headers=("Registration Date", "Ticket Type", "Attendee Name", "Event Name"),
    ),
    DatasetSignature(
        dataset_type="generic_contact_export",
        purpose="baseline_identity_enrichment",
        recommended_action="review_before_ingest",  # too generic to auto-run unattended
        ingest_script="",
        filename_hints=(),
        required_headers=("First Name", "Last Name"),
        distinctive_headers=("Email", "Company", "Title", "Phone"),
    ),
    DatasetSignature(
        dataset_type="mcdonalds_nsn_lookup_workbook",
        purpose="enterprise_micro_graph_enhancement",
        recommended_action="run_micro_graph_mcdonalds",
        ingest_script="system/scripts/micro_graph_mcdonalds.py",
        filename_hints=("nsnlookup", "nsn"),
        # StoreTech is the store-level base table micro_graph_mcdonalds.py
        # always looks for; the full set is its sheet-name fingerprint.
        required_sheets=("StoreTech",),
        distinctive_sheets=("Markets", "FO OTM-STIM", "RFM", "COOP2", "Entity", "StoreTech"),
    ),
    DatasetSignature(
        dataset_type="master_account_plan_workbook",
        purpose="vendor_portfolio_plan",
        recommended_action="run_master_account_plan_ingest",
        ingest_script="master_account_plans/_engine/create_plan.py",
        # RB-2026-08-28: real incident -- a portfolio-level, multi-account
        # vendor plan (Worldpay_Master_Account_Plan_2026-08-14.xlsx) was
        # misrouted by the model into createBlueSheetAccount (a single-
        # account, single-objective artifact), producing an empty,
        # falsely-authorized Blue Sheet. "Ranked Portfolio" (a scored,
        # multi-account table) + "RM Portfolio" (a named-relationship-
        # manager roster) together is a combination no single-account
        # Blue Sheet or Account Research document has -- required so this
        # never false-positive-matches an unrelated workbook, and so this
        # document shape is recognized and routed correctly by
        # construction rather than depending on the model's judgment.
        required_sheets=("Ranked Portfolio", "RM Portfolio"),
        distinctive_sheets=("Executive Summary", "Dashboard", "Ranked Portfolio",
                             "Stack Intelligence", "Evidence Ledger", "Conflict Register",
                             "RM Portfolio", "Scoring Model"),
    ),
)


def _normalize_filename(path: Path) -> str:
    return "".join(ch for ch in path.stem.lower() if ch.isalnum())


def _read_headers(path: Path) -> list[str]:
    """Best-effort CSV header row read."""
    try:
        text = path.read_text(encoding="utf-8-sig", errors="ignore")
    except OSError:
        return []
    reader = csv.reader(io.StringIO(text))
    for row in reader:
        if any(cell.strip() for cell in row):
            return [cell.strip() for cell in row]
    return []


def _read_xlsx_info(path: Path) -> tuple[set[str], set[str]]:
    """Best-effort sheet-name + header-row read for xlsx/xlsm.

    Read-only, row 1 of each sheet only — cheap even against large
    multi-sheet workbooks, since it never touches data rows.

    openpyxl is imported lazily so this module (and anything that imports
    it at module scope, e.g. server.py) stays importable in environments
    without it; xlsx classification just degrades to "no signal" instead
    of crashing the whole process at startup.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        return set(), set()
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return set(), set()
    sheet_names = set(wb.sheetnames)
    headers: set[str] = set()
    for ws in wb.worksheets:
        row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not row:
            continue
        for cell in row:
            if cell is not None and str(cell).strip():
                headers.add(str(cell).strip())
    wb.close()
    return sheet_names, headers


def _score_signature(
    sig: DatasetSignature, filename_norm: str, headers: set[str], sheets: set[str]
) -> tuple[float, list[str]]:
    if sig.required_headers and not all(h in headers for h in sig.required_headers):
        return 0.0, []
    if sig.required_sheets and not all(s in sheets for s in sig.required_sheets):
        return 0.0, []

    signals: list[str] = []
    score = 0.0

    filename_hit = any(hint in filename_norm for hint in sig.filename_hints)
    if filename_hit:
        score += 0.4
        signals.append("filename_match")

    if sig.required_headers:
        score += 0.15
        signals.append("required_headers_present")
    if sig.required_sheets:
        # Weighted higher than a couple of required headers: a multi-sheet
        # fingerprint like StoreTech/Markets/RFM is distinctive enough on
        # its own to recognize the source without needing the filename too.
        score += 0.4
        signals.append("required_sheets_present")

    if sig.distinctive_headers:
        hits = [h for h in sig.distinctive_headers if h in headers]
        if hits:
            score += 0.45 * (len(hits) / len(sig.distinctive_headers))
            signals.append(f"distinctive_headers:{','.join(hits)}")
    if sig.distinctive_sheets:
        hits = [s for s in sig.distinctive_sheets if s in sheets]
        if hits:
            score += 0.45 * (len(hits) / len(sig.distinctive_sheets))
            signals.append(f"distinctive_sheets:{','.join(hits)}")

    return min(score, 1.0), signals


def classify(path: str | Path) -> Classification:
    """Classify a structured file against the signature registry.

    Signatures are tried in registry order; the highest-scoring match wins.
    Falls back to `unknown_structured_dataset` (confidence 0.0) when nothing
    scores above zero — callers should treat that as "ask the operator,"
    never as licence to guess.
    """
    p = Path(path)
    suffix = p.suffix.lower()
    headers: set[str] = set()
    sheets: set[str] = set()
    if suffix == ".csv":
        headers = set(_read_headers(p))
    elif suffix in (".xlsx", ".xlsm"):
        sheets, headers = _read_xlsx_info(p)
    filename_norm = _normalize_filename(p)

    best: Classification | None = None
    for sig in DATASET_SIGNATURES:
        score, signals = _score_signature(sig, filename_norm, headers, sheets)
        if best is None or score > best.confidence:
            best = Classification(
                dataset_type=sig.dataset_type,
                confidence=score,
                purpose=sig.purpose,
                recommended_action=sig.recommended_action,
                ingest_script=sig.ingest_script,
                matched_signals=signals,
                already_handled_by=sig.already_handled_by,
            )

    if best is None or best.confidence <= 0.0:
        return Classification(
            dataset_type=UNKNOWN_TYPE,
            confidence=0.0,
            purpose="unclassified",
            recommended_action="ask_operator",
            ingest_script="",
        )
    return best


def classify_bytes(content: bytes, filename: str) -> Classification:
    """Classify upload bytes before they've been written to their inbox path.

    Server-side upload routing (`/ingest/upload`) has the raw bytes and a
    filename but no on-disk file yet; this spares callers from writing to a
    real destination just to find out what the file is before the routing
    decision is even made.
    """
    import tempfile

    safe_name = Path(filename).name or "upload.bin"
    with tempfile.TemporaryDirectory() as td:
        tmp_path = Path(td) / safe_name
        tmp_path.write_bytes(content)
        return classify(tmp_path)


def confidence_threshold(settings: dict | None = None) -> float:
    settings = settings if settings is not None else core.load_settings()
    return float(
        settings.get("structured_ingest", {}).get(
            "confidence_threshold", DEFAULT_CONFIDENCE_THRESHOLD
        )
    )


def should_auto_ingest(classification: Classification, settings: dict | None = None) -> bool:
    if not classification.ingest_script:
        return False
    return classification.confidence >= confidence_threshold(settings)


# ---------------------------------------------------------------------------
# Smoke test
# ---------------------------------------------------------------------------

def _write_tmp_csv(tmp_dir: Path, name: str, headers: list[str]) -> Path:
    path = tmp_dir / name
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerow(["Jane" for _ in headers])
    return path


def _smoke() -> bool:
    import tempfile

    errors: list[str] = []
    with tempfile.TemporaryDirectory() as td:
        tmp_dir = Path(td)

        hubspot_path = _write_tmp_csv(
            tmp_dir,
            "hubspot-crm-exports-all-contacts-2026-07-07.csv",
            ["First Name", "Last Name", "Email", "Company Name", "Job Title",
             "Contact owner", "Lead Status", "Marketing contact status"],
        )
        result = classify(hubspot_path)
        if result.dataset_type != "hubspot_crm_export":
            errors.append(f"expected hubspot_crm_export, got {result.dataset_type}")
        if result.confidence < DEFAULT_CONFIDENCE_THRESHOLD:
            errors.append(f"hubspot export confidence too low: {result.confidence}")

        linkedin_path = _write_tmp_csv(
            tmp_dir, "Connections.csv", ["First Name", "Last Name", "URL", "Email Address", "Connected On"],
        )
        result2 = classify(linkedin_path)
        if result2.dataset_type != "linkedin_connections_export":
            errors.append(f"expected linkedin_connections_export, got {result2.dataset_type}")

        generic_path = _write_tmp_csv(tmp_dir, "some_export.csv", ["First Name", "Last Name"])
        result3 = classify(generic_path)
        if should_auto_ingest(result3, settings={"structured_ingest": {}}):
            errors.append(f"bare name-only CSV should not clear the auto-ingest bar (confidence={result3.confidence})")

        empty_path = tmp_dir / "notes.txt"
        empty_path.write_text("just some prose, not a dataset")
        result4 = classify(empty_path)
        if result4.dataset_type != UNKNOWN_TYPE:
            errors.append(f"non-CSV prose file should classify unknown, got {result4.dataset_type}")

    if errors:
        for e in errors:
            print(f"  FAIL: {e}")
        return False
    print("dataset_classifier smoke: all checks passed")
    return True


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(description="RB structured-dataset classifier")
    parser.add_argument("--file", metavar="PATH", help="Classify a single file")
    parser.add_argument("--smoke", action="store_true", help="Run smoke tests")
    args = parser.parse_args()

    if args.smoke:
        sys.exit(0 if _smoke() else 1)

    if args.file:
        result = classify(args.file)
        print(json.dumps(result.to_dict(), indent=2))
        return

    parser.print_help()


if __name__ == "__main__":
    main()
