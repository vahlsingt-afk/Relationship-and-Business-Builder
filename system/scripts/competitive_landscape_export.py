#!/usr/bin/env python3
"""
competitive_landscape_export.py — RB-2026-09-01.

xlsx renderer for competitive_landscape.py's analysis engine -- same
Workbook/tab/metadata pattern as ecosystem_export.py (reused, not
reinvented). Original 4 tabs, plus 6 more added 2026-09-01 (Phase 1 of
Todd's competitive-intelligence design request -- schema/plumbing only,
never fabricated content):

  Market Share by Category   -- one section per TECH_STACK_CATEGORIES entry
  Battle Cards                -- one row per (category, competitor) with a
                                  real, sourced positioning + a computed
                                  delta line against Genius
  Brand x Category Grid       -- one row per tracked brand, one column per
                                  category -- the actual "fashioned after
                                  the canonical sheet" matrix
  Competitor Research Status  -- which of the tracked competitors have real
                                  battle-card content, which are stale, and
                                  real (never invented) new-competitor
                                  candidates
  RM Battle Cards             -- one row per (competitor, category) with a
                                  category_battle_cards entry -- the
                                  RM-facing plain-English content
  Category Coverage Plan      -- one row per valid battle-card category,
                                  real vendor/target counts
  Competitor Master           -- one summary row per tracked competitor
  Vendor Lookup                -- one row per tracked competitor, the real
                                  lookup key (vendor_entity_id, categories
                                  with a battle card on file)
  Evidence Ledger              -- every real evidence record across every
                                  tracked competitor's evidence.jsonl
  Research Queue                -- the real pending review-queue items from
                                  competitor_intelligence_review.py's daily
                                  scan -- never invented, editable landing
                                  zone for Todd's review

Usage:
    python3 competitive_landscape_export.py --output /path/to/file.xlsx
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

SCRIPTS_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPTS_DIR))
import rb_core as core  # noqa: E402
import ecosystem_intelligence as ei  # noqa: E402
import ecosystem_export as eco_export  # noqa: E402
import competitive_landscape as cl  # noqa: E402
import competitor_intelligence as compintel  # noqa: E402
import competitor_intelligence_common as cic  # noqa: E402
import xlsx_safety  # noqa: E402

from openpyxl import Workbook

EXPORTS_DIR = core.SYSTEM_DIR / "exports"

TAB_MARKET_SHARE = "Market Share by Category"
TAB_BATTLE_CARDS = "Battle Cards"
TAB_BRAND_GRID = "Brand x Category Grid"
TAB_RESEARCH_STATUS = "Competitor Research Status"
TAB_RM_BATTLE_CARDS = "RM Battle Cards"
TAB_CATEGORY_COVERAGE_PLAN = "Category Coverage Plan"
TAB_COMPETITOR_MASTER = "Competitor Master"
TAB_VENDOR_LOOKUP = "Vendor Lookup"
TAB_EVIDENCE_LEDGER = "Evidence Ledger"
TAB_RESEARCH_QUEUE = "Research Queue"


def _join_points(points: list[dict]) -> str:
    return " | ".join(p.get("point", "") for p in points if p.get("point"))


def _join_extended_entries(entries: list[dict]) -> str:
    """Renders a list of competitor_intelligence_common.extended_field()
    entries (products/strengths/weaknesses/vulnerabilities/key_customers/
    recent_news) as one delimited string, same join style as
    _join_points() above."""
    return " | ".join(
        (e.get("value") or "") + (f" (as of {e['as_of']})" if e.get("as_of") else "")
        for e in entries if e.get("value")
    )


def _safe_append(ws, row):
    """RB-SECURITY-2026-09-05: rows here can carry externally-influenced
    strings (a vendor/brand name, a battle-card positioning summary sourced
    from a press release, a research-queue reason string) into an exported
    workbook. Same formula/CSV-injection guard (CWE-1236) already applied
    to Blue Sheets, via the shared xlsx_safety.py primitive."""
    ws.append(xlsx_safety.sanitize_row(row))


def build_workbook(graph: dict | None = None) -> Workbook:
    graph = graph or ei._read_graph()
    by_id = ei._index_by_id(graph.get("entities") or [])
    metadata = eco_export.build_export_metadata(graph, export_type="internal")

    wb = Workbook()
    wb.remove(wb.active)

    # --- Market Share by Category ---
    ws = wb.create_sheet(TAB_MARKET_SHARE)
    _safe_append(ws, ["RB Competitive Tech-Stack Analysis — Market Share by Category"])
    _safe_append(ws, [])
    for key, label in (
        ("graph_version", "Graph version"), ("graph_hash", "Graph hash"),
        ("generated_at", "Generated at"), ("evidence_cutoff", "Evidence cutoff"),
    ):
        _safe_append(ws, [label, metadata.get(key)])
    _safe_append(ws, [])
    for category in cl.TECH_STACK_CATEGORIES:
        share = cl.category_market_share(graph, category, by_id=by_id)
        _safe_append(ws, [
            f"Category: {category}",
            f"{share['brands_with_known_vendor']} of {share['total_brands_tracked']} brands tracked "
            f"({share['coverage_pct']}% coverage)",
        ])
        if not share["vendors"]:
            _safe_append(ws, ["  (no confirmed vendor relationships in this category yet)"])
        else:
            _safe_append(ws, ["Rank", "Vendor", "Brand count", "Share of known %", "Share of universe %", "Genius family?"])
            for row in share["vendors"]:
                _safe_append(ws, [
                    row["rank"], row["vendor_name"], row["brand_count"],
                    row["share_of_known_pct"], row["share_of_universe_pct"],
                    "Yes" if row["is_genius_family"] else "",
                ])
        _safe_append(ws, [])

    # --- Battle Cards ---
    ws = wb.create_sheet(TAB_BATTLE_CARDS)
    bc_headers = [
        "category", "genius_brand_count", "genius_share_of_known_pct",
        "competitor", "competitor_brand_count", "competitor_share_of_known_pct",
        "delta_line", "has_competitor_profile", "positioning_summary", "todds_pov",
        "genius_advantages", "competitor_advantages", "recent_financial_health",
    ]
    _safe_append(ws, bc_headers)
    for category in cl.TECH_STACK_CATEGORIES:
        card = cl.battle_card_for_category(graph, category, by_id=by_id)
        for bc in card["battle_cards"]:
            # RB-2026-09-06: real read from the vendor's own public earnings
            # (see competitive_landscape._recent_financial_health_for_vendor)
            # -- honest "(no public earnings)" for the majority of vendors
            # that are privately held, never a fabricated placeholder.
            health = bc["recent_financial_health"]
            _safe_append(ws, [
                category, card["genius_brand_count"], card["genius_share_of_known_pct"],
                bc["vendor_name"], bc["brand_count"], bc["share_of_known_pct"],
                bc["delta_line"], "Yes" if bc["has_competitor_profile"] else "No",
                bc["positioning_summary"] or "(none on file)",
                bc["todds_pov"] or "(none on file)",
                _join_points(bc["genius_advantages"]) or "(none on file)",
                _join_points(bc["competitor_advantages"]) or "(none on file)",
                health["answer"] if health else "(no public earnings)",
            ])

    # --- Brand x Category Grid ---
    ws = wb.create_sheet(TAB_BRAND_GRID)
    grid_headers = ["brand_name"] + cl.TECH_STACK_CATEGORIES
    _safe_append(ws, grid_headers)
    for row in cl.brand_category_grid(graph, by_id=by_id):
        _safe_append(ws, [row["brand_name"]] + [row.get(c) for c in cl.TECH_STACK_CATEGORIES])

    # --- Competitor Research Status ---
    ws = wb.create_sheet(TAB_RESEARCH_STATUS)
    status = cl.competitor_research_status(graph, by_id=by_id)
    _safe_append(ws, ["Tracked Competitors"])
    _safe_append(ws, [
        "competitor_slug", "display_name", "positioning_filled", "competes_on_filled",
        "vs_genius_point_count", "last_evidence_date", "needs_research",
    ])
    for c in status["competitors"]:
        _safe_append(ws, [
            c["competitor_slug"], c["display_name"],
            "Yes" if c["positioning_filled"] else "No",
            "Yes" if c["competes_on_filled"] else "No",
            c["vs_genius_point_count"], c["last_evidence_date"] or "(none)",
            "Yes" if c["needs_research"] else "No",
        ])
    _safe_append(ws, [])
    _safe_append(ws, ["New Competitor Candidates (real vendors already tracked, no profile yet — confirm via createCompetitor)"])
    _safe_append(ws, ["vendor_id", "vendor_name", "primary_category"])
    for cand in status["new_competitor_candidates"]:
        _safe_append(ws, [cand["vendor_id"], cand["vendor_name"], cand["primary_category"]])

    reg = cic.load_registry()
    competitors: list[dict] = []
    for entry in reg.get("registry", []):
        slug = entry.get("competitor_slug", "")
        comp_path = cic.competitor_dir(slug) / "competitor.json"
        if comp_path.exists():
            competitors.append(cic.load_json(comp_path))

    # --- RM Battle Cards (2026-09-01) ---
    ws = wb.create_sheet(TAB_RM_BATTLE_CARDS)
    _safe_append(ws, [
        "competitor_slug", "display_name", "category", "status", "confidence_pct",
        "rm_plain_english_posture", "listen_for", "discovery_questions",
        "when_to_bring_todd_in", "red_flags", "evidence_ids", "source_urls", "last_validated",
    ])
    for comp in competitors:
        for category, card in sorted((comp.get("category_battle_cards") or {}).items()):
            _safe_append(ws, [
                comp.get("competitor_slug"), comp.get("display_name"), category,
                card.get("status"), card.get("confidence_pct"),
                card.get("rm_plain_english_posture") or "(none on file)",
                " | ".join(card.get("listen_for") or []) or "(none on file)",
                " | ".join(card.get("discovery_questions") or []) or "(none on file)",
                card.get("when_to_bring_todd_in") or "(none on file)",
                " | ".join(card.get("red_flags") or []) or "(none on file)",
                ", ".join(card.get("evidence_ids") or []),
                ", ".join(card.get("source_urls") or []),
                card.get("last_validated") or "(never validated)",
            ])

    # --- Category Coverage Plan (2026-09-01) ---
    ws = wb.create_sheet(TAB_CATEGORY_COVERAGE_PLAN)
    _safe_append(ws, ["category", "real_vendors_tracked", "competitors_with_battle_card", "coverage_status"])
    battle_card_counts: dict[str, int] = {}
    for comp in competitors:
        for category in (comp.get("category_battle_cards") or {}):
            battle_card_counts[category] = battle_card_counts.get(category, 0) + 1
    for category in sorted(compintel.valid_battle_card_categories()):
        if category in cl.TECH_STACK_CATEGORIES:
            real_vendor_count = len(cl.category_market_share(graph, category, by_id=by_id)["vendors"])
        else:
            real_vendor_count = None  # RM-only category, no per-brand data source yet
        bc_count = battle_card_counts.get(category, 0)
        if bc_count == 0:
            coverage_status = "Needs Research"
        elif bc_count < 5:
            coverage_status = "Partial Coverage"
        else:
            coverage_status = "Covered"
        _safe_append(ws, [category, real_vendor_count, bc_count, coverage_status])

    # --- Competitor Master (2026-09-01; extended 2026-09-25 with the
    # Ecosystem Lookup Tool's Phase B fields -- products/strengths/
    # weaknesses/vulnerabilities/key_customers/recent_news/trends. This
    # whole workbook is already Todd's internal-only competitive-
    # intelligence copy (todds_pov is already unredacted above), so unlike
    # Team Portal's shareable_extended_view(), weaknesses/vulnerabilities
    # are included here too -- cic.get_extended_profile() backfills any of
    # the 153 pre-project competitor.json files still missing these keys,
    # same read-time-only backward-compat this export already relies on
    # implicitly for every other field.) ---
    ws = wb.create_sheet(TAB_COMPETITOR_MASTER)
    status_by_slug = {c["competitor_slug"]: c for c in status["competitors"]}
    _safe_append(ws, [
        "competitor_slug", "display_name", "aliases", "primary_category", "competes_on",
        "positioning_summary", "todds_pov", "vs_genius_point_count",
        "category_battle_card_count", "last_evidence_date", "needs_research",
        "products", "strengths", "weaknesses", "vulnerabilities", "key_customers",
        "recent_news", "trends",
    ])
    for comp in competitors:
        slug = comp.get("competitor_slug")
        s = status_by_slug.get(slug, {})
        vs_genius = comp.get("vs_genius") or {}
        extended = cic.get_extended_profile(comp)
        trends = extended.get("trends") or {}
        _safe_append(ws, [
            slug, comp.get("display_name"), ", ".join(comp.get("aliases") or []),
            comp.get("primary_category"), ", ".join(sorted(comp.get("competes_on") or [])),
            comp.get("positioning_summary") or "(none on file)",
            comp.get("todds_pov") or "(none on file)",
            len(vs_genius.get("genius_advantages") or []) + len(vs_genius.get("competitor_advantages") or []),
            len(comp.get("category_battle_cards") or {}),
            comp.get("last_evidence_date") or "(none)",
            "Yes" if s.get("needs_research", True) else "No",
            _join_extended_entries(extended.get("products") or []) or "(none on file)",
            _join_extended_entries(extended.get("strengths") or []) or "(none on file)",
            _join_extended_entries(extended.get("weaknesses") or []) or "(none on file)",
            _join_extended_entries(extended.get("vulnerabilities") or []) or "(none on file)",
            _join_extended_entries(extended.get("key_customers") or []) or "(none on file)",
            _join_extended_entries(extended.get("recent_news") or []) or "(none on file)",
            trends.get("value") or "(none on file)",
        ])

    # --- Vendor Lookup (2026-09-01) ---
    ws = wb.create_sheet(TAB_VENDOR_LOOKUP)
    _safe_append(ws, ["competitor_slug", "display_name", "vendor_entity_id", "primary_category", "categories_with_battle_card"])
    for comp in competitors:
        _safe_append(ws, [
            comp.get("competitor_slug"), comp.get("display_name"), comp.get("vendor_entity_id"),
            comp.get("primary_category"),
            ", ".join(sorted((comp.get("category_battle_cards") or {}).keys())) or "(none yet)",
        ])

    # --- Evidence Ledger (2026-09-01) ---
    ws = wb.create_sheet(TAB_EVIDENCE_LEDGER)
    _safe_append(ws, ["competitor_slug", "evidence_id", "category", "summary", "source", "confidence", "logged_at"])
    for entry in reg.get("registry", []):
        slug = entry.get("competitor_slug", "")
        evidence_path = cic.competitor_dir(slug) / "evidence.jsonl"
        for record in cic.load_jsonl(evidence_path):
            _safe_append(ws, [
                slug, record.get("evidence_id"), record.get("category"),
                record.get("summary"), record.get("source"), record.get("confidence"),
                record.get("logged_at"),
            ])

    # --- Research Queue (2026-09-01) ---
    ws = wb.create_sheet(TAB_RESEARCH_QUEUE)
    _safe_append(ws, ["review_id", "competitor_slug", "category", "kind", "reason", "status", "queued_at", "resolution"])
    review_queue = cic.load_review_queue()
    for item in review_queue.get("pending_reviews", []):
        _safe_append(ws, [
            item.get("review_id"), item.get("competitor_slug"), item.get("category"),
            item.get("kind"), item.get("reason"), item.get("status"),
            item.get("queued_at"), item.get("resolution") or "",
        ])

    return wb


def export_workbook(*, output_path: Path | None = None, graph: dict | None = None) -> Path:
    wb = build_workbook(graph=graph)
    if output_path is None:
        from datetime import datetime, timezone
        EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
        tag = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
        output_path = EXPORTS_DIR / f"competitive-landscape-{tag}.xlsx"
    else:
        output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--output", type=Path, default=None)
    args = p.parse_args()
    out = export_workbook(output_path=args.output)
    print(f"Wrote {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
