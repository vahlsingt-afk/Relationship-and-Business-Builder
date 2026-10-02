from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "system/account_intelligence/Worldpay_Backbook_ABM_Meeting_Prep_2026-09-18.docx"
NAVY = "17365D"
BLUE = "2F75B5"
PALE = "EAF2F8"
GRAY = "F3F5F7"
BORDER = "D9D9D9"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table):
    tbl_pr = table._tbl.tblPr
    node = tbl_pr.find(qn("w:tblBorders"))
    if node is None:
        node = OxmlElement("w:tblBorders")
        tbl_pr.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), BORDER)
        node.append(tag)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, color="000000", size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for edge in ("top", "left", "bottom", "right"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:w"), "100")
        elem.set(qn("w:type"), "dxa")
        mar.append(elem)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t)
    repeat_header(t.rows[0])
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, "FFFFFF", 9)
        shade(t.rows[0].cells[i], NAVY)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value), False, "000000", 8.7)
            if ridx % 2:
                shade(cells[i], GRAY)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.add_run(item)


def numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.add_run(f"{index}.  ").bold = True
        p.add_run(item)


def heading(doc, text, level=1):
    doc.add_heading(text, level=level)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.68)
sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.75)
sec.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor.from_string("222222")
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.08
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(27)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
for name, size in (("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)):
    styles[name].font.name = "Aptos Display"
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
    styles[name].paragraph_format.space_before = Pt(12)
    styles[name].paragraph_format.space_after = Pt(5)
    styles[name].paragraph_format.keep_with_next = True
for name in ("List Bullet", "List Bullet 2", "List Number"):
    styles[name].font.name = "Aptos"
    styles[name].font.size = Pt(10.5)
    styles[name].paragraph_format.space_after = Pt(3)

header = sec.header.paragraphs[0]
header.text = "WORLDPAY BACKBOOK ABM MEETING PREP"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.name = "Aptos"
header.runs[0].font.size = Pt(8)
header.runs[0].font.bold = True
header.runs[0].font.color.rgb = RGBColor.from_string(NAVY)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("Prepared from RB live Worldpay intelligence  •  September 14 2026  •  Internal working document")
run.font.name = "Aptos"
run.font.size = Pt(8)
run.font.color.rgb = RGBColor.from_string("666666")

p = doc.add_paragraph(style="Title")
p.add_run("Worldpay Backbook ABM Meeting Prep")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run("Friday September 18 2026   |   9:00 to 10:00 AM Central")
r.bold = True
r.font.size = Pt(12)
r.font.color.rgb = RGBColor.from_string(NAVY)
doc.add_paragraph("Prepared for Todd Vahlsing for the Worldpay and Genius sales and marketing alignment discussion.")

heading(doc, "Purpose and recommendation", 1)
doc.add_paragraph("The objective is to agree on a disciplined 90 day program that converts Worldpay customer access into qualified Genius opportunities without disrupting active pursuits, customer relationships or protected channel motions.")
p = doc.add_paragraph()
r = p.add_run("Recommendation  ")
r.bold = True
p.add_run("Do not launch a conventional campaign against the entire Worldpay customer base. Use Worldpay relationship intelligence to identify the specific problem each customer may be willing to solve, then activate a narrowly tailored Genius motion around that problem.")
p = doc.add_paragraph()
r = p.add_run("Meeting thesis  ")
r.bold = True
p.add_run("Worldpay gives us access. RB gives us intelligence. ABM converts both into an account specific wedge. The first successful wedge earns the right to expand.")

heading(doc, "Meeting details", 2)
table(doc, ["Item", "Detail"], [
    ("Purpose", "Align Worldpay and Genius on a targeted cross sell program for Worldpay restaurant customers"),
    ("Organizer", "Claire Whelen"),
    ("Participants", "Claire Whelen, Tracy Gallimore, Aaron Thein, Todd Vahlsing, Erica Chivers, Bryan Bailey and Ryan Hildebrand"),
    ("Desired result", "Pilot cohort, account protections, RM sponsors, two campaign concepts, owners, deadlines and scorecard"),
], [1.45, 5.35])

heading(doc, "What the live portfolio says", 1)
doc.add_paragraph("RB’s live Worldpay Master Account Plan contains 30 ranked accounts across six RM portfolios, four P1 accounts, 28 RM or field intelligence updates, 33 evidence events and 10 unresolved conflicts. The scores identify where attention may be valuable. They do not establish ABM eligibility.")
table(doc, ["Account", "Current RB intelligence", "Correct treatment"], [
    ("Five Guys", "High strategic value, but the customer imposed a hard engagement gate after uncoordinated Genius outreach.", "No ABM. Only customer approved activity coordinated through Todd and Avery."),
    ("Church’s", "Jeff Caplin has DMB interest while Church’s finalizes a five year Qu agreement.", "Protect the active DMB discovery motion. No POS or broad backbook campaign."),
    ("Pollo Campero", "A DMB RFP and defined procurement workstream are already active.", "Manage as an opportunity. Suppress parallel ABM touches."),
], [1.15, 3.3, 2.35])

heading(doc, "ABM eligibility gate", 1)
doc.add_paragraph("An account enters the backbook ABM cohort only when every condition below is satisfied.")
numbered(doc, [
    "No active Genius opportunity, RFP, proposal or customer specific campaign already owns the motion.",
    "No customer engagement restriction or unresolved account ownership conflict exists.",
    "No dealer, reseller or channel protection blocks direct activation.",
    "The Worldpay RM explicitly sponsors the account and outreach route.",
    "RB has enough current stack and relationship intelligence to form a credible customer problem hypothesis.",
    "The proposed campaign complements rather than duplicates current sales activity.",
])
doc.add_paragraph("Failing the gate does not make an account unimportant. It routes the account to opportunity management, RM nurture, channel management or intelligence development instead of campaign activation.")

heading(doc, "Recommended account segmentation", 1)
heading(doc, "Wedge ready activation candidates", 2)
table(doc, ["Account", "Primary hypothesis", "Campaign and first step"], [
    ("GoTo Foods", "Multi brand complexity across DMB, CMS, payments and deployment despite an established Qu position.", "Simplify the multi brand estate. Map DMB ownership, incumbents and payment exposure with Brian Wood."),
    ("Steak ’n Shake", "Kiosk transition and NCR reporting problems create measurable operating friction.", "Fix the digital guest journey. Document vendors, failures, owners and remediation milestones."),
    ("Papa Johns", "PAR closes the near term POS door, but DMB, deployment and payment adjacencies may remain open.", "Modernize without reopening POS. Verify DMB lifecycle and PAR deployment dependencies."),
    ("Cracker Barrel", "The Oracle transition may leave operational, payment or integration gaps.", "Make the new core perform better. Confirm rollout status, FreedomPay scope and unresolved issues."),
    ("Little Caesars", "A technical support relationship may enable recovery and discovery around payments, DMB or franchise exceptions.", "Resolve the open issues first. Consolidate ownership and use the recovered relationship to learn."),
    ("Red Robin", "Improving conditions may make pay at table or payment funded modernization timely, but the trigger is unconfirmed.", "Validate the signal, Ziosk pain, NCR term and customer ownership before activation."),
], [1.15, 3.05, 2.6])

heading(doc, "POS and platform window", 2)
p = doc.add_paragraph()
r = p.add_run("Del Taco is the strongest platform oriented candidate. ")
r.bold = True
p.add_run("Its aging NCR estate, Worldpay extension activity and Yadav acquisition create a credible governance and consolidation window. The campaign should offer a post acquisition architecture workshop focused on multi brand governance, POS and payments simplification, shared services, reporting and franchise consistency.")
bullets(doc, [
    "Complete the Worldpay extension before broadening the motion.",
    "Consolidate all Global Payments and Genius outreach.",
    "Confirm whether a sourcing or vendor consolidation event actually exists.",
    "Map technology decision ownership under Yadav.",
    "Determine whether NCR replacement is genuinely open.",
])
p = doc.add_paragraph()
r = p.add_run("Intelligence stage POS candidates  ")
r.bold = True
p.add_run("Eat’n Park, Miller’s Ale House, LaRosa’s, Bob Evans and Bullritos should receive RM assisted research before campaign activation. An aging or fragmented estate is a hypothesis, not proof of an open decision.")

heading(doc, "Existing customer channel and active opportunity motions", 2)
table(doc, ["Account", "Route", "Required action"], [
    ("Subway", "Installed base verification and expansion", "Confirm Genius product, sites, deployment, sponsor, ownership, commitments and support gaps."),
    ("RDS five unit rebrand", "Channel pilot", "Define customer, authority, dates, economics, support, protections, success metrics and case study rights."),
    ("Freddy’s", "RDS led only", "Wait for a customer defined problem and channel approved route."),
    ("Slim Chickens", "RDS led only", "Use FSTEC relationship development to validate stakeholders and modular needs."),
    ("Choice Hotels", "Separate hospitality win back", "Refresh the RFP status and keep it outside the restaurant ABM pilot."),
], [1.35, 1.85, 3.6])

heading(doc, "Protected and no fly accounts", 2)
table(doc, ["Account", "Reason", "Permitted activity"], [
    ("Five Guys", "Customer hard gate after uncoordinated Genius outreach", "Todd and Avery coordinated relationship activity only"),
    ("Church’s", "Active DMB conversation and five year Qu renewal", "Maggie and Todd DMB discovery only"),
    ("Pollo Campero", "Active DMB RFP", "Existing opportunity and procurement workstream only"),
    ("Freddy’s and Slim Chickens", "Protected RDS route", "RDS sponsored activity only"),
    ("Wendy’s and Dutch Bros", "Existing Genius motions", "Established account team only"),
    ("Domino’s", "Entrenched proprietary estate with no defined Genius problem", "RM nurture and adjacency listening"),
    ("Firehouse", "Strategic Toast position with no verified opening", "Investigate uncovered layers or franchise exceptions only"),
    ("Barbara B Mann and Festiva", "Entity, scope and fit unresolved", "Research only"),
], [1.5, 3.0, 2.3])

heading(doc, "Campaign concepts", 1)
heading(doc, "What is the crack in the stack", 2)
doc.add_paragraph("This should be the umbrella campaign. Offer CIO, CTO, restaurant technology, operations and payments leaders a 30 minute assessment to identify one area where the current stack creates unnecessary cost, complexity, risk or operational friction. The customer receives a current state map, verified friction points, implications and a recommended next investigation. The call to action is an assessment, not a product demonstration.")

heading(doc, "Modernize without ripping everything out", 2)
doc.add_paragraph("Use this with accounts whose core POS is recently selected or strategically entrenched. Lead with DMB, drive thru technology, kiosks, hardware, deployment remediation, reporting, integrations or payment interoperability. The message is that Genius can solve a specific problem while preserving the customer’s selected core platform.")

heading(doc, "Fix the digital guest journey", 2)
doc.add_paragraph("Use this where kiosks, pay at table, digital ordering or channel fragmentation create operational pain. The assessment follows the order from origination through authorization, POS and KDS routing, reporting and reconciliation, then identifies where ownership breaks across multiple vendors.")

heading(doc, "Fund modernization through the payment relationship", 2)
doc.add_paragraph("For a tightly qualified subset, assess whether payment economics, avoided gateway costs, service consolidation or a transaction metered commercial structure could help fund modernization. Initially offer an economic assessment only. Do not promise payment funded hardware or per transaction software economics until finance, legal, product, billing and leadership approve the construct.")

heading(doc, "Protect the payment and data path", 2)
doc.add_paragraph("Use this where Toast, PAR Pay, Qu Pay, Olo Pay or other platform attached payment models threaten customer choice. Frame the value around transparency, token and transaction data control, interoperability, flexibility and reduced platform dependency—not around defending Worldpay revenue.")

heading(doc, "How Worldpay RMs participate", 1)
doc.add_paragraph("The Worldpay RM is the campaign’s intelligence and access layer. Before activation, each RM should answer:")
bullets(doc, [
    "How strong is the relationship and would the RM personally sponsor an introduction?",
    "Who owns restaurant technology and who owns payments?",
    "What has the customer complained about or asked for?",
    "What change event, contract, remodel, RFP or leadership shift exists?",
    "What is the incumbent stack and what must Genius avoid saying or doing?",
    "Is another sales, customer, reseller or channel motion already active?",
])
doc.add_paragraph("The operating sequence is Worldpay customer to RM intelligence to RB assessment to identified crack to personalized hypothesis to RM sponsored introduction to discovery to qualified wedge to expansion.")

heading(doc, "The one page ABM card", 1)
doc.add_paragraph("Every activated account should have one current card maintained through RB. It should include:")
bullets(doc, [
    "Account, Worldpay RM and activation status",
    "Relationship strength and customer stakeholders",
    "Current POS, payments, DMB, ordering, loyalty and adjacent stack",
    "Confirmed business catalyst and incumbent lifecycle",
    "Known pain, primary Genius wedge and personalized hypothesis",
    "Relevant proof point, approved outreach route and desired next commitment",
    "Suppression instructions, confidence and unresolved intelligence gaps",
])

doc.add_page_break()
heading(doc, "Recommended 90 day pilot", 1)
doc.add_paragraph("Start with eight to ten activated accounts. Do not activate the entire backbook.")
table(doc, ["Pilot status", "Accounts"], [
    ("High touch candidates", "Del Taco, GoTo Foods, Steak ’n Shake, Papa Johns, Cracker Barrel and Little Caesars"),
    ("Conditional after rapid qualification", "Red Robin, Eat’n Park, Miller’s Ale House and LaRosa’s or Bob Evans"),
    ("Explicitly excluded", "Five Guys, Church’s, Pollo Campero, Freddy’s, Slim Chickens, Wendy’s, Dutch Bros, Choice Hotels and any account without RM sponsorship"),
], [2.0, 4.8])

heading(doc, "First 30 days qualification", 2)
bullets(doc, ["Validate the cohort with every RM.", "Complete the ABM cards.", "Resolve ownership and suppression rules.", "Confirm one customer problem per account.", "Choose one campaign and one call to action per account.", "Develop customer specific messaging."])
heading(doc, "Days 31 through 60 activation", 2)
bullets(doc, ["RM makes or sponsors the introduction.", "Todd leads the operator oriented assessment.", "Marketing supplies one relevant personalized artifact.", "Sales records new intelligence and confirms or disproves the problem.", "RB re-scores the account and updates downstream plans."])
heading(doc, "Days 61 through 90 conversion", 2)
bullets(doc, ["Convert validated problems into discovery or scoped opportunities.", "Stop campaigns where the hypothesis is disproven.", "Expand only when the customer confirms the need.", "Capture reusable proof, objections and messaging patterns.", "Select the wedge that deserves broader scale."])

heading(doc, "Measurement", 1)
doc.add_paragraph("The primary scorecard should measure commercial learning and customer progress rather than clicks, opens or MQLs.")
table(doc, ["Leading indicators", "Commercial outcomes", "Governance signals"], [
    ("Accounts reviewed; completed cards; RM sponsorship; warm introductions", "Assessment meetings; verified problems; qualified wedges; opportunities and pipeline", "Hypotheses disproven; accounts suppressed; ownership conflicts resolved; time from signal to action"),
], [2.25, 2.3, 2.25])
doc.add_paragraph("Correctly eliminating an account or preventing harmful outreach is a positive pilot result.")

heading(doc, "Decisions to secure Friday", 1)
numbered(doc, [
    "Confirm that MAP priority and ABM eligibility are separate decisions.",
    "Name the owner of final activation and suppression decisions.",
    "Select the initial eight to ten accounts.",
    "Confirm active opportunities, channel protected accounts and customer no fly rules.",
    "Assign a Worldpay RM sponsor to every selected account.",
    "Agree on the minimum intelligence required before activation.",
    "Commission the first two campaign concepts and assign messaging approval.",
    "Choose where RM feedback and customer intelligence will be recorded.",
    "Establish the weekly sales, marketing and RM scorecard review.",
])

heading(doc, "Recommended first two campaigns", 1)
p = doc.add_paragraph()
r = p.add_run("1  What is the crack in the stack  ")
r.bold = True
p.add_run("An RM sponsored, account specific assessment that identifies one credible operational or economic problem.")
p = doc.add_paragraph()
r = p.add_run("2  Modernize without ripping everything out  ")
r.bold = True
p.add_run("A wedge campaign for DMB, kiosks, hardware, integrations, deployment health and payment control where POS displacement is inappropriate.")
doc.add_paragraph("Develop the payment funded technology campaign next, after the commercial structure and claims receive internal approval.")

heading(doc, "Suggested opening", 1)
doc.add_paragraph("I reviewed the Worldpay portfolio account by account. The opportunity is real, but the customer list itself is not the campaign. Some accounts are already in active pursuits, some are channel protected, and at least one has explicitly gated Genius engagement. My recommendation is a small, RM sponsored pilot. We use Worldpay’s access and account knowledge to identify one credible crack in each customer’s stack, then build a specific campaign around that problem. We measure introductions, discovery and qualified opportunities—not email engagement.")

heading(doc, "Desired meeting outcome", 1)
bullets(doc, [
    "An approved campaign operating model",
    "A preliminary eight to ten account cohort",
    "A named RM sponsor for every account",
    "A confirmed suppression and protection list",
    "Two campaign concepts assigned to marketing",
    "Owners and deadlines for ABM cards",
    "A weekly sales, marketing and RM review cadence",
])
p = doc.add_paragraph()
r = p.add_run("Bottom line  ")
r.bold = True
p.add_run("The addressable ABM market is not the Worldpay backbook. It is the set of accounts where Worldpay access, customer permission, current intelligence and a real business problem intersect.")

doc.save(OUT)
print(OUT)
