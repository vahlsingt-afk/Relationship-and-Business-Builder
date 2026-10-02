#!/usr/bin/env python3
"""ecosystem_export.py — Shareable/internal workbook exporter (RB Unified
Restaurant-Tech Graph, 2026-07-31), Phase 6.

Generates a multi-tab .xlsx workbook FROM system/ecosystem_intelligence.json
-- the graph is the only source of truth; this script never hand-maintains
export content separately, and never reads the original research workbook.
One codepath (build_workbook()) produces two projections of the same data:

  - "internal": every field, including strategic_note and per-relationship
    confidence rationale text -- for Todd's own use.
  - "shareable": public-evidence-only. strategic_note and confidence
    rationale are stripped; everything else (brand, vendor, category,
    deployment status, classification, sources) is public-safe evidence
    that already justified the graph's own claim.

Tabs (per the Unified Restaurant-Tech Graph request):
  Executive Summary, Canonical Restaurant Tech Stack, Vendor Customer List,
  Evidence Ledger, Macro Views, Update Playbook.

"Current win" eligibility (used to build the Vendor Customer List tab and the
Executive Summary's win count) reuses the exact rule Phase 2's workbook
migration tool already enforces (active status, not a historical/superseded
deployment_status, not an unconfirmed McDonald's-style working profile) --
see is_current_win() below, which reads eco._HISTORICAL_DEPLOYMENT_STATUSES
directly rather than re-encoding the list a second time.

Usage:
    python3 ecosystem_export.py --internal
    python3 ecosystem_export.py --shareable
    python3 ecosystem_export.py --shareable --output /path/to/file.xlsx
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook

SCRIPTS_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPTS_DIR))
import rb_core as core  # noqa: E402
import ecosystem_intelligence as eco  # noqa: E402
import brand_profile_common as bpc  # noqa: E402
import xlsx_safety  # noqa: E402

EXPORTS_DIR = core.SYSTEM_DIR / "exports"
SCHEMA_VERSION = "rb_ecosystem_intelligence_v1"

EXPORT_TYPES = ("internal", "shareable")

TAB_EXECUTIVE_SUMMARY = "Executive Summary"
TAB_TECH_STACK = "Canonical Restaurant Tech Stack"
TAB_VENDOR_CUSTOMER_LIST = "Vendor Customer List"
TAB_EVIDENCE_LEDGER = "Evidence Ledger"
TAB_MACRO_VIEWS = "Macro Views"
TAB_UPDATE_PLAYBOOK = "Update Playbook"
TAB_BRAND_PROFILES = "Brand Profiles"

_BRAND_PROFILE_HEADERS = [
    "brand_id", "brand_name", "parent_ownership", "hq_city_state", "founded_year",
    "synopsis", "trajectory_badge", "trajectory_value_pct", "trajectory_as_of_year",
    "total_units", "franchised_units", "company_owned_units", "franchisee_count",
    "leadership_confirmed", "recent_signals",
]


def _graph_path() -> Path:
    return core.SYSTEM_DIR / "ecosystem_intelligence.json"


def load_graph() -> dict:
    return json.loads(_graph_path().read_text(encoding="utf-8"))


def _graph_hash(graph: dict) -> str:
    payload = json.dumps(graph, sort_keys=True, default=str).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()[:16]


def _index_entities(graph: dict) -> dict[str, dict]:
    return {e.get("id"): e for e in graph.get("entities") or [] if e.get("id")}


def is_current_win(rel: dict) -> bool:
    """A relationship is an eligible "current win" only when it's active AND
    not a historical/superseded claim AND not an unconfirmed working profile
    (the McDonald's NEWPOS/QSRSoft/etc. safeguard) -- the identical rule
    Phase 2's migration tool uses to decide working_profile_not_promoted."""
    if rel.get("status") != "active":
        return False
    deployment_status = rel.get("deployment_status")
    if deployment_status in eco._HISTORICAL_DEPLOYMENT_STATUSES:
        return False
    if deployment_status == "working_profile_public_source_required":
        return False
    return True


def build_export_metadata(graph: dict, *, export_type: str) -> dict:
    return {
        "graph_version": graph.get("version"),
        "graph_contract": graph.get("contract"),
        "graph_hash": _graph_hash(graph),
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "evidence_cutoff": graph.get("last_updated"),
        "export_type": export_type,
        "schema_version": SCHEMA_VERSION,
    }


def _relationship_row(rel: dict, entities: dict[str, dict], *, shareable: bool) -> dict[str, Any]:
    brand = entities.get(rel.get("from_entity_id")) or {}
    vendor = entities.get(rel.get("to_entity_id")) or {}
    classification = rel.get("relationship_classification") or {}
    confidence = rel.get("confidence") or {}
    deployment = rel.get("deployment") or {}
    row: dict[str, Any] = {
        "brand": brand.get("name") or rel.get("from_entity_id"),
        "vendor": vendor.get("name") or rel.get("to_entity_id"),
        "category": rel.get("category"),
        "vendor_role": rel.get("vendor_role"),
        "status": rel.get("status"),
        "deployment_status": rel.get("deployment_status"),
        "deployment_stage": deployment.get("stage"),
        "deployed_units": deployment.get("deployed_units"),
        "penetration_pct": deployment.get("penetration_pct"),
        "classification_level": classification.get("level"),
        "classification_label": classification.get("level_name"),
        "confidence_level": confidence.get("level"),
        "ai_application": rel.get("ai_application"),
        "current_win": is_current_win(rel),
        "sources": ", ".join(rel.get("sources") or []),
        "updated_at": rel.get("updated_at"),
    }
    # Privacy filter: strategic_note is Todd's own sales-strategy commentary
    # (e.g. "relevant for vendors seeking integration access rather than
    # displacement") and confidence.rationale can carry internal review
    # framing -- neither belongs in a shareable export. Everything else on
    # this row is the public evidence that already justified the claim.
    if not shareable:
        row["strategic_note"] = rel.get("strategic_note") or ""
        row["confidence_rationale"] = confidence.get("rationale") or ""
    return row


def _brand_profile_row(entity: dict, graph: dict, *, shareable: bool) -> dict[str, Any]:
    """Ecosystem Lookup Tool (2026-09-25): the Company Profile fields
    (identity, synopsis, footprint, live-computed trajectory, leadership,
    recent_signals) for one brand, projected the same internal/shareable
    way as _relationship_row above -- reuses bpc.shareable_view() directly
    rather than re-deriving which leadership entries are safe to export, so
    this can never silently drift from what Team Portal itself shows."""
    profile = bpc.get_profile(entity["id"], graph=graph)
    view = bpc.shareable_view(profile) if shareable else profile
    identity = view.get("identity") or {}
    footprint = view.get("footprint") or {}
    trajectory = view.get("trajectory") or {}
    leadership = view.get("leadership") or {}
    signals = view.get("recent_signals") or []

    def _val(f: dict | None) -> Any:
        return (f or {}).get("value")

    def _join_leadership(entries: list[dict]) -> str:
        return "; ".join(
            (e.get("name") or "") + (f" ({e['title']})" if e.get("title") else "")
            for e in entries
        )

    def _join_signals(entries: list[dict]) -> str:
        return " | ".join(
            f"[{s.get('signal_type')}] {s.get('value')}" + (f" (as of {s['as_of']})" if s.get("as_of") else "")
            for s in entries
        )

    row: dict[str, Any] = {
        "brand_id": view.get("brand_id"),
        "brand_name": view.get("brand_name"),
        "parent_ownership": _val(identity.get("parent_ownership")),
        "hq_city_state": _val(identity.get("hq_city_state")),
        "founded_year": _val(identity.get("founded_year")),
        "synopsis": _val(view.get("synopsis")),
        "trajectory_badge": trajectory.get("badge"),
        "trajectory_value_pct": trajectory.get("value_pct"),
        "trajectory_as_of_year": trajectory.get("as_of_year"),
        "total_units": _val(footprint.get("total_units")),
        "franchised_units": _val(footprint.get("franchised_units")),
        "company_owned_units": _val(footprint.get("company_owned_units")),
        "franchisee_count": _val(footprint.get("franchisee_count")),
        "leadership_confirmed": _join_leadership(leadership.get("confirmed") or []),
        "recent_signals": _join_signals(signals),
    }
    # leadership.reported_unverified is exactly what bpc.shareable_view()
    # already excludes -- an unconfirmed claim about a real person is
    # riskier to hand outside Todd's own working view than a sourced fact,
    # per brand_profile_common.py's own documented reasoning. Only ever
    # present on the internal export, matching _relationship_row's
    # strategic_note/confidence_rationale convention above.
    if not shareable:
        row["leadership_reported_unverified"] = _join_leadership(leadership.get("reported_unverified") or [])
    return row


def _evidence_rows(rel: dict, entities: dict[str, dict]) -> list[dict[str, Any]]:
    brand = entities.get(rel.get("from_entity_id")) or {}
    vendor = entities.get(rel.get("to_entity_id")) or {}
    base = {"brand": brand.get("name"), "vendor": vendor.get("name"), "category": rel.get("category")}
    assertions = rel.get("source_assertions") or []
    if assertions:
        return [
            {
                **base,
                "source_id": a.get("source_id"),
                "url": a.get("url"),
                "title": a.get("title"),
                "publisher": a.get("publisher"),
                "discovered_at": a.get("discovered_at"),
                "posture": a.get("posture"),
            }
            for a in assertions
        ]
    # Fall back to the plain sources[] list when no structured assertion exists.
    return [{**base, "source_id": s, "url": "", "title": "", "publisher": "", "discovered_at": "", "posture": ""}
            for s in (rel.get("sources") or [])]


def _macro_counts(rows: list[dict[str, Any]]) -> dict[str, dict[str, int]]:
    by_category: dict[str, int] = {}
    by_classification: dict[str, int] = {}
    for row in rows:
        cat = row.get("category") or "unknown"
        by_category[cat] = by_category.get(cat, 0) + 1
        label = row.get("classification_label") or "unclassified"
        by_classification[label] = by_classification.get(label, 0) + 1
    return {"by_category": by_category, "by_classification": by_classification}


_UPDATE_PLAYBOOK_ROWS = [
    ("Daily refresh", "python3 system/scripts/refresh_sources.py --all",
     "Fetches earnings/trade-press signals and applies mutation-worthy ones to the graph."),
    ("Workbook reconciliation (dry-run)", "python3 system/scripts/ecosystem_intelligence.py migrate-workbook <path>",
     "Preview how a new research workbook would reconcile against the graph. Review before --confirm."),
    ("Workbook reconciliation (apply)", "python3 system/scripts/ecosystem_intelligence.py migrate-workbook <path> --confirm",
     "Apply a reviewed workbook migration to the live graph."),
    ("Schema validation", "python3 system/schemas/validate.py --ecosystem-only",
     "Confirm the graph is schema-valid. Runs automatically before every graph write."),
    ("Rebuild classifications", "python3 system/scripts/relationship_classification.py classify-all",
     "Recompute the 6-level relationship classification for every relationship."),
    ("Regenerate this export", "python3 system/scripts/ecosystem_export.py --shareable",
     "Regenerate the shareable workbook from the current graph state."),
]


def _safe_append(ws, row):
    """RB-SECURITY-2026-09-05: rows here can carry externally-influenced
    strings (an entity/brand name, a vendor claim). Same formula/CSV-
    injection guard (CWE-1236) already applied to Blue Sheets, via the
    shared xlsx_safety.py primitive -- this is a shareable export type too,
    the same class of exposure that fix was written for."""
    ws.append(xlsx_safety.sanitize_row(row))


def build_workbook(*, export_type: str, graph: dict | None = None) -> tuple[Workbook, dict]:
    if export_type not in EXPORT_TYPES:
        raise ValueError(f"export_type must be one of {EXPORT_TYPES}, got {export_type!r}")
    if graph is None:
        graph = load_graph()
    shareable = export_type == "shareable"
    entities = _index_entities(graph)
    metadata = build_export_metadata(graph, export_type=export_type)

    rows = [_relationship_row(rel, entities, shareable=shareable) for rel in (graph.get("relationships") or [])]
    current_wins = [r for r in rows if r["current_win"]]
    macro = _macro_counts(rows)

    wb = Workbook()
    wb.remove(wb.active)

    # --- Executive Summary ---
    ws = wb.create_sheet(TAB_EXECUTIVE_SUMMARY)
    _safe_append(ws, ["RB Restaurant-Tech Graph Export"])
    _safe_append(ws, [])
    for key, label in (
        ("export_type", "Export type"), ("graph_version", "Graph version"),
        ("graph_hash", "Graph hash"), ("generated_at", "Generated at"),
        ("evidence_cutoff", "Evidence cutoff"), ("schema_version", "Schema version"),
    ):
        _safe_append(ws, [label, metadata.get(key)])
    _safe_append(ws, [])
    _safe_append(ws, ["Total entities", len(graph.get("entities") or [])])
    _safe_append(ws, ["Total relationships", len(rows)])
    _safe_append(ws, ["Current wins (active, evidence-backed, not a working profile)", len(current_wins)])

    # --- Brand Profiles (2026-09-25, Ecosystem Lookup Tool) ---
    # One row per brand entity -- ALL of them, not just the ones with real
    # research on file. An honest "not yet researched" blank for most
    # fields on most rows is the intended state today (Todd's decision #1:
    # populate what's computable now, never fabricate the rest) -- this
    # tab is the offline mirror of exactly what Team Portal's Company
    # Profile card shows, including the coverage gap itself.
    ws = wb.create_sheet(TAB_BRAND_PROFILES)
    brand_entities = [e for e in (graph.get("entities") or []) if e.get("entity_type") == "brand"]
    profile_rows = [_brand_profile_row(e, graph, shareable=shareable) for e in brand_entities]
    bp_headers = list(profile_rows[0].keys()) if profile_rows else (
        _BRAND_PROFILE_HEADERS if shareable else _BRAND_PROFILE_HEADERS + ["leadership_reported_unverified"]
    )
    _safe_append(ws, bp_headers)
    for row in profile_rows:
        _safe_append(ws, [row.get(h) for h in bp_headers])

    # --- Canonical Restaurant Tech Stack ---
    ws = wb.create_sheet(TAB_TECH_STACK)
    headers = list(rows[0].keys()) if rows else [
        "brand", "vendor", "category", "vendor_role", "status", "deployment_status",
        "deployment_stage", "deployed_units", "penetration_pct", "classification_level",
        "classification_label", "confidence_level", "ai_application", "current_win",
        "sources", "updated_at",
    ]
    _safe_append(ws, headers)
    for row in rows:
        _safe_append(ws, [row.get(h) for h in headers])

    # --- Vendor Customer List (current wins only) ---
    ws = wb.create_sheet(TAB_VENDOR_CUSTOMER_LIST)
    vcl_headers = ["brand", "vendor", "category", "classification_label", "deployment_status", "ai_application"]
    _safe_append(ws, vcl_headers)
    for row in current_wins:
        _safe_append(ws, [row.get(h) for h in vcl_headers])

    # --- Evidence Ledger ---
    ws = wb.create_sheet(TAB_EVIDENCE_LEDGER)
    ev_headers = ["brand", "vendor", "category", "source_id", "url", "title", "publisher", "discovered_at", "posture"]
    _safe_append(ws, ev_headers)
    for rel in (graph.get("relationships") or []):
        for ev_row in _evidence_rows(rel, entities):
            _safe_append(ws, [ev_row.get(h) for h in ev_headers])

    # --- Macro Views ---
    ws = wb.create_sheet(TAB_MACRO_VIEWS)
    _safe_append(ws, ["By category"])
    _safe_append(ws, ["category", "count"])
    for cat, count in sorted(macro["by_category"].items()):
        _safe_append(ws, [cat, count])
    _safe_append(ws, [])
    _safe_append(ws, ["By classification level"])
    _safe_append(ws, ["classification_label", "count"])
    for label, count in sorted(macro["by_classification"].items()):
        _safe_append(ws, [label, count])
    _safe_append(ws, [])
    _safe_append(ws, ["Total relationships (reconciliation check)", len(rows)])

    # --- Update Playbook ---
    ws = wb.create_sheet(TAB_UPDATE_PLAYBOOK)
    _safe_append(ws, ["Action", "Command", "Notes"])
    for action, command, notes in _UPDATE_PLAYBOOK_ROWS:
        _safe_append(ws, [action, command, notes])

    return wb, metadata


def export_workbook(*, export_type: str, output_path: Path | None = None, graph: dict | None = None) -> dict:
    wb, metadata = build_workbook(export_type=export_type, graph=graph)
    if output_path is None:
        EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
        tag = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
        output_path = EXPORTS_DIR / f"ecosystem_export-{export_type}-{tag}.xlsx"
    else:
        output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return {"path": str(output_path), "metadata": metadata}


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__)
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--internal", action="store_true", help="Full internal export (includes strategic notes).")
    g.add_argument("--shareable", action="store_true", help="Privacy-filtered export safe to share externally.")
    p.add_argument("--output", metavar="PATH", help="Output .xlsx path. Defaults to system/exports/.")
    args = p.parse_args(argv)

    export_type = "internal" if args.internal else "shareable"
    result = export_workbook(
        export_type=export_type,
        output_path=Path(args.output) if args.output else None,
    )
    print(json.dumps(result, indent=2, default=str))
    return 0


if __name__ == "__main__":
    sys.exit(main())
