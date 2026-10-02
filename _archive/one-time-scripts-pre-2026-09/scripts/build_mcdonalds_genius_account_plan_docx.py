from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "system/account_intelligence/2026-08-03-mcdonalds-genius-account-plan.md"
OUTPUT = ROOT / "system/account_intelligence/McDonalds_Global_Payments_Genius_Capture_Plan_REVISED_2026-08-06.docx"

INK = "1E1E1E"
RED = "C51A2E"
GOLD = "FFBC0D"
BLUE = "244A67"
MUTED = "66717A"
LIGHT = "F3F5F7"
PALE_GOLD = "FFF7D6"
WHITE = "FFFFFF"


def rgb(hexstr):
    return RGBColor.from_string(hexstr)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D7DCE0", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[i]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = rgb(MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def new_decimal_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids or [0]) + 1
    num_id = max(num_ids or [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start"); start.set(qn("w:val"), "1"); lvl.append(start)
    num_fmt = OxmlElement("w:numFmt"); num_fmt.set(qn("w:val"), "decimal"); lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText"); lvl_text.set(qn("w:val"), "%1."); lvl.append(lvl_text)
    suff = OxmlElement("w:suff"); suff.set(qn("w:val"), "tab"); lvl.append(suff)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs"); tab = OxmlElement("w:tab"); tab.set(qn("w:val"), "num"); tab.set(qn("w:pos"), "605"); tabs.append(tab); ppr.append(tabs)
    ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "605"); ind.set(qn("w:hanging"), "288"); ppr.append(ind)
    lvl.append(ppr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id))
    abs_id = OxmlElement("w:abstractNumId"); abs_id.set(qn("w:val"), str(abstract_id)); num.append(abs_id)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0"); num_pr.append(ilvl)
    num = OxmlElement("w:numId"); num.set(qn("w:val"), str(num_id)); num_pr.append(num)
    ppr.append(num_pr)


def add_inline_runs(paragraph, text, size=10.2, color=INK):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        clean = part[2:-2] if bold else part[1:-1] if code else part
        run = paragraph.add_run(clean)
        run.font.name = "Arial"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(color)
        run.bold = bold
        if code:
            run.font.name = "Consolas"
            run.font.size = Pt(size - 0.5)


def add_table(doc, headers, rows):
    n = len(headers)
    widths_map = {
        2: [2500, 6860],
        3: [2200, 3500, 3660],
        4: [1750, 2400, 2600, 2610],
    }
    widths = widths_map.get(n, [9360 // n] * n)
    table = doc.add_table(rows=1, cols=n)
    hdr = table.rows[0]
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, BLUE)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_inline_runs(p, text, size=8.7, color=WHITE)
        for run in p.runs:
            run.bold = True
    set_repeat_table_header(hdr)
    for r_i, row_data in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row_data):
            cell = cells[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if r_i % 2:
                set_cell_shading(cell, LIGHT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            add_inline_runs(p, text, size=8.5)
    set_table_geometry(table, widths)
    set_table_borders(table)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_profile_blocks(doc, headers, rows):
    for row in rows:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        keep_with_next(p)
        add_inline_runs(p, row[0], size=11, color=BLUE)
        for run in p.runs:
            run.bold = True
        for label, value in zip(headers[1:], row[1:]):
            q = doc.add_paragraph()
            q.paragraph_format.left_indent = Inches(0.18)
            q.paragraph_format.first_line_indent = Inches(-0.18)
            q.paragraph_format.space_after = Pt(2)
            lead = q.add_run(f"{label}: ")
            lead.bold = True
            lead.font.name = "Arial"
            lead.font.size = Pt(9.2)
            lead.font.color.rgb = rgb(MUTED)
            add_inline_runs(q, value, size=9.2)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.82)
    sec.bottom_margin = Inches(0.78)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.10
    for style_name, size, color, before, after in (
        ("Heading 1", 16, RED, 15, 7),
        ("Heading 2", 12.8, BLUE, 11, 5),
        ("Heading 3", 11.2, BLUE, 8, 4),
    ):
        st = styles[style_name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = rgb(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    # Running header/footer.
    header = sec.header
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("McDONALD'S / GLOBAL PAYMENTS + GENIUS  |  STRATEGIC ACCOUNT PLAN")
    hr.font.name = "Arial"
    hr.font.size = Pt(8.5)
    hr.font.bold = True
    hr.font.color.rgb = rgb(MUTED)
    footer = sec.footer
    add_page_number(footer.paragraphs[0])

    # Customer-pack first page.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("STRATEGIC ACCOUNT PLAN")
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = rgb(RED)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("McDonald's / Global Payments + Genius")
    r.font.name = "Arial"
    r.font.size = Pt(27)
    r.font.bold = True
    r.font.color.rgb = rgb(INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("From CosMc's production proof to a core-estate and global pathway")
    r.font.name = "Arial"
    r.font.size = Pt(13)
    r.font.color.rgb = rgb(BLUE)

    meta = doc.add_table(rows=2, cols=2)
    meta_data = [
        ("OWNER", "Todd Vahlsing"),
        ("DATE", "August 6, 2026"),
        ("ACCOUNT", "McDonald's Corporation"),
        ("OFFER", "Unified commerce powered by the Genius restaurant technology portfolio"),
    ]
    for idx, (label, value) in enumerate(meta_data):
        cell = meta.cell(idx // 2, idx % 2)
        set_cell_shading(cell, LIGHT if idx != 1 else PALE_GOLD)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        rr = p.add_run(label + "\n")
        rr.font.name = "Arial"; rr.font.size = Pt(8); rr.font.bold = True; rr.font.color.rgb = rgb(MUTED)
        rr = p.add_run(value)
        rr.font.name = "Arial"; rr.font.size = Pt(10); rr.font.bold = True; rr.font.color.rgb = rgb(INK)
    set_table_geometry(meta, [4680, 4680])
    set_table_borders(meta, color="E1E4E7")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(3)
    rr = p.add_run("ACCOUNT THESIS")
    rr.font.name = "Arial"; rr.font.size = Pt(9); rr.font.bold = True; rr.font.color.rgb = rgb(RED)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(8)
    add_inline_runs(p, "Global Payments is already competitively selected and proven inside the McDonald's ecosystem through CosMc's. The pursuit is an expansion and transfer motion—not a cold vendor introduction.", size=12, color=INK)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(15)
    add_inline_runs(p, "Evidence boundary: CosMc's publicly establishes competitive selection, single-supplier trust, broad solution scope, and ecosystem credibility. It does not, by itself, establish global-scale readiness for the core McDonald's estate.", size=9.5, color=MUTED)

    # Skip markdown title and metadata already represented.
    i = next(idx for idx, value in enumerate(lines) if value.startswith("## 1."))
    active_num_id = None
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            active_num_id = None
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
            i += 1
            continue
        if line.startswith("> "):
            quote = []
            while i < len(lines) and lines[i].startswith(">"):
                quote.append(lines[i].lstrip("> "))
                i += 1
            table = doc.add_table(rows=1, cols=1)
            cell = table.cell(0, 0)
            set_cell_shading(cell, PALE_GOLD)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline_runs(p, " ".join(quote), size=10.3, color=INK)
            set_table_geometry(table, [9360])
            set_table_borders(table, color=GOLD, size="10")
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            parsed = [[c.strip() for c in row.strip().strip("|").split("|")] for row in block]
            headers = parsed[0]
            rows = [r for r in parsed[2:] if len(r) == len(headers)]
            if len(headers) >= 5:
                add_profile_blocks(doc, headers, rows)
            else:
                add_table(doc, headers, rows)
            continue
        m_num = re.match(r"^(\d+)\.\s+(.*)$", line)
        if m_num:
            if active_num_id is None:
                active_num_id = new_decimal_numbering(doc)
            p = doc.add_paragraph()
            apply_numbering(p, active_num_id)
            p.paragraph_format.left_indent = Inches(0.42)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            p.paragraph_format.space_after = Pt(4)
            add_inline_runs(p, m_num.group(2))
            i += 1
            continue
        active_num_id = None
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.42)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            p.paragraph_format.space_after = Pt(4)
            add_inline_runs(p, line[2:])
            i += 1
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        add_inline_runs(p, line)
        i += 1

    # Set core properties.
    doc.core_properties.title = "McDonald's / Global Payments + Genius Strategic Account Plan"
    doc.core_properties.subject = "CosMc's-to-McDonald's account expansion and relationship strategy"
    doc.core_properties.author = "Todd Vahlsing"
    doc.core_properties.keywords = "McDonald's, Genius, CosMc's, account plan, restaurant technology"
    doc.core_properties.comments = "Editable strategic account plan generated from the Relationship Builder baseline."

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
