#!/usr/bin/env python3
"""Build the RB productization decision memo."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "RB_PRODUCTIZATION_TODD_SPECIFIC_AUDIT_MEMO_2026-06-09.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 98, 108)
BLACK = RGBColor(0, 0, 0)
LIGHT = "F2F4F7"
CALLOUT = "E8EEF5"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.10

    header = section.header.paragraphs[0]
    header.text = "RELATIONSHIP BUILDER | PRODUCTIZATION"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = header.runs[0]
    hr.font.name = "Calibri"
    hr.font.size = Pt(9)
    hr.font.bold = True
    hr.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = footer.add_run("Confidential working memo | Page ")
    fr.font.name = "Calibri"
    fr.font.size = Pt(9)
    fr.font.color.rgb = MUTED
    add_page_number(footer)


def add_metadata(doc, rows):
    for label, value in rows:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{label}: ")
        r.bold = True
        r.font.size = Pt(10.5)
        v = p.add_run(value)
        v.font.size = Pt(10.5)


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, CALLOUT)
    set_cell_margins(cell, 140, 180, 140, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}: ")
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(widths[idx])
        set_cell_shading(cell, LIGHT)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        run.bold = True
        run.font.size = Pt(9.5)
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            cells[idx].width = Inches(widths[idx])
            set_cell_margins(cells[idx])
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, len(row_data) - 1) else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(value))
            run.font.size = Pt(9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build():
    doc = Document()
    style_document(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("DECISION MEMO")
    r.font.name = "Calibri"
    r.font.size = Pt(23)
    r.font.bold = True
    r.font.color.rgb = BLACK

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Productizing Relationship Builder Without Losing Its Personal Intelligence")
    r.font.size = Pt(14)
    r.font.color.rgb = MUTED

    add_metadata(doc, [
        ("To", "Todd Vahlsing and the RB build team"),
        ("From", "Codex"),
        ("Date", "June 9, 2026"),
        ("Re", "Architecture risk, Todd-specific assumption audit, and first productization workstream"),
        ("Decision", "Begin with a controlled catalog and tenant-separation program before adding major product surface area"),
    ])

    add_callout(
        doc,
        "Recommendation",
        "Treat the current system as a successful single-user reference implementation. Preserve its intelligence logic and Todd's private corpus, while moving identity, configuration, storage, scheduling, authentication, and connectors behind explicit user and tenant boundaries.",
    )

    doc.add_heading("Why This Matters", level=1)
    doc.add_paragraph(
        "The concern is valid: RB currently delivers meaningful value, but its operating environment still depends on Todd's Mac, accounts, profile, schedules, paths, credentials, and industry context. That creates a real risk that a second user would receive a renamed version of Todd's system rather than a genuinely personal RB instance."
    )
    doc.add_paragraph(
        "This is a productization problem, not evidence that the intelligence engine was built incorrectly. The system already contains the beginnings of the right architecture, including profile directories, an active-profile loader, and tests against person-specific schema fields. The next work should complete and enforce those boundaries."
    )

    doc.add_heading("The Critical Distinction", level=1)
    add_bullets(doc, [
        "Shared product logic should be user-neutral and reusable.",
        "Tenant configuration should hold identity, timezone, voice, boundaries, connected accounts, schedules, and opportunity context.",
        "Todd's private corpus should be preserved intact, but isolated as one tenant's data rather than reused as product defaults.",
        "Historical sprint documents and Todd-specific test traces should remain as provenance, but must not function as current universal contracts.",
    ])
    add_callout(
        doc,
        "Safety rule",
        "Do not globally replace or delete the word Todd. The repository also contains legitimate tenant history, test fixtures, and unrelated contacts named Todd. Every occurrence must be classified before action.",
    )

    doc.add_heading("Initial Evidence", level=1)
    doc.add_paragraph(
        "A targeted scan of shared Python runtime and API files found 119 candidate references across 30 files. This count covers Todd-specific dispositions, schema fields, identity values, profile paths, machine paths, and timezone assumptions. It is a candidate inventory, not a final defect count."
    )
    add_table(
        doc,
        ["Area", "Observed pattern", "Required direction", "Priority"],
        [
            ("Profile resolution", "Hard-coded todd_vahlsing paths", "Authenticated tenant/profile context", "P0"),
            ("Canonical schema", "why_this_matters_to_todd; new_to_todd_likely", "Neutral fields with compatibility readers", "P0"),
            ("User decisions", "ask_todd; needs_todd", "Neutral machine values; personalized rendering", "P1"),
            ("Self identity", "Todd name and email embedded in engines", "Profile and connector self aliases", "P0"),
            ("Host runtime", "Todd home paths, LaunchAgents, localhost topology", "Configurable roots and deployment adapters", "P0"),
            ("Authentication", "One shared static API key", "User/workspace authentication and scoped tokens", "P0"),
            ("State", "Shared JSON, cache, and database files", "Tenant-scoped storage abstraction", "P0"),
            ("Tests", "Todd/Chicago/restaurant-tech assumptions", "Cross-profile isolation suite", "P0"),
        ],
        [1.25, 2.0, 2.55, 0.7],
    )

    doc.add_heading("What Must Be Cataloged", level=1)
    add_bullets(doc, [
        "Names, email aliases, profile IDs, career details, BridgePoint language, industry assumptions, and drafting voice.",
        "Todd-specific schema and vocabulary embedded in APIs, prompts, canonical responses, and persistence.",
        "Absolute paths, local ports, application-support directories, Downloads folders, and macOS-only scheduling.",
        "Credentials, API-key handling, connector grants, and any secret copied into generated configuration.",
        "Shared state that lacks tenant, workspace, user, or profile ownership.",
        "Historical documents and generated tenant artifacts that must be preserved but removed from universal product instructions.",
    ])

    doc.add_heading("Recommended Architecture", level=1)
    add_bullets(doc, [
        "RB Cloud: authentication, tenant storage, relationship graph, intelligence jobs, audit records, and the Bridgepoint API.",
        "RB clients: web application and ChatGPT integration, both acting as authenticated clients of RB.",
        "Optional local companion: only for Apple or other sources that genuinely require local access.",
        "Connector layer: OAuth-based Google, Microsoft, and future providers plus explicit import workflows.",
        "Codex: engineering and administration tooling, not a runtime dependency for ordinary customers.",
    ])

    doc.add_heading("Work Sequence", level=1)
    for title, text in [
        ("1. Freeze the context contract", "Define tenant_id, workspace_id, user_id, profile_id, request context, job context, configuration roots, and secret ownership."),
        ("2. Build a repeatable catalog", "Scan shared runtime, API schemas, prompts, installation files, tests, tenant data, and historical artifacts as separate scopes. Classify each candidate before changing it."),
        ("3. Neutralize shared contracts", "Migrate Todd-named schema fields and dispositions using compatibility readers, neutral writers, receipts, and regression tests."),
        ("4. Remove identity and host coupling", "Inject profile, self aliases, timezone, paths, schedules, and connector configuration."),
        ("5. Prove tenant isolation", "Run the same workflows for two synthetic users in different industries and timezones. Fail on any cross-user language, data, path, or recommendation leakage."),
        ("6. Design product delivery", "Build cloud scheduling, OAuth connectors, onboarding, first-value delivery, upgrades, observability, and an optional local agent."),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(f"{title}. ")
        r.bold = True
        p.add_run(text)

    doc.add_heading("Immediate Assignment for Claude", level=1)
    add_bullets(doc, [
        "Create a line-level Todd-specific candidate catalog with classification, proposed action, risk, and status.",
        "Create an allowlist for Todd tenant data, historical records, intentional fixtures, and unrelated contacts named Todd.",
        "Draft the tenant-context and profile-resolution architecture.",
        "Propose the neutral schema migration map and compatibility window.",
        "Design a two-user isolation test plan.",
        "Stop for architecture review before moving persisted data or performing broad renames.",
    ])

    doc.add_heading("Decision", level=1)
    doc.add_paragraph(
        "Make the Todd-specific assumption audit the first productization workstream. Continue feature iteration only where it does not deepen single-user coupling. The objective is not to make RB less personal; it is to make personalization a property of each user's configuration and data rather than a hidden property of the shared codebase."
    )

    doc.add_paragraph()
    final = doc.add_paragraph()
    final.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = final.add_run(
        "RB should feel deeply personal because it is grounded for each user, not because the product is Todd's system with the names changed."
    )
    r.bold = True
    r.italic = True
    r.font.color.rgb = DARK_BLUE

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
