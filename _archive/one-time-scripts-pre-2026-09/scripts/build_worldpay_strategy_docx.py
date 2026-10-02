#!/usr/bin/env python3
"""Create the revised strategy by making restrained edits to the original DOCX."""
from copy import deepcopy
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path("/Users/toddvahlsing/Documents/Genius Info/Worldpay_Genius_Enterprise_Account_Development_Strategy_2026-07.docx")
OUTPUT = ROOT / "system/account_intelligence/Worldpay_Genius_Enterprise_Account_Development_Strategy_REVISED_2026-08-03.docx"


def find_paragraph(doc, startswith):
    return next(p for p in doc.paragraphs if p.text.startswith(startswith))


def replace_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def insert_after(paragraph, text, style=None, bold=False):
    new_p = deepcopy(paragraph._p)
    for child in list(new_p):
        if child.tag.endswith("}pPr"):
            continue
        new_p.remove(child)
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    result = Paragraph(new_p, paragraph._parent)
    if style:
        result.style = style
    run = result.add_run(text)
    run.bold = bold
    return result


def set_cell(cell, text):
    cell.text = text
    for p in cell.paragraphs:
        p.paragraph_format.keep_together = True


def add_formatted_row(table, values):
    row = table.add_row()
    template = table.rows[-2]
    for idx, value in enumerate(values):
        row.cells[idx]._tc.get_or_add_tcPr().clear()
        row.cells[idx]._tc.get_or_add_tcPr().extend(deepcopy(template.cells[idx]._tc.get_or_add_tcPr()))
        set_cell(row.cells[idx], value)
    return row


def build():
    doc = Document(SOURCE)
    doc.core_properties.title = "Worldpay + Genius Enterprise Account Development Strategy — Revised"
    doc.core_properties.author = "Todd Vahlsing"

    replace_text(find_paragraph(doc, "July 2026"), "Revised 3 August 2026  |  Internal discussion draft")

    operating = doc.tables[2]
    add_formatted_row(operating, [
        "8  REFINE",
        "Re-score the account and adjust the storyline, stakeholder plan, product posture, and resource level as facts change.",
        "Current priority and next-best action",
    ])

    enablement = find_paragraph(doc, "Todd joins the RM team meeting")
    p = insert_after(enablement, "Blue Sheet Adoption", "Heading 2")
    p = insert_after(p, "Start with a one-page minimum viable Blue Sheet and deepen it only when the relationship, opportunity, or risk warrants. The goal is a better decision or customer advance—not completion of a document.", "Normal")
    p = insert_after(p, "Events as Customer Access", "Heading 2")
    insert_after(p, "Treat FS/TEC, customer conferences, QBRs, SBRs, and Worldpay events as account-development channels. Identify the target stakeholder, RM-sponsored introduction, meeting objective, and follow-up. Attendance alone is not an outcome.", "Normal")

    blue = find_paragraph(doc, "The Blue Sheet is the proactive")
    replace_text(blue, "The Blue Sheet is the proactive Worldpay-Genius account plan. It is created before an active RFP or known pain point is required, begins as a one-page working plan, and expands only as evidence and customer movement warrant.")
    product_heading = find_paragraph(doc, "Product Strategy")
    product_heading.paragraph_format.page_break_before = False
    previous = product_heading._p.getprevious()
    if previous is not None and not "".join(previous.itertext()).strip():
        for node in previous.xpath('.//w:br[@w:type="page"]'):
            node.getparent().remove(node)
    for run in product_heading.runs:
        for child in list(run._r):
            if child.tag.endswith("}lastRenderedPageBreak"):
                run._r.remove(child)

    replace_text(find_paragraph(doc, "Build relationships across marketing/digital"),
        "Build relationships across marketing/digital, technology, operations, and payments while confirming the DMB environment and Qu commercial position. Treat Qu Pay as a potential Worldpay volume-defense issue; Ryan will help clarify internal ownership before broader outreach.")
    replace_text(find_paragraph(doc, "Confirm the current Global Payments/Genius engagement"),
        "Confirm the current Global Payments/Genius engagement, ownership, deployment status, and customer commitments. Subway uses a proprietary POS and is deploying Genius kitchen-screen software globally; build from that position rather than assuming a generic POS displacement motion.")
    replace_text(find_paragraph(doc, "Both brands are PAR customers"),
        "Both brands are PAR customers with strong RDS relationships. Identify the internal dealer owner and agree whether each motion is referral-led, jointly sold, or dealer-led before customer outreach.")

    measures = doc.tables[10]
    set_cell(measures.cell(3, 0), "Blue Sheets that produce a verified fact, relationship advance, disposition, or next commitment")
    set_cell(measures.cell(4, 0), "QBR/SBR and event meetings with a defined account outcome")

    last_priority = find_paragraph(doc, "Establish a simple scorecard")
    p = insert_after(last_priority, "Build the first FS/TEC and customer-event introduction plan.", "List Number")
    insert_after(p, "Begin a shared capability map for gateway, processor, fraud, dispute, gift, and other payment add-ons.", "List Number")

    set_cell(doc.tables[11].cell(0, 0), "Be proactive, thoughtful, and patient. Build relationship width before expecting a transaction—and do not mistake activity for achievement.")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
