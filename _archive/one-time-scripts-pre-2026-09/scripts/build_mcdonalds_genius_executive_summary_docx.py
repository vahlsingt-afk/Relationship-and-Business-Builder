from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_mcdonalds_genius_account_plan_docx import (
    BLUE, INK, LIGHT, MUTED, PALE_GOLD, RED,
    add_inline_runs, add_page_number, add_profile_blocks, add_table,
    apply_numbering, keep_with_next, new_decimal_numbering, rgb,
    set_cell_shading, set_table_borders, set_table_geometry,
)


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "system/account_intelligence/2026-08-04-mcdonalds-genius-executive-summary.md"
OUTPUT = ROOT / "system/account_intelligence/McDonalds_Global_Payments_Genius_Executive_Summary_REVISED_2026-08-06.docx"


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.82)
    sec.bottom_margin = Inches(0.78)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in (
        ("Heading 1", 16, RED, 15, 7),
        ("Heading 2", 12.8, BLUE, 11, 5),
        ("Heading 3", 11.2, BLUE, 8, 4),
    ):
        st = doc.styles[name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = rgb(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    hp = sec.header.paragraphs[0]
    hr = hp.add_run("McDONALD'S / GLOBAL PAYMENTS + GENIUS  |  EXECUTIVE SUMMARY")
    hr.font.name = "Arial"; hr.font.size = Pt(8.5); hr.font.bold = True; hr.font.color.rgb = rgb(MUTED)
    add_page_number(sec.footer.paragraphs[0])

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(4)
    r = p.add_run("EXECUTIVE SUMMARY"); r.font.name = "Arial"; r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = rgb(RED)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
    r = p.add_run("McDonald's / Global Payments + Genius"); r.font.name = "Arial"; r.font.size = Pt(25); r.font.bold = True; r.font.color.rgb = rgb(INK)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(13)
    r = p.add_run("Three contained approval paths into the core McDonald's estate"); r.font.name = "Arial"; r.font.size = Pt(13); r.font.color.rgb = rgb(BLUE)

    meta = doc.add_table(rows=2, cols=2)
    for idx, (label, value) in enumerate((("OWNER", "Todd Vahlsing"), ("DATE", "August 6, 2026"), ("ACCOUNT", "McDonald's Corporation"), ("PURPOSE", "Executive capture-plan summary"))):
        cell = meta.cell(idx // 2, idx % 2); set_cell_shading(cell, PALE_GOLD if idx == 1 else LIGHT)
        q = cell.paragraphs[0]; q.paragraph_format.space_after = Pt(1)
        a = q.add_run(label + "\n"); a.font.name = "Arial"; a.font.size = Pt(8); a.font.bold = True; a.font.color.rgb = rgb(MUTED)
        b = q.add_run(value); b.font.name = "Arial"; b.font.size = Pt(10); b.font.bold = True; b.font.color.rgb = rgb(INK)
    set_table_geometry(meta, [4680, 4680]); set_table_borders(meta, color="E1E4E7")

    i = next(idx for idx, value in enumerate(lines) if value.startswith("## Bottom line"))
    active_num_id = None
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            active_num_id = None; i += 1; continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1); i += 1; continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2); i += 1; continue
        if line.startswith("> "):
            quote = []
            while i < len(lines) and lines[i].startswith(">"):
                quote.append(lines[i].lstrip("> ")); i += 1
            table = doc.add_table(rows=1, cols=1); cell = table.cell(0, 0); set_cell_shading(cell, PALE_GOLD)
            q = cell.paragraphs[0]; q.paragraph_format.space_after = Pt(0); add_inline_runs(q, " ".join(quote), size=10.3)
            set_table_geometry(table, [9360]); set_table_borders(table, color="FFBC0D", size="10")
            doc.add_paragraph().paragraph_format.space_after = Pt(1); continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i]); i += 1
            parsed = [[c.strip() for c in row.strip().strip("|").split("|")] for row in block]
            headers, rows = parsed[0], [r for r in parsed[2:] if len(r) == len(parsed[0])]
            add_profile_blocks(doc, headers, rows) if len(headers) >= 5 else add_table(doc, headers, rows)
            continue
        m = re.match(r"^(\d+)\.\s+(.*)$", line)
        if m:
            if active_num_id is None: active_num_id = new_decimal_numbering(doc)
            p = doc.add_paragraph(); apply_numbering(p, active_num_id)
            p.paragraph_format.left_indent = Inches(0.42); p.paragraph_format.first_line_indent = Inches(-0.22); p.paragraph_format.space_after = Pt(4)
            add_inline_runs(p, m.group(2)); i += 1; continue
        active_num_id = None
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.42); p.paragraph_format.first_line_indent = Inches(-0.22); p.paragraph_format.space_after = Pt(4)
            add_inline_runs(p, line[2:]); i += 1; continue
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5); add_inline_runs(p, line); i += 1

    doc.core_properties.title = "McDonald's / Global Payments + Genius Executive Summary"
    doc.core_properties.subject = "Three approval paths: DMB, POS hardware, and camera-based drive-thru timing"
    doc.core_properties.author = "Todd Vahlsing"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
