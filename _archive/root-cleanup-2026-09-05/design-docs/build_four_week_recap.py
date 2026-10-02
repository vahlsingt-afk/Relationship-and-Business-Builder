from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = "Todd_Vahlsing_Four_Week_Genius_Experience_Recap.docx"
BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
GRAY = RGBColor(90, 98, 108)
BLACK = RGBColor(31, 31, 31)


def set_font(run, size=11, bold=False, color=BLACK, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_font(run, size=9, color=GRAY)


def add_meta(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    set_font(p.add_run(f"{label}: "), bold=True)
    set_font(p.add_run(value))


def add_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), bold=True)
        set_font(p.add_run(text[len(bold_lead):]))
    else:
        set_font(p.add_run(text))
    return p


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.78)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.88)
sec.right_margin = Inches(0.88)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.35)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(10.5)
normal.font.color.rgb = BLACK
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.06

h1 = doc.styles["Heading 1"]
h1.font.name = "Calibri"
h1._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
h1._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
h1.font.size = Pt(14)
h1.font.bold = True
h1.font.color.rgb = BLUE
h1.paragraph_format.space_before = Pt(10)
h1.paragraph_format.space_after = Pt(4)
h1.paragraph_format.keep_with_next = True

# Quiet running header and footer.
hp = sec.header.paragraphs[0]
hp.text = "GENIUS | FOUR-WEEK EXPERIENCE RECAP"
set_font(hp.runs[0], size=8.5, bold=True, color=GRAY)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(fp.add_run("Todd Vahlsing  •  "), size=9, color=GRAY)
add_page_number(fp)

# Memo masthead.
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
set_font(p.add_run("FOUR-WEEK EXPERIENCE RECAP"), size=22, bold=True, color=DARK)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(11)
set_font(p.add_run("Global Payments and the Genius Team"), size=14, color=GRAY)

add_meta(doc, "To", "Ryan Hildebrand")
add_meta(doc, "From", "Todd Vahlsing")
add_meta(doc, "Date", "August 3, 2026")
add_meta(doc, "Purpose", "Constructive feedback on the hiring and onboarding experience")

# Overall callout.
table = doc.add_table(rows=1, cols=1)
table.autofit = False
table.columns[0].width = Inches(6.5)
cell = table.cell(0, 0)
cell.width = Inches(6.5)
set_cell_shading(cell, "E8EEF5")
cell.margin_top = Inches(0.08)
cell.margin_bottom = Inches(0.08)
cell.margin_left = Inches(0.12)
cell.margin_right = Inches(0.12)
cp = cell.paragraphs[0]
cp.paragraph_format.space_after = Pt(0)
set_font(cp.add_run("Overall assessment. "), bold=True, color=DARK)
set_font(cp.add_run("The first four weeks have been positive. I feel confident that joining Genius was the right decision, and I am increasingly excited about what we can build."))

add_heading(doc, "What Worked Especially Well")
add_body(doc, "Connection during hiring. Ryan, Mike, and I connected quickly as restaurant people and industry veterans. The conversations felt natural and genuine, and I came away believing that my experience would be valued and that I had a place at Genius.", "Connection during hiring. ")
add_body(doc, "Leadership support. Ryan’s morning and evening check-ins during the first two weeks were especially valuable. They provided access, direction, and confidence while I was finding my footing.", "Leadership support. ")
add_body(doc, "A welcoming team. People across the organization have been open and willing to help. Dale McKee has been especially helpful, and my early interaction with Nick was valuable. The team has also been receptive to the outside perspective I bring.", "A welcoming team. ")
add_body(doc, "A solid learning foundation. The online Genius curriculum was useful, and the live demonstrations and product sessions have been even more effective because they provide context, allow questions, and connect capabilities to customer needs.", "A solid learning foundation. ")

add_heading(doc, "What Was Different Than Expected")
add_body(doc, "The role is more fluid than I anticipated. That creates some ambiguity, but it also gives me room to contribute beyond a narrowly defined box and begin having a broader organizational impact.")
add_body(doc, "I expected a larger Worldpay portfolio. My initial understanding was approximately 90 accounts; the actionable portfolio appears closer to 25 solid accounts. That remains enough to build from, but it is materially different from the original expectation.")
add_body(doc, "I understood that parts of the role were aspirational, and I remain committed to helping build them. FSTEC is emerging as an important milestone for advancing the Worldpay cross-sell motion, and I have a good understanding of what success should look like over the next 30–90 days.")

add_heading(doc, "Where the Most Friction Exists")
add_body(doc, "The largest day-to-day challenge has been navigating the different logins, email environments, SSO requirements, and mobile-device policies. This continues to consume more time and energy than it should.")
add_body(doc, "The organization also relies heavily on tribal knowledge instead of consistently documented and maintained CRM information. People are helpful, but finding information often depends on knowing whom to ask. Silos and differing understandings of the restaurant marketplace create additional friction.")

doc.add_page_break()
add_heading(doc, "Practical Improvements for Future Hires")
items = [
    ("Create a tool map.", "Give each new hire a simple guide to the systems, what each one is used for, where to log in, which identity or email is required, and whom to contact for help."),
    ("Curate the marketing library.", "Organize the strongest and most current materials so a new hire can find the right asset quickly and understand when to use it."),
    ("Preload the calendar.", "Place essential team meetings, training sessions, and customer-partner events on the new hire’s calendar at the outset."),
    ("Prioritize live product learning.", "Pair the online curriculum with live demonstrations, product training, and opportunities to ask questions of experienced team members."),
    ("Keep introductions role-specific and resilient.", "Use a role-based introduction plan that can be adjusted quickly when personnel or responsibilities change."),
]
for label, detail in items:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run("• "), bold=True, color=BLUE)
    set_font(p.add_run(label + " "), bold=True)
    set_font(p.add_run(detail))

add_heading(doc, "Looking Ahead")
add_body(doc, "In many ways, this reminds me of PAR when I joined in 2016: meaningful opportunity, an evolving organization, and the chance to help shape how the company goes to market. The important difference is that Global Payments brings the backing and mindset to overcome many of the challenges that growing restaurant-technology businesses typically face.")
add_body(doc, "None of these observations will be surprising. My intent is to provide a candid view of the onboarding experience so we can preserve what worked and make the path easier for future hires. I appreciate the support from Ryan and the team, and I am excited about what comes next.")

doc.core_properties.title = "Four-Week Genius Experience Recap"
doc.core_properties.subject = "Hiring and onboarding feedback for Ryan Hildebrand"
doc.core_properties.author = "Todd Vahlsing"
doc.core_properties.keywords = "Genius, Global Payments, onboarding, four-week recap"
doc.save(OUT)
print(OUT)
